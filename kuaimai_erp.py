#!/usr/bin/env python3
"""
快麦 ERP 基础资料自动编辑程序。

读取产品文件夹中的产品信息.xlsx 和图片子目录，使用 Playwright 操作快麦通。
商品查询优先调用页面同源 API，页面导航、文件上传和表单校验使用 DOM。
"""

from __future__ import annotations

import argparse
import asyncio
import base64
import getpass
import json
import logging
import os
import re
import signal
import shutil
import subprocess
import sys
import tempfile
import time
import uuid
from collections.abc import Mapping
from contextlib import asynccontextmanager
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple
from urllib.parse import unquote, unquote_to_bytes, urljoin, urlsplit

from category_profile import category_profile
from douyin_data import DouyinAssets, DouyinDataError, DouyinFields, field_lookup, parse_douyin_fields, read_douyin_assets
from douyin_listing import DouyinListing, DouyinListingError
from jd_data import JdFields, parse_jd_fields
from jd_form_listing import JdFormListing, JdFormListingError
from learning_models import (
    ProductFingerprint,
    RunCheckpoint,
    checkpoint_event_payload,
    StageResult,
    canonical_sha256,
)
from learning_store import LearningStore
from learning_assets import PreparedLearningAsset, prepare_learning_assets
from learning_client import CloudLearningClient, flush_outbox_async
from attribute_runtime import AttributeRuntime, ReviewBatchRequired, ReviewRequired
from historical_readbacks import load_verified_attribute_history
from review_resume import persist_resume_and_ack, validate_resume
from pdd_data import PddFields, parse_pdd_fields
from pdd_form_listing import PddFormListing, PddFormListingError
from platform_discovery import PlatformDiscoveryError
from platform_inspection import RedactingFormatter, SensitiveLogRedactor, run_platform_inspection
from platform_registry import (
    PlatformSpec,
    expand_platform_selection,
    get_platform_spec,
    normalize_platform_selection,
)
from size_image_recognition import (
    ClothingSkuRecommendation,
    RecognitionError,
    SizeLength,
    SkuRecommendation,
    recognize_clothing_recommendations,
    recognize_recommendations,
    recognize_size_names,
    recognize_size_lengths,
)
from taobao_data import TaobaoFields, parse_taobao_fields
from taobao_listing import TaobaoListing, TaobaoListingError
from tmall_data import TmallAssets, TmallDataError, TmallFields, parse_tmall_fields, read_tmall_assets
from tmall_form_listing import (
    TmallFormListing,
    TmallFormListingError,
    TmallProductWriteRequired,
)
from tmall_api_index import TMALL_READ_ONLY_ENDPOINTS, TmallApiJsonIndex
from tmall_size_sources import (
    TmallSizeReviewRequired,
    TmallSizeSourceError,
    resolve_tmall_size_sources,
)
from xhs_data import XhsDataError, XhsFields, parse_xhs_fields
from xhs_form_listing import XhsFormListing, XhsFormListingError
from youzan_data import YouzanDataError, YouzanFields, parse_youzan_fields
from youzan_form_listing import YouzanFormListing, YouzanFormListingError
from wxsph_data import WxsphFields, parse_wxsph_fields
from wxsph_form_listing import WxsphFormListing, WxsphFormListingError


ERP_ENTRY_URL = "https://erp.superboss.cc/index.html#/index/"
CENTER_URL = "https://scm.superboss.cc/supplier/prod/center"
# 快麦现在会按账号/环境把登录页跳到 erp、erpa 或 viperp，
# 商品中心则可能使用 scm、scma 或 vipscm。这些都属于同一套快麦会话，
# 不应因域名切换而把已登录页面误判为仍在登录。
ERP_HOSTS = {"erp.superboss.cc", "erpa.superboss.cc", "viperp.superboss.cc"}
SCM_HOSTS = {"scm.superboss.cc", "scma.superboss.cc", "vipscm.superboss.cc"}
ERP_LOGIN_KEYCHAIN_SERVICE = "kuaimai-erp-auto-login"
ERP_LOGIN_KEYCHAIN_LABEL = "快麦 ERP 自动登录"
ERP_LOGIN_CREDENTIAL_RETRY_SECONDS = 30.0
DEFAULT_EXCEL_URL = "smb://gongxiang/共享文件/谭/products/战浮+NGBL-2064/产品信息.xlsx"
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"}
NEW_PRODUCT_SIZES = ("S", "M", "L", "XL", "2XL")

_ERP_LOGIN_KEYCHAIN_PROMPTED_ACCOUNTS: set[str] = set()
_ERP_LOGIN_CREDENTIAL_RETRY_AT: Dict[str, float] = {}
_ERP_LOGIN_MISSING_CREDENTIAL_WARNED: set[str] = set()
_ERP_LOGIN_PENDING_KEYCHAIN_ACCOUNTS: set[str] = set()

ERP_LOGIN_REJECTION_PATTERN = re.compile(
    r"账号或密码(?:错误|不正确)|用户名或密码(?:错误|不正确)|"
    r"密码错误|登录失败|登录名或密码错误|账户或密码错误"
)


def shared_runtime_root(script_dir: Path) -> Path:
    """让隔离 worktree 与主项目共用登录配置，避免每个 worktree 重登。"""
    for parent in (script_dir, *script_dir.parents):
        if parent.name == ".worktrees":
            return parent.parent
    return script_dir


SCRIPT_DIR = Path(__file__).resolve().parent
RUNTIME_ROOT = shared_runtime_root(SCRIPT_DIR)
DEFAULT_PROFILE_DIR = RUNTIME_ROOT / "output/kuaimai/chrome-profile"
DEFAULT_AUTH_STATE = RUNTIME_ROOT / "output/kuaimai/auth-state.json"
DEFAULT_DOUYIN_PUBLISH_SHOPS = (
    "钊叔 NEIGBORL 制",
    "啊亮穿搭",
    "老朱和NEIGBORL",
    "泰美了穿搭",
    "NEIGBORL钊哥小店",
    "小马客服",
)
DEFAULT_TAOBAO_PUBLISH_SHOPS = ("钊叔制",)
DEFAULT_TMALL_PUBLISH_SHOPS = ("NEIGBORL批判家专卖店",)
DEFAULT_PDD_PUBLISH_SHOPS = ("NEIGBORL鞋服旗舰店",)
DEFAULT_WXSPH_PUBLISH_SHOPS = (
    "NEIGBORL钊叔制鞋服",
    "NEIGBORL钊叔制造局",
)
DEFAULT_XHS_PUBLISH_SHOPS = (
    "钊叔的店",
    "钊叔制NEIGBORL的店",
    "NEIGBORL钊叔劳伦的店",
)
# These aliases are read-only name normalization.  A shop is selected only
# when its canonical name is also present in DEFAULT_XHS_PUBLISH_SHOPS.
XHS_PUBLISH_SHOP_ALIASES = {
    "啊亮熟NEIGBORL的店": "啊亮製NEIGBORL的店",
    "NEIGBORL钊叔旁伦的店": "NEIGBORL钊叔劳伦的店",
}
DEFAULT_YOUZAN_PUBLISH_SHOPS = ("NEIGBORL官方旗舰店",)
DEFAULT_JD_PUBLISH_SHOPS = ("NEIGBORL服饰旗舰店",)
COMMERCE_PUBLISH_TARGETS = {
    "taobao": ("淘宝", DEFAULT_TAOBAO_PUBLISH_SHOPS),
    "tmall": ("天猫", DEFAULT_TMALL_PUBLISH_SHOPS),
    "pdd": ("拼多多", DEFAULT_PDD_PUBLISH_SHOPS),
    "wxsph": ("微信小店（视频号）", DEFAULT_WXSPH_PUBLISH_SHOPS),
    "xhs": ("小红书", DEFAULT_XHS_PUBLISH_SHOPS),
    "youzan": ("有赞", DEFAULT_YOUZAN_PUBLISH_SHOPS),
    "jd": ("京东", DEFAULT_JD_PUBLISH_SHOPS),
}
KNOWN_DOUYIN_SHOPS = (
    "钊叔 NEIGBORL 制",
    "夏一制",
    "啊亮穿搭",
    "老朱和NEIGBORL",
    "泰美了穿搭",
    "杰叔抖店",
    "Li Napping",
    "NEIGBORL钊哥小店",
    "小马客服",
)

# 只有这些抖音店铺按产品 Excel 的“运费设置”匹配模板；
# 其他授权店铺统一使用“包邮”。
DOUYIN_FREIGHT_TEMPLATE_SHOPS = (
    "钊叔 NEIGBORL 制",
    "啊亮穿搭",
    "泰美了穿搭",
    "老朱和NEIGBORL",
    "NEIGBORL钊哥小店",
)


class AutomationError(RuntimeError):
    """可向用户直接展示的自动化异常。"""


def is_scm_url(value: str) -> bool:
    return urlsplit(value).hostname in SCM_HOSTS


def cdp_url_for_profile_processes(profile_dir: Path, process_text: str) -> Optional[str]:
    """从进程列表中找出占用指定专用配置目录的可接管 Chrome。"""
    profile_flag = f"--user-data-dir={profile_dir}"
    for line in process_text.splitlines():
        if profile_flag not in line:
            continue
        match = re.search(r"--remote-debugging-port(?:=|\s+)(\d+)", line)
        if match:
            return f"http://127.0.0.1:{match.group(1)}"
    return None


def discover_existing_profile_cdp_url(profile_dir: Path) -> Optional[str]:
    """仅接管当前脚本专用配置目录，不碰用户的日常 Chrome。"""
    try:
        result = subprocess.run(
            ["ps", "-axo", "command="],
            check=False,
            capture_output=True,
            text=True,
        )
    except OSError:
        return None
    if result.returncode != 0:
        return None
    return cdp_url_for_profile_processes(profile_dir, result.stdout)


def browser_pids_for_profile_processes(
    profile_dir: Path, process_text: str
) -> Tuple[int, ...]:
    """找出占用脚本专用配置目录的 Chrome 主进程，不匹配日常 Chrome。"""
    profile_flag = f"--user-data-dir={profile_dir}"
    pids = []
    for line in process_text.splitlines():
        stripped = line.strip()
        if profile_flag not in stripped or "--type=" in stripped:
            continue
        match = re.match(r"(\d+)\s+(.+)$", stripped)
        if not match or "Google Chrome" not in match.group(2):
            continue
        pids.append(int(match.group(1)))
    return tuple(dict.fromkeys(pids))


async def close_stale_profile_browsers(
    profile_dir: Path,
    logger: logging.Logger,
    *,
    timeout_seconds: float = 5.0,
) -> None:
    """关闭上次遗留的专用 Chrome，让本次运行获得全新的受控窗口。"""
    try:
        result = subprocess.run(
            ["ps", "-axo", "pid=,command="],
            check=False,
            capture_output=True,
            text=True,
        )
    except OSError:
        return
    if result.returncode != 0:
        return
    pids = browser_pids_for_profile_processes(profile_dir, result.stdout)
    if not pids:
        return

    logger.info("检测到上次遗留的专用 Chrome，正在关闭后启动新窗口")
    for pid in pids:
        try:
            os.kill(pid, signal.SIGTERM)
        except ProcessLookupError:
            pass
    deadline = asyncio.get_running_loop().time() + timeout_seconds
    remaining = set(pids)
    while remaining and asyncio.get_running_loop().time() < deadline:
        for pid in tuple(remaining):
            try:
                os.kill(pid, 0)
            except ProcessLookupError:
                remaining.discard(pid)
        if remaining:
            await asyncio.sleep(0.1)
    for pid in remaining:
        try:
            os.kill(pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
    logger.info("旧专用 Chrome 已关闭，本次将使用新窗口")


def center_url_for(value: str) -> str:
    parsed = urlsplit(value)
    if parsed.hostname in SCM_HOSTS:
        return f"{parsed.scheme}://{parsed.netloc}/supplier/prod/center"
    return CENTER_URL


def center_navigation_url_for(value: str) -> str:
    """首次从供应商首页进入商品中心时保留站点要求的 Cookie 标记。"""
    return f"{center_url_for(value)}?hasCookie=true"


def is_scm_cookie(cookie: Dict[str, Any]) -> bool:
    domain_label = (cookie.get("domain") or "").lstrip(".").split(".", 1)[0].casefold()
    return "scm" in domain_label or "scm" in (cookie.get("name") or "").casefold()


@dataclass(frozen=True)
class ProductData:
    excel_path: Path
    product_dir: Path
    title: str
    style_code: str
    base_price: str
    main_images: List[Path]
    main_images_34: List[Path]
    detail_images: List[Path]
    sku_images: List[Path]
    douyin_fields: Optional[DouyinFields] = None
    douyin_assets: Optional[DouyinAssets] = None
    taobao_fields: Optional[TaobaoFields] = None
    tmall_fields: Optional[TmallFields] = None
    pdd_fields: Optional[PddFields] = None
    wxsph_fields: Optional[WxsphFields] = None
    xhs_fields: Optional[XhsFields] = None
    youzan_fields: Optional[YouzanFields] = None
    jd_fields: Optional[JdFields] = None
    category_hints: Tuple[str, ...] = ()
    garment_kind: str = "generic"
    derived_size_names: Tuple[str, ...] = ()
    colors: Tuple[str, ...] = ()

    def __post_init__(self) -> None:
        if (self.douyin_fields is None) != (self.douyin_assets is None):
            raise ValueError("抖音字段与素材必须成对提供或同时省略")


@dataclass
class LearningRunContext:
    store: LearningStore
    run_id: str
    product_version: str
    device_id: str
    image_version: str
    client: Optional[CloudLearningClient] = None
    attribute_runtime: Optional[AttributeRuntime] = None
    initialized: bool = False


def create_learning_context(
    args: argparse.Namespace,
    product: ProductData,
) -> Optional[LearningRunContext]:
    if not getattr(args, "learning_enabled", False):
        return None
    fingerprint = ProductFingerprint.from_inputs(
        product.style_code,
        product.title,
        tuple(product.main_images)
        + tuple(product.main_images_34)
        + tuple(product.detail_images),
    )
    store = LearningStore(Path(args.learning_db))
    store.migrate()
    store.upsert_product(fingerprint)
    store.enqueue(
        f"product.upsert:{fingerprint.product_version}",
        "product.upsert",
        {
            "product_version": fingerprint.product_version,
            "style_code": fingerprint.style_code,
            "title": fingerprint.title,
            "category_json": {"hints": list(getattr(product, "category_hints", ()))},
        },
    )
    run_id = uuid.uuid4().hex
    client: Optional[CloudLearningClient] = None
    runtime: Optional[AttributeRuntime] = None
    if str(getattr(args, "learning_api_url", "")).strip():
        client = CloudLearningClient(args.learning_api_url)
        runtime = AttributeRuntime(
            store,
            client,
            run_id,
            fingerprint.product_version,
            device_id=store.get_or_create_device_id(),
            verified_history=load_verified_attribute_history(
                SCRIPT_DIR / "output/kuaimai/runs", product.style_code
            ),
        )
    return LearningRunContext(
        store,
        run_id,
        fingerprint.product_version,
        store.get_or_create_device_id(),
        fingerprint.image_version,
        client,
        runtime,
    )


async def initialize_learning_run(
    context: Optional[LearningRunContext],
    product: ProductData,
    logger: logging.Logger,
) -> None:
    """Upload de-duplicated evidence and populate the one-shot vision cache."""
    if (
        context is None
        or getattr(context, "client", None) is None
        or getattr(context, "initialized", False)
    ):
        return
    verified_history = getattr(
        getattr(context, "attribute_runtime", None), "verified_history", {}
    )
    if verified_history:
        logger.info(
            "已加载同款历史保存回读：%s 个平台属性，可直接复用时不再重复审核",
            len(verified_history),
        )
    await flush_outbox_async(
        context.store,
        context.client,
        datetime.now(timezone.utc),
    )
    paths = tuple(
        dict.fromkeys(
            str(path)
            for path in (
                tuple(product.main_images)
                + tuple(product.main_images_34)
                + tuple(product.detail_images)
            )
        )
    )
    prepared: List[PreparedLearningAsset] = []
    for raw_path in paths:
        original, thumbnail = await asyncio.to_thread(
            prepare_learning_assets, Path(raw_path)
        )
        prepared.append(thumbnail)
        if original.content_type in {"image/jpeg", "image/png", "image/webp"}:
            prepared.append(original)
    unique_assets = {
        (asset.kind, asset.sha256): asset for asset in prepared
    }
    for asset in unique_assets.values():
        await asyncio.to_thread(
            context.client.upload_asset,
            context.product_version,
            asset.sha256,
            asset.kind,
            asset.content_type,
            asset.body,
        )
    analysis = await asyncio.to_thread(
        context.client.analyze, context.product_version
    )
    context.initialized = True
    logger.info(
        "AI 学习素材已同步：%s 个文件、%s 个去重对象；识图状态=%s%s",
        len(paths),
        len(unique_assets),
        analysis.get("status", "unknown"),
        "（命中缓存）" if analysis.get("cached") is True else "",
    )


def natural_key(path: Path) -> List[Any]:
    # 文件名中的连字符和空格经常只是不同软件导出的分组分隔符。
    # 例如黑朗姆详情图同时存在 ``未标题-1-恢复的_01`` 和
    # ``未标题 2_01``；若直接比较原始前缀，空格会排在连字符前，
    # 导致第 2 组跑到第 1 组前面。统一这两类分隔符后再做自然排序，
    # 可按组号 1、2 以及组内编号 01、02 排列，同时保持普通文件名的
    # ``1、2、10`` 数字排序行为。
    # 这些分隔符不参与分组编号本身，直接忽略它们，避免
    # ``未标题-1-恢复的`` 的连字符/后缀把组 1 排到 ``未标题2`` 后面。
    name = re.sub(r"[-\s]+", "", path.name)
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", name)]


def normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_category_hints(value: Any) -> Tuple[str, ...]:
    return tuple(
        part.strip()
        for part in re.split(r"[/／]", normalize_cell(value))
        if part.strip()
    )


def parse_spec_values(value: Any) -> Tuple[str, ...]:
    """Read an ordered Excel specification list without silently deduplicating it."""
    values = tuple(
        part.strip()
        for part in re.split(r"[/／,，、;；]", normalize_cell(value))
        if part.strip()
    )
    normalized = tuple(re.sub(r"\s+", "", item).casefold() for item in values)
    if len(set(normalized)) != len(normalized):
        raise AutomationError("Excel 颜色规格存在重复值，无法安全匹配 SKU 图片顺序")
    return values


def normalize_price(value: Any) -> str:
    raw_text = normalize_cell(value).replace(",", "")
    # 兼容运营在 Excel 中写成“690元”；仍然拒绝“690元起”等含糊价格。
    text = raw_text[:-1].strip() if raw_text.endswith("元") else raw_text
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise AutomationError(f"基本售价不是有效数字：{raw_text!r}") from exc
    if number < 0:
        raise AutomationError(f"基本售价不能小于 0：{raw_text!r}")
    return format(number.normalize(), "f")


EXCEL_FIELD_COLUMN_HEADERS = frozenset(
    ("属性", "属性筛选", "字段", "字段名", "属性名")
)
EXCEL_VALUE_COLUMN_HEADERS = frozenset(("内容", "值", "字段值", "属性值"))


def _normalized_excel_header(value: Any) -> str:
    return re.sub(r"\s+", "", normalize_cell(value)).casefold()


def read_excel_fields(rows: Iterable[Sequence[Any]]) -> Dict[str, Any]:
    """按表头列名读取字段，并兼容旧版无表头的键值行。"""
    materialized_rows = [tuple(row) for row in rows]
    field_column: Optional[int] = None
    value_column: Optional[int] = None
    data_start = 0

    for row_index, row in enumerate(materialized_rows):
        normalized = [_normalized_excel_header(cell) for cell in row]
        field_matches = [
            index
            for index, value in enumerate(normalized)
            if value in EXCEL_FIELD_COLUMN_HEADERS
        ]
        value_matches = [
            index
            for index, value in enumerate(normalized)
            if value in EXCEL_VALUE_COLUMN_HEADERS
        ]
        if field_matches and value_matches and field_matches[0] != value_matches[0]:
            field_column = field_matches[0]
            value_column = value_matches[0]
            data_start = row_index + 1
            break

    fields: Dict[str, Any] = {}
    source_rows: Dict[str, int] = {}
    for row_index, row in enumerate(materialized_rows[data_start:], start=data_start + 1):
        if not row:
            continue
        if field_column is None:
            key = normalize_cell(row[0] if row else None)
            value = next(
                (
                    cell
                    for cell in row[1:]
                    if cell is not None and normalize_cell(cell)
                ),
                None,
            )
        else:
            key = normalize_cell(
                row[field_column] if field_column < len(row) else None
            )
            value = row[value_column] if value_column < len(row) else None
        if not key:
            continue

        normalized_key = re.sub(r"\s+", "", key).casefold()
        if normalized_key in source_rows:
            raise AutomationError(
                f"Excel 中存在重复字段 {key!r}："
                f"第 {source_rows[normalized_key]} 行和第 {row_index} 行"
            )
        fields[key] = value
        source_rows[normalized_key] = row_index
    return fields


def resolve_excel_path(value: str) -> Path:
    """将 smb://host/share/path 映射到 macOS 已挂载的 /Volumes/share/path。"""
    if value.lower().startswith("smb://"):
        parsed = urlsplit(value)
        parts = [unquote(part) for part in parsed.path.split("/") if part]
        if len(parts) < 2:
            raise AutomationError(f"SMB 路径缺少共享名或文件路径：{value}")
        path = Path("/Volumes").joinpath(*parts)
    else:
        path = Path(value).expanduser()

    path = path.resolve(strict=False)
    if not path.is_file():
        if value.lower().startswith("smb://"):
            raise AutomationError(
                f"找不到 Excel：{path}\n"
                f"请先在 Finder 中连接服务器 {urlsplit(value).hostname}，并挂载共享文件夹。"
            )
        raise AutomationError(f"找不到 Excel：{path}")
    return path


def find_image_dir(product_dir: Path, candidates: Sequence[str], label: str) -> Path:
    for relative in candidates:
        candidate = product_dir / relative
        if candidate.is_dir():
            return candidate
    names = "、".join(candidates)
    raise AutomationError(f"找不到{label}文件夹，已尝试：{names}")


def list_images(directory: Path, label: str) -> List[Path]:
    files = sorted(
        (path for path in directory.iterdir() if path.is_file() and path.suffix.casefold() in IMAGE_SUFFIXES),
        key=natural_key,
    )
    if not files:
        raise AutomationError(f"{label}文件夹中没有可上传图片：{directory}")
    return files


def read_product_data(
    excel_path: Path,
    *,
    include_douyin: bool = True,
) -> ProductData:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise AutomationError("缺少 openpyxl，请先运行：python3 -m pip install -r requirements.txt") from exc

    workbook = load_workbook(excel_path, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        raise AutomationError(f"Excel 内容不完整：{excel_path}")

    fields = read_excel_fields(rows)

    def field_value(*aliases: str) -> Any:
        found = field_lookup(fields, *aliases)
        return found[1] if found else None

    title = normalize_cell(field_value("商品标题", "商品名称", "宝贝标题"))
    style_code = normalize_cell(field_value("货号", "商家外部编码", "款式编码"))
    price_value = field_value("基本售价", "吊牌价", "商品价格", "价格")

    if not title:
        raise AutomationError("Excel 中找不到商品标题/商品名称/宝贝标题")
    if not style_code:
        raise AutomationError("Excel 中找不到货号/款式编码")
    if price_value is None or normalize_cell(price_value) == "":
        raise AutomationError("Excel 中找不到“吊牌价/价格/基本售价”")

    product_dir = excel_path.parent
    category_hints = parse_category_hints(field_value("商品分类"))
    colors = parse_spec_values(field_value("颜色"))
    profile = category_profile(category_hints, title)
    derived_size_names: Tuple[str, ...] = ()
    size_chart_dir = product_dir / "尺码信息表"
    if (
        field_value("尺码") is None
        and profile.supports_letter_size_chart
        and size_chart_dir.is_dir()
    ):
        size_chart_images = list_images(size_chart_dir, "尺码信息表")
        if len(size_chart_images) != 1:
            raise AutomationError(
                "Excel 缺少尺码时，尺码信息表文件夹必须恰好 1 张图片，"
                f"当前为 {len(size_chart_images)} 张"
            )
        derived_size_names = recognize_size_names(
            size_chart_images[0], profile.garment_kind
        )
        fields = dict(fields)
        fields["尺码"] = "/".join(derived_size_names)
    main_dir = find_image_dir(
        product_dir,
        ("1：1主图", "1:1主图", "主图/1：1", "主图/1:1"),
        "1:1 主图",
    )
    main_34_dir = find_image_dir(
        product_dir,
        ("3：4主图", "3:4主图", "主图/3：4", "主图/3:4"),
        "3:4 主图",
    )
    detail_dir = find_image_dir(product_dir, ("详情页图", "商品详情图"), "商品详情图")
    sku_dir = find_image_dir(product_dir, ("SKU图", "sku图", "Sku图"), "SKU 图")
    has_douyin_signals = field_lookup(fields, "导购短标题") is not None or any(
        (product_dir / directory_name).is_dir()
        for directory_name in ("水洗标图片", "尺码信息表", "身高体重推荐表")
    )

    return ProductData(
        excel_path=excel_path,
        product_dir=product_dir,
        title=title,
        style_code=style_code,
        base_price=normalize_price(price_value),
        main_images=list_images(main_dir, "1:1 主图"),
        main_images_34=list_images(main_34_dir, "3:4 主图"),
        detail_images=list_images(detail_dir, "商品详情图"),
        sku_images=list_images(sku_dir, "SKU 图"),
        douyin_fields=(
            parse_douyin_fields(fields)
            if include_douyin and has_douyin_signals
            else None
        ),
        douyin_assets=(
            read_douyin_assets(product_dir)
            if include_douyin and has_douyin_signals
            else None
        ),
        taobao_fields=parse_taobao_fields(fields),
        tmall_fields=parse_tmall_fields(fields),
        pdd_fields=parse_pdd_fields(fields),
        wxsph_fields=parse_wxsph_fields(fields),
        xhs_fields=parse_xhs_fields(fields),
        youzan_fields=parse_youzan_fields(fields),
        jd_fields=parse_jd_fields(fields),
        category_hints=category_hints,
        garment_kind=profile.garment_kind,
        derived_size_names=derived_size_names,
        colors=colors,
    )


def read_new_product_seed(excel_path: Path) -> ProductData:
    """只建基础链接：名称 1、一张主图、Excel 颜色、固定尺码、价格 0。"""
    from openpyxl import load_workbook

    workbook = load_workbook(excel_path, data_only=True, read_only=True)
    try:
        fields = read_excel_fields(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()
    style_field = field_lookup(fields, "货号", "款式编码", "商家外部编码")
    style_code = normalize_cell(style_field[1] if style_field else None)
    if not style_code:
        raise AutomationError("新增链接需要 Excel 中的货号/款式编码")
    color_field = field_lookup(fields, "颜色")
    colors = parse_spec_values(color_field[1] if color_field else None)
    if not colors:
        raise AutomationError("新增链接需要 Excel 中的颜色规格")
    directory = find_image_dir(
        excel_path.parent, ("1：1主图", "1:1主图", "主图/1：1", "主图/1:1", "主图"), "主图"
    )
    return ProductData(
        excel_path, excel_path.parent, "1", style_code, "0",
        list_images(directory, "主图")[:1], [], [], [], colors=colors,
    )


def product_summary(product: ProductData) -> Dict[str, Any]:
    def serialize(value: Any) -> Any:
        if isinstance(value, Path):
            return str(value)
        if isinstance(value, Mapping):
            return {key: serialize(item) for key, item in value.items()}
        if isinstance(value, (list, tuple)):
            return [serialize(item) for item in value]
        return value

    payload = asdict(product)
    taobao_fields = payload.pop("taobao_fields", None)
    tmall_fields = payload.pop("tmall_fields", None)
    pdd_fields = payload.pop("pdd_fields", None)
    wxsph_fields = payload.pop("wxsph_fields", None)
    xhs_fields = payload.pop("xhs_fields", None)
    youzan_fields = payload.pop("youzan_fields", None)
    jd_fields = payload.pop("jd_fields", None)
    result = serialize(payload)
    if taobao_fields is not None and product.taobao_fields is not None:
        result["taobao_field_count"] = len(product.taobao_fields.fields)
    if tmall_fields is not None and product.tmall_fields is not None:
        result["tmall_field_count"] = len(product.tmall_fields.fields)
    if pdd_fields is not None and product.pdd_fields is not None:
        result["pdd_field_count"] = len(product.pdd_fields.fields)
    if wxsph_fields is not None and product.wxsph_fields is not None:
        result["wxsph_field_count"] = len(product.wxsph_fields.fields)
    if xhs_fields is not None and product.xhs_fields is not None:
        result["xhs_field_count"] = len(product.xhs_fields.fields)
    if youzan_fields is not None and product.youzan_fields is not None:
        result["youzan_field_count"] = len(product.youzan_fields.fields)
    if jd_fields is not None and product.jd_fields is not None:
        result["jd_field_count"] = len(product.jd_fields.fields)
    return result


def recognize_product_recommendations(
    product: ProductData,
    artifact_dir: Path,
) -> tuple[Any, ...]:
    """在打开浏览器前完成本地尺码识别，失败则不进入页面。"""
    if product.douyin_fields is None or product.douyin_assets is None:
        return ()
    if product.garment_kind == "pants":
        recommendations = recognize_recommendations(
            product.douyin_assets.size_chart_image,
            product.douyin_assets.height_weight_image,
            product.douyin_fields.sizes,
        )
    elif product.garment_kind == "clothing":
        recommendations = recognize_clothing_recommendations(
            product.douyin_assets.size_chart_image,
            product.douyin_assets.height_weight_image,
            product.douyin_fields.sizes,
        )
    else:
        (artifact_dir / "ocr-result.json").write_text(
            json.dumps(
                {
                    "status": "skipped",
                    "reason": "unsupported_garment_kind",
                    "garment_kind": product.garment_kind,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        return ()
    payload = [asdict(item) for item in recommendations]
    (artifact_dir / "ocr-result.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return recommendations


def build_tmall_size_rows(
    recommendations: Sequence[SkuRecommendation],
) -> Tuple[Mapping[str, Any], ...]:
    """Convert verified OCR output into Tmall's dynamic size-table sources.

    Height and weight stay as explicit ranges.  The form adapter decides which
    page-supported columns to enable; keeping the source lossless prevents a
    minimum or maximum value from being guessed away.
    """
    return tuple(
        {
            "尺码": item.size,
            "身高(cm)": (item.height_min, item.height_max),
            "体重(kg)": (item.weight_min, item.weight_max),
            "腰围(cm)": item.waist,
            "臀围(cm)": item.hip,
            "裤长(cm)": item.length,
        }
        for item in recommendations
    )


def tmall_preview_summary(report: Mapping[str, Any]) -> Dict[str, Any]:
    """Return a disk-safe Tmall report without product values or platform IDs."""

    def child_mapping(name: str) -> Mapping[str, Any]:
        value = report.get(name)
        return value if isinstance(value, Mapping) else {}

    def keys_of(value: Any) -> Tuple[str, ...]:
        return tuple(str(key) for key in value) if isinstance(value, Mapping) else ()

    identity = child_mapping("product_identity")
    attributes = child_mapping("attributes")
    specifications = child_mapping("specifications")
    sku_batch = child_mapping("sku_batch")
    size_chart = child_mapping("size_chart")
    sales = child_mapping("sales_and_logistics")
    images = child_mapping("images")
    after_sales = child_mapping("after_sales")
    api_json_validation = child_mapping("api_json_validation")

    summary: Dict[str, Any] = {
        "platform": "tmall",
        "status": str(report.get("status") or "unknown"),
        "category": str(report.get("category") or ""),
        "saved": bool(report.get("saved", False)),
        "published": bool(report.get("published", False)),
        "product_identity_fields": keys_of(identity.get("values", identity)),
    }
    if report.get("reason_code"):
        summary["reason_code"] = str(report["reason_code"])
    summary["form_mode"] = str(report.get("form_mode") or "unknown")
    for report_key in (
        "initial_product_images_ready",
        "existing_product_images_before",
        "existing_product_images_ready",
        "matched_product_images_ready",
    ):
        if not isinstance(report.get(report_key), Mapping):
            continue
        image_ready = report[report_key]
        summary[report_key] = {
            key: image_ready[key]
            for key in (
                "count",
                "expected_count",
                "slot_count",
                "decoded",
                "stable_seconds",
                "waited_seconds",
            )
            if key in image_ready
        }
    if attributes:
        summary["attribute_fields"] = keys_of(
            attributes.get("attributes", attributes)
        )
    if specifications:
        summary["specifications"] = {
            "changed": int(specifications.get("changed") or 0),
            "dimension_count": len(keys_of(specifications.get("dimensions", {}))),
        }
    if sku_batch:
        summary["sku_batch"] = {
            "batch_clicked": bool(sku_batch.get("batch_clicked", False)),
            "row_count": int(sku_batch.get("row_count") or 0),
            "fields": keys_of(sku_batch.get("values", {})),
            "platform_codes_preserved": bool(
                sku_batch.get("platform_codes_preserved", False)
            ),
        }
    if size_chart:
        summary["size_chart"] = {
            "row_count": int(size_chart.get("row_count") or 0),
            "checked_parameters": tuple(
                str(value) for value in size_chart.get("checked_parameters", ())
            ),
        }
    if sales:
        summary["sales_and_logistics_fields"] = keys_of(
            sales.get("values", sales)
        )
    if images:
        summary["images"] = {
            str(label): str(action) for label, action in images.items()
        }
    if after_sales:
        summary["after_sales_fields"] = keys_of(
            after_sales.get("values", after_sales)
        )
        summary["new_product_declaration_applied"] = (
            after_sales.get("new_product_declaration") == "是"
        )
    if api_json_validation:
        endpoint_summaries = []
        endpoints = api_json_validation.get("endpoints")
        if isinstance(endpoints, (tuple, list)):
            for endpoint in endpoints:
                if not isinstance(endpoint, Mapping):
                    continue
                path = urlsplit(str(endpoint.get("path") or "")).path
                if path not in TMALL_READ_ONLY_ENDPOINTS:
                    continue
                endpoint_summaries.append(
                    {
                        "path": path,
                        "method": str(endpoint.get("method") or ""),
                        "http_status": endpoint.get("http_status"),
                        "json": bool(endpoint.get("json", False)),
                        "field_count": int(endpoint.get("field_count") or 0),
                    }
                )
        summary["api_json_validation"] = {
            "status": str(api_json_validation.get("status") or "not_observed"),
            "captured_endpoint_count": int(
                api_json_validation.get("captured_endpoint_count") or 0
            ),
            "json_endpoint_count": int(
                api_json_validation.get("json_endpoint_count") or 0
            ),
            "api_field_count": int(api_json_validation.get("api_field_count") or 0),
            "dom_field_count": int(api_json_validation.get("dom_field_count") or 0),
            "matched_field_count": int(
                api_json_validation.get("matched_field_count") or 0
            ),
            "endpoints": tuple(endpoint_summaries),
        }
    return summary


async def write_tmall_review_required(
    page: Any,
    artifact_dir: Path,
    report: Dict[str, Any],
    *,
    reason_code: str,
) -> None:
    """写入不含商品值的天猫人工复核报告。"""

    report.update(
        {
            "status": "review_required",
            "reason_code": reason_code,
            "saved": False,
            "published": False,
        }
    )
    (artifact_dir / "tmall-review-required.json").write_text(
        json.dumps(
            tmall_preview_summary(report),
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    await safe_screenshot(page, artifact_dir / "tmall-review-required.png")


def resolve_taobao_category_mode(product: ProductData, category_mode: str) -> str:
    if category_mode != "auto":
        return category_mode
    return "casual-pants" if product.garment_kind == "pants" else "recommended"


def taobao_garment_kind(product: ProductData, category_mode: str) -> str:
    if category_mode == "casual-pants":
        return "pants"
    if product.garment_kind in {"pants", "clothing"}:
        return product.garment_kind
    raise AutomationError(
        "当前品类无法确定淘宝尺码表是衣长还是裤长，请进入运营审核"
    )


def recognize_taobao_size_lengths(
    product: ProductData,
    category_mode: str,
    artifact_dir: Path,
) -> tuple[str, tuple[SizeLength, ...]]:
    if product.taobao_fields is None:
        raise AutomationError("淘宝尺码识别缺少 Excel 字段")
    size_source = field_lookup(dict(product.taobao_fields.fields), "尺码")
    if size_source is None:
        raise AutomationError("Excel 中缺少淘宝尺码字段")
    sizes = tuple(
        value.strip()
        for value in re.split(r"[/／,，、;；]", str(size_source[1]))
        if value.strip()
    )
    if not sizes:
        raise AutomationError("Excel 中的淘宝尺码为空")
    size_chart_dir = product.product_dir / "尺码信息表"
    if not size_chart_dir.is_dir():
        raise AutomationError(f"找不到淘宝尺码信息表文件夹：{size_chart_dir}")
    images = list_images(size_chart_dir, "尺码信息表")
    if len(images) != 1:
        raise AutomationError(
            f"尺码信息表文件夹必须恰好 1 张图片，当前为 {len(images)} 张"
        )
    garment_kind = taobao_garment_kind(product, category_mode)
    lengths = recognize_size_lengths(images[0], sizes, garment_kind)
    (artifact_dir / "taobao-size-lengths.json").write_text(
        json.dumps(
            {
                "garment_kind": garment_kind,
                "source": str(images[0]),
                "rows": [asdict(item) for item in lengths],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return garment_kind, lengths


def setup_logging(
    artifact_dir: Path,
    redactor: Optional[SensitiveLogRedactor] = None,
) -> logging.Logger:
    artifact_dir.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("kuaimai_erp")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    if redactor is None:
        formatter = logging.Formatter(
            "%(asctime)s | %(levelname)s | %(message)s",
            "%Y-%m-%d %H:%M:%S",
        )
    else:
        formatter = RedactingFormatter(
            redactor,
            "%(asctime)s | %(levelname)s | %(message)s",
            "%Y-%m-%d %H:%M:%S",
        )
    stream = logging.StreamHandler(sys.stdout)
    stream.setFormatter(formatter)
    file_handler = logging.FileHandler(artifact_dir / "run.log", encoding="utf-8")
    file_handler.setFormatter(formatter)
    logger.addHandler(stream)
    logger.addHandler(file_handler)
    return logger


async def save_auth_state(
    context: Any,
    state_path: Path,
    logger: logging.Logger,
    scm_verified: bool = True,
) -> None:
    """显式保存包括会话 Cookie 在内的认证状态，避免 Chrome 关闭时清理。"""
    state = await context.storage_state()
    state["cookies"] = [
        cookie
        for cookie in state.get("cookies") or []
        if (cookie.get("domain") or "").lstrip(".").endswith("superboss.cc")
    ]
    state["origins"] = [
        origin_state
        for origin_state in state.get("origins") or []
        if (urlsplit(origin_state.get("origin") or "").hostname or "").endswith("superboss.cc")
    ]
    if not scm_verified:
        state["cookies"] = [
            cookie
            for cookie in state["cookies"]
            if not is_scm_cookie(cookie)
        ]
        state["origins"] = [
            origin_state
            for origin_state in state["origins"]
            if urlsplit(origin_state.get("origin") or "").hostname in ERP_HOSTS
        ]
    session_storage: Dict[str, List[Dict[str, str]]] = {}
    for page in getattr(context, "pages", []):
        try:
            snapshot = await page.evaluate(
                """
                () => ({
                  origin: location.origin,
                  entries: Object.entries(sessionStorage).map(([name, value]) => ({name, value}))
                })
                """
            )
            origin = snapshot.get("origin")
            hostname = urlsplit(origin or "").hostname or ""
            if hostname.endswith("superboss.cc") and (scm_verified or hostname in ERP_HOSTS):
                session_storage[origin] = snapshot.get("entries") or []
        except Exception:
            continue
    state["sessionStorage"] = session_storage
    state["scmVerified"] = scm_verified

    state_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = state_path.with_suffix(state_path.suffix + ".tmp")
    temporary_path.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")
    temporary_path.chmod(0o600)
    temporary_path.replace(state_path)
    state_path.chmod(0o600)
    logger.info(
        "已保存登录状态：%s 个 Cookie，%s 个站点，SCM 已验证=%s",
        len(state.get("cookies") or []),
        len(state.get("origins") or []),
        "是" if scm_verified else "否",
    )


async def restore_auth_state(context: Any, state_path: Path, logger: logging.Logger) -> bool:
    """恢复 storage_state 中不会由持久化 Chrome 自动保留的会话认证数据。"""
    if not state_path.is_file():
        logger.info("未找到已保存的登录状态，本次需要手工登录")
        return False
    try:
        state = json.loads(state_path.read_text(encoding="utf-8"))
        cookies = state.get("cookies") or []
        scm_verified = state.get("scmVerified") is True
        if not scm_verified:
            cookies = [
                cookie
                for cookie in cookies
                if not is_scm_cookie(cookie)
            ]
        if cookies:
            await context.add_cookies(cookies)

        storage_by_origin: Dict[str, Dict[str, List[Dict[str, str]]]] = {}
        for origin_state in state.get("origins") or []:
            origin = origin_state.get("origin")
            hostname = urlsplit(origin or "").hostname or ""
            if origin and (scm_verified or hostname in ERP_HOSTS):
                storage_by_origin.setdefault(origin, {})["localStorage"] = (
                    origin_state.get("localStorage") or []
                )
        for origin, entries in (state.get("sessionStorage") or {}).items():
            hostname = urlsplit(origin or "").hostname or ""
            if scm_verified or hostname in ERP_HOSTS:
                storage_by_origin.setdefault(origin, {})["sessionStorage"] = entries or []

        if storage_by_origin:
            serialized = json.dumps(storage_by_origin, ensure_ascii=False)
            await context.add_init_script(
                """
                (() => {
                  const states = %s;
                  const state = states[location.origin];
                  if (!state) return;
                  for (const item of state.localStorage || []) localStorage.setItem(item.name, item.value);
                  for (const item of state.sessionStorage || []) sessionStorage.setItem(item.name, item.value);
                })();
                """
                % serialized
            )
        logger.info(
            "已恢复登录状态：%s 个 Cookie，SCM 已验证=%s",
            len(cookies),
            "是" if scm_verified else "否",
        )
        return True
    except Exception as exc:
        logger.warning("登录状态文件无法恢复，将改为手工登录：%s", exc)
        return False


def saved_scm_state_is_verified(state_path: Optional[Path]) -> bool:
    """只有之前经商品 API 验证过的 SCM 状态才允许走快速路径。"""
    if state_path is None or not state_path.is_file():
        return False
    try:
        return json.loads(state_path.read_text(encoding="utf-8")).get("scmVerified") is True
    except Exception:
        return False


async def visible(locator: Any) -> bool:
    try:
        return await locator.is_visible()
    except Exception:
        return False


async def safe_screenshot(page: Any, path: Path) -> None:
    try:
        await page.screenshot(path=str(path), full_page=False)
    except Exception:
        pass


async def visible_text_across_frames(page: Any, text: str) -> Optional[Any]:
    """在主页面及所有同/跨域 iframe 中查找可见文本定位器。"""
    # 登录跳转会替换首页 iframe。frame 列表是一个瞬时快照，查询期间旧 frame
    # 可能已经脱离；忽略该快照并让外层状态轮询读取新 frame。
    for frame in list(page.frames):
        try:
            candidate = await first_visible(frame.get_by_text(text, exact=True))
            if candidate is None:
                # 部分账号的 ERP 首页会在文案后附加环境标识（例如“快麦通商品中心”）；
                # 仅放宽为以菜单名开头，避免用整个 body 的模糊文本误触发。
                candidate = await first_visible(
                    frame.get_by_text(
                        re.compile(
                            rf"^\s*{re.escape(text)}"
                            rf"(?:\s|[|｜>/（(]|商品中心|$)"
                        )
                    )
                )
            if candidate is not None:
                return candidate
        except Exception:
            continue
    return None


def _read_macos_keychain_password(account: str) -> str:
    """从 macOS 钥匙串读取快麦 ERP 密码，不输出密码或错误原文。"""
    if sys.platform != "darwin" or not account:
        return ""
    try:
        result = subprocess.run(
            [
                "/usr/bin/security",
                "find-generic-password",
                "-a",
                account,
                "-s",
                ERP_LOGIN_KEYCHAIN_SERVICE,
                "-w",
            ],
            check=False,
            capture_output=True,
            text=True,
            timeout=5,
        )
    except (OSError, subprocess.TimeoutExpired):
        return ""
    if result.returncode != 0:
        return ""
    return result.stdout.rstrip("\r\n")


def _delete_macos_keychain_password(account: str) -> bool:
    """Delete only this ERP account's saved credential."""
    if sys.platform != "darwin" or not account:
        return False
    try:
        result = subprocess.run(
            [
                "/usr/bin/security",
                "delete-generic-password",
                "-a",
                account,
                "-s",
                ERP_LOGIN_KEYCHAIN_SERVICE,
            ],
            check=False,
            capture_output=True,
            text=True,
            timeout=5,
        )
    except (OSError, subprocess.TimeoutExpired):
        return False
    return result.returncode == 0


def _prompt_and_store_macos_keychain_password(
    account: str,
    logger: logging.Logger,
) -> bool:
    """Prompt securely and keep the Keychain item only after confirmation."""
    if (
        sys.platform != "darwin"
        or not account
        or not sys.stdin.isatty()
        or account in _ERP_LOGIN_KEYCHAIN_PROMPTED_ACCOUNTS
    ):
        return False
    _ERP_LOGIN_KEYCHAIN_PROMPTED_ACCOUNTS.add(account)
    logger.warning(
        "配置快麦 ERP 自动登录：系统将提示输入密码，随后程序会要求再次输入确认；"
        "两次一致才保存到 macOS 钥匙串，密码不会写入项目文件或日志"
    )
    for attempt in range(1, 4):
        try:
            result = subprocess.run(
                [
                    "/usr/bin/security",
                    "add-generic-password",
                    "-U",
                    "-a",
                    account,
                    "-s",
                    ERP_LOGIN_KEYCHAIN_SERVICE,
                    "-l",
                    ERP_LOGIN_KEYCHAIN_LABEL,
                    "-w",
                ],
                check=False,
            )
        except OSError:
            return False
        if result.returncode != 0:
            logger.warning("未完成 macOS 钥匙串密码录入，本次保留手工登录等待")
            return False

        stored_password = _read_macos_keychain_password(account)
        if not stored_password:
            _delete_macos_keychain_password(account)
            logger.warning("无法回读刚保存的 ERP 凭据，已撤销本次保存")
            return False
        try:
            repeated_password = getpass.getpass("请再次输入快麦 ERP 登录密码以确认: ")
        except (EOFError, KeyboardInterrupt):
            repeated_password = ""
        matches = bool(repeated_password) and repeated_password == stored_password
        del repeated_password
        del stored_password
        if matches:
            logger.info("两次密码输入一致，macOS 钥匙串已保存快麦 ERP 自动登录凭据")
            return True

        _delete_macos_keychain_password(account)
        logger.warning("两次密码输入不一致，已撤销保存（第 %s/3 次）", attempt)

    _ERP_LOGIN_KEYCHAIN_PROMPTED_ACCOUNTS.discard(account)
    logger.warning("密码连续三次确认不一致，本次保留手工登录等待")
    return False


def resolve_erp_login_password(
    account: str,
    logger: logging.Logger,
) -> Tuple[str, str]:
    """按临时环境变量、macOS 钥匙串的顺序获取登录密码。"""
    environment_password = os.environ.get("KUAIMAI_ERP_PASSWORD", "")
    if environment_password:
        return environment_password, "environment"

    account = account.strip()
    if not account:
        return "", ""
    now = time.monotonic()
    if now < _ERP_LOGIN_CREDENTIAL_RETRY_AT.get(account, 0.0):
        return "", ""

    password = _read_macos_keychain_password(account)
    if password:
        _ERP_LOGIN_CREDENTIAL_RETRY_AT.pop(account, None)
        return password, "keychain"

    if _prompt_and_store_macos_keychain_password(account, logger):
        password = _read_macos_keychain_password(account)
        if password:
            _ERP_LOGIN_CREDENTIAL_RETRY_AT.pop(account, None)
            return password, "keychain"

    _ERP_LOGIN_CREDENTIAL_RETRY_AT[account] = (
        time.monotonic() + ERP_LOGIN_CREDENTIAL_RETRY_SECONDS
    )
    return "", ""


async def _erp_login_account(frame: Any) -> str:
    """从当前登录页回读账号，避免将公司名误当账号。"""
    for selector in (
        'input[placeholder*="账号"]:visible',
        'input[autocomplete="username"]:visible',
    ):
        candidate = await first_visible(frame.locator(selector))
        if candidate is not None:
            value = (await candidate.input_value()).strip()
            if value:
                return value
    return os.environ.get("KUAIMAI_ERP_ACCOUNT", "").strip()


async def invalidate_rejected_erp_keychain_credential(
    page: Any,
    logger: logging.Logger,
) -> bool:
    """Remove a Keychain password only after the ERP page explicitly rejects it."""
    if not _ERP_LOGIN_PENDING_KEYCHAIN_ACCOUNTS:
        return False

    rejection_seen = False
    rejected_account = ""
    selectors = (
        ".el-message--error:visible, .el-form-item__error:visible, "
        "[role='alert']:visible, .login-error:visible, .error-msg:visible"
    )
    for frame in list(page.frames):
        try:
            messages = frame.locator(selectors)
            for index in range(await messages.count()):
                text = re.sub(r"\s+", "", (await messages.nth(index).inner_text()).strip())
                if ERP_LOGIN_REJECTION_PATTERN.search(text):
                    rejection_seen = True
                    account = await _erp_login_account(frame)
                    if account in _ERP_LOGIN_PENDING_KEYCHAIN_ACCOUNTS:
                        rejected_account = account
                    break
        except Exception:
            continue
        if rejection_seen:
            break

    if not rejection_seen:
        return False
    if not rejected_account and len(_ERP_LOGIN_PENDING_KEYCHAIN_ACCOUNTS) == 1:
        rejected_account = next(iter(_ERP_LOGIN_PENDING_KEYCHAIN_ACCOUNTS))
    if not rejected_account:
        logger.warning("ERP 页面拒绝了自动登录凭据，但无法唯一确定对应账号；已停止自动重试")
        _ERP_LOGIN_PENDING_KEYCHAIN_ACCOUNTS.clear()
        return True

    deleted = _delete_macos_keychain_password(rejected_account)
    _ERP_LOGIN_PENDING_KEYCHAIN_ACCOUNTS.discard(rejected_account)
    _ERP_LOGIN_KEYCHAIN_PROMPTED_ACCOUNTS.discard(rejected_account)
    _ERP_LOGIN_CREDENTIAL_RETRY_AT.pop(rejected_account, None)
    _ERP_LOGIN_MISSING_CREDENTIAL_WARNED.discard(rejected_account)
    if deleted:
        logger.warning(
            "快麦 ERP 明确返回账号或密码错误；已删除该账号的错误钥匙串凭据，"
            "请按终端提示重新输入并确认"
        )
    else:
        _ERP_LOGIN_CREDENTIAL_RETRY_AT[rejected_account] = float("inf")
        logger.warning(
            "快麦 ERP 明确返回账号或密码错误，但钥匙串凭据删除失败；"
            "已停止使用该凭据，请手工登录后重试"
        )
    return True


async def try_erp_login_with_agreement(page: Any, logger: logging.Logger) -> bool:
    """在已跳转的快麦 ERP 登录页补全密码、勾选协议并提交。

    密码可由本机环境变量或 macOS 钥匙串提供，仅填入 HTTPS ERP
    登录页，不写入项目文件或日志。没有密码时不提交空表单。
    """
    agreement_pattern = re.compile(r"我已阅读.*同意")
    for frame in list(page.frames):
        try:
            agreement = await first_visible(frame.get_by_text(agreement_pattern))
            if agreement is None:
                continue

            password_source = ""
            credential_account = ""
            passwords = frame.locator('input[type="password"]:visible')
            if await passwords.count():
                if await passwords.count() != 1:
                    continue
                password_input = passwords.first
                if not await password_input.input_value():
                    parsed = urlsplit(frame.url)
                    if parsed.scheme != "https" or parsed.hostname not in ERP_HOSTS:
                        continue
                    account = await _erp_login_account(frame)
                    password, password_source = resolve_erp_login_password(account, logger)
                    if not password:
                        warning_key = account or "<unknown-account>"
                        if warning_key not in _ERP_LOGIN_MISSING_CREDENTIAL_WARNED:
                            logger.warning(
                                "已识别快麦 ERP 登录页，但未获取到自动登录密码；"
                                "将等待钥匙串录入或手工登录，不会提交空密码"
                            )
                            _ERP_LOGIN_MISSING_CREDENTIAL_WARNED.add(warning_key)
                        continue
                    await password_input.fill(password)
                    del password
                    credential_account = account
                    _ERP_LOGIN_MISSING_CREDENTIAL_WARNED.discard(account)
                    logger.info(
                        "已从%s读取快麦 ERP 自动登录凭据",
                        "macOS 钥匙串" if password_source == "keychain" else "本次运行环境",
                    )

            checkbox = None
            # 快麦当前登录页的 #reading 与 label[for=reading]
            # 是兄弟节点，且页面还有一个隐藏的二维码绑定复选框。
            # 优先用稳定 ID，避免误依赖“input 嵌套在 label 里”
            # 或“全页只有一个 checkbox”这两个不成立的假设。
            current_checkbox = frame.locator("#reading")
            if await current_checkbox.count() == 1:
                checkbox = current_checkbox.first
            else:
                label = agreement.locator("xpath=ancestor::label[1]")
                if await label.count() == 1:
                    candidate = label.locator('input[type="checkbox"]').first
                    if await candidate.count() == 1:
                        checkbox = candidate
            if checkbox is None:
                checkboxes = frame.locator('input[type="checkbox"]')
                if await checkboxes.count() == 1:
                    checkbox = checkboxes.first
            if checkbox is None:
                continue

            if not await checkbox.is_checked():
                try:
                    await checkbox.check(force=True)
                except Exception:
                    # Element UI 的真实 input 可能透明，由包含文案的 label
                    # 触发同一控件，仍须回读 checked 状态后才继续提交。
                    await agreement.click(force=True)
                if not await checkbox.is_checked():
                    continue

            login = await first_visible(frame.locator("#login-btn"))
            if login is None:
                login = await first_visible(
                    frame.get_by_role("button", name=re.compile(r"^\s*登\s*录\s*$"))
                )
            if login is None:
                continue
            await login.click()
            if password_source == "keychain" and credential_account:
                _ERP_LOGIN_PENDING_KEYCHAIN_ACCOUNTS.add(credential_account)
            logger.info("检测到快麦 ERP 登录页：已勾选用户协议并点击登录")
            return True
        except Exception as exc:
            # 页面登录框常在跳转期间销毁；下一轮登录检测会读取新的 frame。
            logger.debug("快麦 ERP 自动登录本轮未完成：%s", exc)
            continue
    return False


async def dismiss_erp_blocking_dialogs(page: Any, logger: logging.Logger) -> int:
    """关闭只用于提醒、但会遮挡快麦通入口的 ERP 首页弹窗。"""
    dismissed = 0
    for frame in list(page.frames):
        try:
            dialogs = frame.locator(".el-dialog:visible").filter(has_text="店铺状态异常确认")
            for index in range(await dialogs.count()):
                dialog = dialogs.nth(index)
                close_button = dialog.locator(
                    "button.el-dialog__headerbtn, .el-dialog__headerbtn, .el-dialog__close"
                ).first
                if await visible(close_button):
                    await close_button.click(force=True)
                    dismissed += 1
        except Exception:
            continue
    if dismissed:
        logger.info("已关闭 %s 个 ERP 首页遮挡弹窗（未进入店铺管理）", dismissed)
    return dismissed


async def dismiss_scm_first_entry_dialog(page: Any, logger: logging.Logger) -> int:
    """首次进入快麦通时，只关闭温馨提示，不进入“立即查看”流程。"""
    dismissed = 0
    for frame in list(page.frames):
        try:
            dialogs = frame.locator(".el-dialog:visible").filter(has_text="温馨提示")
            for index in range(await dialogs.count()):
                dialog = dialogs.nth(index)
                acknowledge = await first_visible(
                    dialog.get_by_role("button", name="知道了", exact=True)
                )
                if acknowledge is None:
                    acknowledge = await first_visible(dialog.get_by_text("知道了", exact=True))
                if acknowledge is not None:
                    await acknowledge.click(force=True)
                    try:
                        await dialog.wait_for(state="hidden", timeout=1500)
                    except Exception:
                        # 弹窗可能在点击后直接销毁，下一轮仍会再次确认页面上已无可见项。
                        pass
                    dismissed += 1
        except Exception:
            continue
    if dismissed:
        logger.info("已关闭 %s 个快麦通首次进入温馨提示", dismissed)
    return dismissed


async def settle_scm_first_entry_dialog(
    page: Any,
    logger: logging.Logger,
    timeout_seconds: float,
) -> int:
    """给异步渲染的首次进入提示留出时间，并在交接页面前确认已关闭。"""
    deadline = time.monotonic() + timeout_seconds
    dismissed = 0
    quiet_since: Optional[float] = None
    while time.monotonic() < deadline:
        current = await dismiss_scm_first_entry_dialog(page, logger)
        dismissed += current
        if current:
            quiet_since = time.monotonic()
        elif dismissed and quiet_since is not None and time.monotonic() - quiet_since >= 0.25:
            return dismissed
        await asyncio.sleep(0.1)
    return dismissed


async def wait_for_erp_entry_stable(page: Any, timeout_seconds: int) -> Any:
    """等待快麦通入口连续可见，避免 ERP 刚跳首页就过早触发 SSO。"""
    deadline = time.monotonic() + timeout_seconds
    stable_since: Optional[float] = None
    while time.monotonic() < deadline:
        entry = await visible_text_across_frames(page, "快麦通")
        if entry is None:
            stable_since = None
        else:
            now = time.monotonic()
            if stable_since is None:
                stable_since = now
            elif now - stable_since >= 0.5:
                return entry
        await asyncio.sleep(0.1)
    raise AutomationError(f"快麦 ERP 首页的“快麦通”入口在 {timeout_seconds} 秒内未稳定加载")


async def wait_for_scm_api_session(
    page: Any,
    style_code: str,
    timeout_seconds: int,
    logger: logging.Logger,
) -> None:
    """以受保护的商品查询 API 为准，确认 SCM 会话确实已经建立。"""
    body = {
        "pageNo": 1,
        "pageSize": 1,
        "outerIds": [style_code] if style_code else [],
        "companyNames": [],
        "shopNames": [],
        "api_name": "item_base_page",
    }
    deadline = time.monotonic() + timeout_seconds
    last_status = "尚未请求"
    next_log_at = 0.0
    while time.monotonic() < deadline:
        await dismiss_scm_first_entry_dialog(page, logger)
        current_path = urlsplit(page.url).path
        if current_path.startswith("/login"):
            raise AutomationError("ERP 单点登录未建立快麦通会话，页面被跳转到快麦通登录页")
        try:
            result = await page.evaluate(
                """
                async (body) => {
                  try {
                    const response = await fetch('/item/base/page.json', {
                      method: 'POST',
                      credentials: 'include',
                      headers: {'Content-Type': 'application/json'},
                      body: JSON.stringify(body)
                    });
                    const contentType = response.headers.get('content-type') || '';
                    let data = null;
                    if (contentType.includes('json')) {
                      try { data = await response.json(); } catch (_) {}
                    }
                    return {
                      status: response.status,
                      finalPath: new URL(response.url).pathname,
                      result: data && data.result
                    };
                  } catch (error) {
                    return {error: String(error)};
                  }
                }
                """,
                body,
            )
            last_status = (
                f"HTTP {result.get('status')}，result={result.get('result')}，"
                f"响应路径={result.get('finalPath') or '-'}"
            )
            if int(result.get("result", 0) or 0) == 1:
                logger.info("快麦通会话已由商品查询 API 确认")
                return
        except Exception as exc:
            last_status = str(exc)

        now = time.monotonic()
        if now >= next_log_at:
            logger.info("等待快麦通会话就绪：%s", last_status)
            next_log_at = now + 10
        await asyncio.sleep(0.5)
    raise AutomationError(
        f"快麦通页面已打开，但会话在 {timeout_seconds} 秒内未通过 API 校验；最后状态：{last_status}"
    )


async def try_reuse_verified_scm_session(
    page: Any,
    style_code: str,
    timeout_seconds: int,
    logger: logging.Logger,
    auth_state_path: Optional[Path] = None,
) -> Any:
    """
    复用已验证的快麦通会话；任何异常由调用方回退到 ERP SSO。
    快速路径仍必须通过受保护商品 API，不以 Cookie 存在与否作为成功依据。
    """
    fast_timeout = max(3, min(timeout_seconds, 12))
    fast_timeout_ms = fast_timeout * 1000
    center_url = center_url_for(page.url)
    await page.goto(
        center_navigation_url_for(page.url),
        wait_until="domcontentloaded",
        timeout=fast_timeout_ms,
    )
    await wait_for_scm_api_session(page, style_code, fast_timeout, logger)

    deadline = time.monotonic() + fast_timeout
    while time.monotonic() < deadline:
        if urlsplit(page.url).path.startswith("/login"):
            raise AutomationError("已保存的快麦通会话已失效")
        if await visible(page.get_by_text("商品中心", exact=True).first):
            # 重新登录后商品中心可能在 API 校验成功之后才渲染“温馨提示”；
            # 在把页面交给商品查询/编辑流程前再关闭一次，只点“知道了”。
            await settle_scm_first_entry_dialog(page, logger, timeout_seconds=2.0)
            if auth_state_path is not None:
                await save_auth_state(page.context, auth_state_path, logger)
            logger.info("已复用通过 API 校验的快麦通会话，跳过 ERP 中转")
            return page
        await asyncio.sleep(0.1)
    raise AutomationError("快速打开商品中心超时")


async def open_product_center_from_supplier(
    page: Any,
    timeout_seconds: int,
    logger: logging.Logger,
) -> None:
    """优先沿页面菜单进入商品中心，避免直接跳转触发重新登录。"""
    timeout_ms = timeout_seconds * 1000
    center_url = center_url_for(page.url)
    if urlsplit(page.url).path == "/supplier/prod/center":
        return

    deadline = time.monotonic() + min(timeout_seconds, 3)
    while time.monotonic() < deadline:
        direct_link = page.locator('a[href*="/supplier/prod/center"]').first
        if await visible(direct_link):
            await direct_link.click()
            break

        management = page.get_by_text("商品管理", exact=True).first
        if await visible(management):
            try:
                await management.click(timeout=2000)
            except Exception:
                pass

        center_entry = page.get_by_text("商品中心", exact=True).first
        if await visible(center_entry):
            await center_entry.click()
            break
        await asyncio.sleep(0.2)
    else:
        logger.info("供应商首页未找到商品中心菜单，使用带 Cookie 标记的页面地址")
        await page.goto(
            center_navigation_url_for(page.url),
            wait_until="domcontentloaded",
            timeout=timeout_ms,
        )
        return

    navigation_deadline = time.monotonic() + min(timeout_seconds, 20)
    while time.monotonic() < navigation_deadline:
        if urlsplit(page.url).path == "/supplier/prod/center":
            return
        if urlsplit(page.url).path.startswith("/login"):
            raise AutomationError("从页面菜单进入商品中心时被跳转到登录页")
        await asyncio.sleep(0.2)

    logger.info("商品中心菜单未完成跳转，使用带 Cookie 标记的页面地址")
    await page.goto(
        f"{center_url}?hasCookie=true",
        wait_until="domcontentloaded",
        timeout=timeout_ms,
    )


async def enter_kuaimai_from_erp(
    page: Any,
    timeout_seconds: int,
    headless: bool,
    logger: logging.Logger,
    operation_timeout_seconds: Optional[int] = None,
    auth_state_path: Optional[Path] = None,
    style_code: str = "",
) -> Any:
    """从快麦 ERP 首页的“快麦通”入口进入 SCM，避免直接访问 SCM 登录页。"""
    operation_timeout = operation_timeout_seconds or min(timeout_seconds, 300)
    operation_timeout_ms = operation_timeout * 1000
    await page.goto(ERP_ENTRY_URL, wait_until="domcontentloaded", timeout=operation_timeout_ms)

    async def erp_home_ready() -> bool:
        parsed = urlsplit(page.url)
        if parsed.hostname not in ERP_HOSTS:
            return False
        if await visible_text_across_frames(page, "快麦通") is not None:
            return True

        # 登录成功后首页的菜单文案可能还在异步渲染。只要已离开登录路径且登录表单不再可见，就先认定 ERP 登录阶段已完成；
        # 后续的“快麦通”入口稳定等待会再负责等待菜单出现。
        if parsed.path.casefold().startswith("/login"):
            return False
        for frame in list(page.frames):
            try:
                password_inputs = frame.locator('input[type="password"]:visible')
                login_buttons = frame.locator(
                    "#login-btn:visible, button:visible"
                ).filter(has_text=re.compile(r"^\s*登\s*录\s*$"))
                if await password_inputs.count() or await login_buttons.count():
                    return False
            except Exception:
                continue
        return parsed.path.casefold() in {
            "",
            "/",
            "/index.html",
            "/simple.html",
            "/index",
        }

    if not await erp_home_ready():
        login_clicked = await try_erp_login_with_agreement(page, logger)
        if login_clicked:
            # 登录请求会异步替换整页；让新首页或登录错误提示先完成一次渲染。
            await asyncio.sleep(0.5)
        if await erp_home_ready():
            login_clicked = True
        elif headless:
            raise AutomationError("当前自动化专用 Chrome 未登录快麦 ERP，请去掉 --headless 后先登录一次")
        else:
            logger.info(
                "等待登录：请在打开的 Chrome 中完成快麦 ERP 登录（最多等待 %s 秒）",
                timeout_seconds,
            )
            deadline = time.monotonic() + timeout_seconds
            while time.monotonic() < deadline:
                if await erp_home_ready():
                    _ERP_LOGIN_PENDING_KEYCHAIN_ACCOUNTS.clear()
                    break
                if await invalidate_rejected_erp_keychain_credential(page, logger):
                    login_clicked = False
                if not login_clicked:
                    login_clicked = await try_erp_login_with_agreement(page, logger)
                await asyncio.sleep(1)
            else:
                raise AutomationError("等待快麦 ERP 登录超时")

    await dismiss_erp_blocking_dialogs(page, logger)
    context = page.context
    if auth_state_path is not None:
        await save_auth_state(context, auth_state_path, logger, scm_verified=False)

    logger.info("已进入快麦 ERP 首页，等待账号与“快麦通”入口稳定")
    entry = await wait_for_erp_entry_stable(page, operation_timeout)
    logger.info("快麦 ERP 首页已稳定，正在从“快麦通”入口进入商品中心")

    existing_scm_pages = {id(candidate) for candidate in context.pages if is_scm_url(candidate.url)}
    # ERP 弹窗关闭后偶尔残留透明 v-modal。原生 click 通常立即返回；若站点
    # 事件处理仍阻塞，最多等待 5 秒后也直接扫描/接管已经打开的 SCM 标签。
    try:
        await asyncio.wait_for(entry.evaluate("element => element.click()"), timeout=5)
        logger.info("已触发 ERP“快麦通”入口，开始接管快麦通标签")
    except asyncio.TimeoutError:
        logger.warning("ERP 入口触发 5 秒未返回，继续接管已打开的快麦通标签")
    except Exception as exc:
        logger.warning("ERP 入口触发返回异常，继续检查已打开标签：%s", exc)

    scm_page = None
    open_deadline = time.monotonic() + operation_timeout
    while time.monotonic() < open_deadline:
        if is_scm_url(page.url):
            scm_page = page
            break
        candidates = [candidate for candidate in context.pages if is_scm_url(candidate.url)]
        new_candidates = [candidate for candidate in candidates if id(candidate) not in existing_scm_pages]
        if new_candidates:
            scm_page = new_candidates[-1]
            break
        reusable_candidates = [
            candidate
            for candidate in candidates
            if urlsplit(candidate.url).path.startswith("/supplier/")
        ]
        if reusable_candidates:
            scm_page = reusable_candidates[-1]
            break
        await asyncio.sleep(0.5)
    if scm_page is None:
        raise AutomationError("点击“快麦通”后未打开快麦通页面")

    # ERP 会先打开 /account/erpVisit.json，再由该页建立 SCM 会话并跳到
    # /supplier/index。不能在中间页出现时就 goto 商品中心，否则会打断 SSO。
    last_path = ""
    sso_deadline = time.monotonic() + operation_timeout
    while time.monotonic() < sso_deadline:
        current_url = scm_page.url
        parsed = urlsplit(current_url)
        current_path = parsed.path
        if current_path != last_path:
            logger.info("快麦通 SSO 跳转：%s", current_path or "/")
            last_path = current_path
        if parsed.hostname in SCM_HOSTS and current_path.startswith("/supplier/"):
            break
        if parsed.hostname in SCM_HOSTS and current_path.startswith("/login"):
            raise AutomationError("ERP 单点登录未建立快麦通会话，页面被跳转到快麦通登录页")
        await asyncio.sleep(0.25)
    else:
        raise AutomationError(
            f"等待快麦通 SSO 完成超时（{operation_timeout} 秒），最后页面：{scm_page.url}"
        )

    await scm_page.wait_for_load_state("domcontentloaded", timeout=operation_timeout_ms)
    await wait_for_scm_api_session(scm_page, style_code, operation_timeout, logger)
    await open_product_center_from_supplier(scm_page, operation_timeout, logger)
    center_deadline = time.monotonic() + operation_timeout
    while time.monotonic() < center_deadline:
        current_path = urlsplit(scm_page.url).path
        if current_path.startswith("/login"):
            raise AutomationError("快麦通 API 曾通过校验，但打开商品中心时会话失效并跳回登录页")
        if await visible(scm_page.get_by_text("商品中心", exact=True).first):
            # 该提示是在商品中心异步初始化时出现，早于供应商首页的关闭时机；
            # 进入中心后再次处理，避免遮挡后续款式查询。
            await settle_scm_first_entry_dialog(scm_page, logger, timeout_seconds=3.0)
            break
        await asyncio.sleep(0.25)
    else:
        raise AutomationError(
            f"已从 ERP 打开快麦通，但商品中心在 {operation_timeout} 秒内未成功加载；"
            f"当前页面：{scm_page.url}"
        )
    if auth_state_path is not None:
        await save_auth_state(context, auth_state_path, logger)
    logger.info("已通过 ERP 单点登录进入快麦通商品中心")
    return scm_page


def find_records(payload: Any) -> Iterable[Dict[str, Any]]:
    if isinstance(payload, dict):
        records = payload.get("records")
        if isinstance(records, list):
            for record in records:
                if isinstance(record, dict):
                    yield record
        for value in payload.values():
            if isinstance(value, (dict, list)):
                yield from find_records(value)
    elif isinstance(payload, list):
        for value in payload:
            if isinstance(value, (dict, list)):
                yield from find_records(value)


async def api_find_product(
    page: Any,
    style_code: str,
    logger: logging.Logger,
    redactor: Optional[SensitiveLogRedactor] = None,
    *,
    strict: bool = False,
) -> Optional[Dict[str, Any]]:
    """在已登录页面内调用快麦同源查询 API，自动携带登录 Cookie。"""
    body = {
        "pageNo": 1,
        "pageSize": 20,
        "outerIds": [style_code],
        "companyNames": [],
        "shopNames": [],
        "api_name": "item_base_page",
    }
    try:
        payload = await page.evaluate(
            """
            async (body) => {
              const response = await fetch('/item/base/page.json', {
                method: 'POST',
                credentials: 'include',
                headers: {'Content-Type': 'application/json'},
                body: JSON.stringify(body)
              });
              const text = await response.text();
              let data = null;
              try { data = JSON.parse(text); } catch (_) {}
              return {ok: response.ok, status: response.status, data, text: text.slice(0, 500)};
            }
            """,
            body,
        )
    except Exception as exc:
        if strict:
            raise AutomationError("新增前查重失败，未确认商品不存在，禁止新增") from exc
        logger.warning("API 查询失败，将使用页面查询兜底：%s", exc)
        return None

    if not payload.get("ok") or not isinstance(payload.get("data"), dict):
        if strict:
            raise AutomationError("新增前查重响应异常，禁止新增")
        logger.warning("API 查询返回异常（HTTP %s），将使用 DOM 兜底", payload.get("status"))
        return None
    data = payload["data"]
    if int(data.get("result", 0) or 0) != 1:
        if strict:
            raise AutomationError("新增前查重未成功，请恢复登录后重试")
        logger.warning("API 查询未成功：%s", data.get("message") or data.get("errmsg") or data.get("result"))
        return None

    records: Iterable[Dict[str, Any]] = find_records(data.get("data"))
    if strict:
        query_result = data.get("data")
        if not isinstance(query_result, dict) or "records" not in query_result:
            raise AutomationError("新增前查重数据结构无法确认，禁止新增")
        try:
            total = int(query_result.get("total"))
        except (TypeError, ValueError) as exc:
            raise AutomationError("新增前查重缺少有效记录总数，禁止新增") from exc
        raw_records = query_result.get("records")
        if raw_records is None:
            # 快麦在精确查询没有结果时返回 records=null，而不是空数组。
            if total != 0:
                raise AutomationError("新增前查重记录总数与结果不一致，禁止新增")
            strict_records: List[Dict[str, Any]] = []
        elif isinstance(raw_records, list) and all(
            isinstance(row, dict) for row in raw_records
        ):
            strict_records = raw_records
            if total < len(strict_records) or (total == 0 and strict_records):
                raise AutomationError("新增前查重记录总数与结果不一致，禁止新增")
        else:
            raise AutomationError("新增前查重数据结构无法确认，禁止新增")
        if any(
            normalize_cell(row.get("outerId") or row.get("outerIds"))
            != style_code
            for row in strict_records
        ):
            raise AutomationError("新增前查询未精确匹配款式编码，禁止新增")
        records = strict_records

    matches = []
    for record in records:
        outer_id = normalize_cell(record.get("outerId") or record.get("outerIds"))
        if outer_id == style_code:
            matches.append(record)
    if len(matches) > 1:
        raise AutomationError(f"API 查到 {len(matches)} 个款式编码 {style_code} 的商品，无法安全确定编辑对象")
    if matches:
        record = matches[0]
        if redactor is not None:
            redactor.add_sensitive_values(
                record.get("baseItemId") or record.get("base_item_id") or record.get("id")
            )
        logger.info(
            "API 已定位目标商品：outerId=%s, baseItemId=%s",
            style_code,
            record.get("baseItemId") or record.get("id"),
        )
        return record
    if strict:
        logger.info("查重查询成功，款式编码 %s 尚不存在", style_code)
    else:
        logger.warning("API 未找到款式编码 %s，将使用页面查询复核", style_code)
    return None


async def first_visible(locator: Any) -> Optional[Any]:
    for index in range(await locator.count()):
        item = locator.nth(index)
        if await visible(item):
            return item
    return None


async def blocking_drawer_loading_mask_count(drawer: Any) -> int:
    """Count only loading masks that can block the editor as a whole.

    The live base form keeps a small loading mask inside every SKU numeric
    input (``.el-input-digit``) even after the form is usable.  Those local
    masks must not hold the entire editor readiness gate open.
    """
    masks = drawer.locator(".el-loading-mask:visible")
    return await masks.evaluate_all(
        "masks => masks.filter(mask => !mask.closest('.el-input-digit')).length"
    )


async def open_product_editor(
    page: Any,
    style_code: str,
    logger: logging.Logger,
    timeout_seconds: int = 300,
) -> Any:
    timeout_ms = timeout_seconds * 1000
    center_url = center_url_for(page.url)
    if page.url.split("?", 1)[0] != center_url:
        await page.goto(center_url, wait_until="domcontentloaded", timeout=timeout_ms)

    style_input = None
    search_deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < search_deadline:
        style_input = await first_visible(page.locator('input[placeholder*="多个款式编码"]'))
        if style_input is None:
            # 有些账号显示的文案不含“多个”。
            style_input = await first_visible(page.locator('input[placeholder*="款式编码"]'))
        if style_input is not None:
            break
        await asyncio.sleep(0.25)
    if style_input is None:
        raise AutomationError(f"商品中心在 {timeout_seconds} 秒内未加载款式编码查询框")

    await style_input.fill(style_code)
    query_button = await first_visible(page.get_by_role("button", name=re.compile(r"^\s*查询\s*$")))
    if query_button is None:
        raise AutomationError("商品中心中找不到“查询”按钮")
    await query_button.click()

    rows = page.locator(".el-table__body-wrapper tbody tr").filter(has_text=style_code)
    try:
        await rows.first.wait_for(state="visible", timeout=timeout_ms)
    except Exception as exc:
        raise AutomationError(f"页面查询后未找到款式编码 {style_code}") from exc

    visible_rows = []
    for index in range(await rows.count()):
        candidate = rows.nth(index)
        if await visible(candidate):
            visible_rows.append(candidate)
    if not visible_rows:
        raise AutomationError(f"找到了 {style_code} 的数据，但表格行不可见")
    if len(visible_rows) > 1:
        raise AutomationError(f"页面查到 {len(visible_rows)} 个款式编码 {style_code} 的商品，无法安全确定编辑对象")
    target_row = visible_rows[0]
    edit_link = await first_visible(target_row.get_by_text("编辑", exact=True))
    if edit_link is None:
        raise AutomationError(f"款式编码 {style_code} 所在行没有“编辑”入口")
    await edit_link.click()

    drawer = page.locator("#prod-center-edit-dialog")
    await drawer.wait_for(state="visible", timeout=timeout_ms)

    base_tab = await first_visible(drawer.get_by_text("基础资料", exact=True))
    if base_tab is not None:
        classes = await base_tab.get_attribute("class") or ""
        if "is-active" not in classes:
            await base_tab.click()
            await page.wait_for_timeout(500)

    logger.info("编辑抽屉已打开，等待基础资料加载完成（最多 %s 秒）", timeout_seconds)
    style_item = await form_item(drawer, "款式编码", timeout_seconds=timeout_seconds)
    style_input_in_drawer = style_item.locator("input").first
    ready_deadline = time.monotonic() + timeout_seconds
    last_value = ""
    while time.monotonic() < ready_deadline:
        try:
            last_value = (await style_input_in_drawer.input_value()).strip()
            loading_masks = await blocking_drawer_loading_mask_count(drawer)
            if last_value == style_code and loading_masks == 0:
                break
        except Exception:
            pass
        await asyncio.sleep(0.25)
    else:
        raise AutomationError(
            f"基础资料在 {timeout_seconds} 秒内未加载完成："
            f"期望款式编码 {style_code}，页面值 {last_value!r}"
        )
    logger.info("已打开 %s 的基础资料编辑页，表单加载完成", style_code)
    return drawer


async def wait_for_base_form_ready_after_save(
    drawer: Any,
    style_code: str,
    title: str,
    timeout_seconds: int,
) -> None:
    """保存后等待当前编辑抽屉稳定，不刷新或重新打开商品。"""
    timeout_ms = timeout_seconds * 1000
    await drawer.wait_for(state="visible", timeout=timeout_ms)
    deadline = time.monotonic() + timeout_seconds
    stable_since: Optional[float] = None
    last_style = ""
    last_title = ""
    while time.monotonic() < deadline:
        try:
            style_item = await form_item(drawer, "款式编码", timeout_seconds=0.5)
            title_item = await form_item(drawer, "商品名称", timeout_seconds=0.5)
            last_style = (
                await style_item.locator("input").first.input_value()
            ).strip()
            last_title = (
                await title_item.locator("input").first.input_value()
            ).strip()
            loading_masks = await blocking_drawer_loading_mask_count(drawer)
            if (
                last_style == style_code
                and last_title == title
                and loading_masks == 0
            ):
                if stable_since is None:
                    stable_since = time.monotonic()
                elif time.monotonic() - stable_since >= 0.5:
                    return
            else:
                stable_since = None
        except Exception:
            stable_since = None
        await asyncio.sleep(0.1)
    raise AutomationError(
        "基础资料保存后当前编辑页未稳定加载："
        f"款式编码={last_style!r}，商品名称={last_title!r}"
    )


async def form_item(
    drawer: Any,
    label: str,
    timeout_seconds: float = 300,
    poll_interval: float = 0.25,
) -> Any:
    """等待慢速编辑表单实际渲染出目标字段，而不是只等待抽屉标题。"""
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        try:
            items = drawer.locator(".el-form-item")
            for index in range(await items.count()):
                item = items.nth(index)
                labels = item.locator(":scope > .el-form-item__label, :scope > label.el-form-item__label")
                if not await labels.count():
                    continue
                text = (await labels.first.inner_text()).strip().rstrip("：:").strip()
                if text == label:
                    return item
        except Exception:
            # Vue 重新渲染表单时 locator 可能短暂失效，下一轮读取新 DOM。
            pass
        await asyncio.sleep(poll_interval)
    raise AutomationError(f"等待基础资料字段“{label}”加载超时（{timeout_seconds:g} 秒）")


async def fill_input(item: Any, value: str, label: str) -> Any:
    input_box = item.locator("input").first
    if not await input_box.count():
        raise AutomationError(f"字段“{label}”中找不到输入框")
    await item.scroll_into_view_if_needed()
    await input_box.fill(value)
    await input_box.press("Tab")
    actual = await input_box.input_value()
    if actual.strip() != value.strip():
        raise AutomationError(f"字段“{label}”填写后校验失败：{actual!r}")
    return input_box


async def delete_uploaded_images(scope: Any, page: Any) -> int:
    deleted = 0
    while True:
        images = scope.locator(".sc-upload .file-img")
        before = await images.count()
        if before == 0:
            return deleted

        image = images.first
        button = image.locator(".del-btn").first
        if not await button.count():
            return deleted

        # 快麦通仅在鼠标悬停图片时显示删除按钮。先走正常悬停交互；
        # 如果动画或遮罩层让按钮仍不可见，则在精确定位后触发同一 DOM click 事件。
        try:
            await image.hover(timeout=2000)
            if await visible(button):
                await button.click(timeout=2000)
            else:
                await button.evaluate("element => element.click()")
        except Exception:
            await button.evaluate("element => element.click()")

        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            if await images.count() < before:
                break
            await asyncio.sleep(0.1)
        else:
            raise AutomationError("点击图片删除按钮后，页面未移除原图")
        deleted += 1


async def wait_for_image_uploads(
    scope: Any,
    expected: int,
    label: str,
    timeout_seconds: int,
    *,
    retry_idle_upload: Any = None,
    idle_retry_seconds: float = 20,
) -> None:
    deadline = time.monotonic() + timeout_seconds
    idle_since = time.monotonic()
    error_since: Optional[float] = None
    error_logged = False
    retried = False
    last_count = -1
    while time.monotonic() < deadline:
        count = await scope.locator(".sc-upload .file-img").count()
        if count != last_count:
            logging.getLogger("kuaimai_erp").info("%s上传进度：%s/%s", label, count, expected)
            last_count = count
        upload_errors = await scope.locator(".file-input.error-warp").count()
        uploading = await scope.get_by_text("上传中", exact=True).count()
        # 页面在必填图片为空时就会预先加上 error-warp；它不是本次
        # 上传失败信号。先判断目标图片是否已经出现，再给前端上传/校验
        # 回调一个宽限窗口，避免把初始红框误判成接口错误。
        if count == expected and uploading == 0:
            return
        if upload_errors:
            if error_since is None:
                error_since = time.monotonic()
            if not error_logged:
                logging.getLogger("kuaimai_erp").info(
                    "%s检测到页面必填红框，等待上传结果后再判定（当前 %s/%s）",
                    label,
                    count,
                    expected,
                )
                error_logged = True
            if (
                time.monotonic() - error_since >= min(10.0, max(3.0, timeout_seconds * 0.1))
                and count == 0
                and uploading == 0
            ):
                raise AutomationError(
                    f"{label}上传失败，页面显示 {upload_errors} 个错误位"
                )
        else:
            error_since = None
        if count or uploading:
            idle_since = time.monotonic()
        elif retry_idle_upload is not None and not retried and time.monotonic() - idle_since >= idle_retry_seconds:
            retried = True
            logging.getLogger("kuaimai_erp").warning(
                "%s：持续 0/%s 且无上传中状态，重新定位上传控件并重试一次", label, expected
            )
            await retry_idle_upload()
        await asyncio.sleep(0.5)
    raise AutomationError(f"{label}上传超时，已完成 {last_count}/{expected}")


def is_blank_sku_placeholder(image_bytes: bytes) -> bool:
    """Return whether image bytes are the near-white default SKU placeholder.

    A decode failure is deliberately treated as a real/unknown image by the
    caller so that an existing user image is never overwritten on a guess.
    """
    if not image_bytes:
        return False
    try:
        import cv2
        import numpy as np

        encoded = np.frombuffer(image_bytes, dtype=np.uint8)
        image = cv2.imdecode(encoded, cv2.IMREAD_UNCHANGED)
        if image is None or image.size == 0:
            return False
        if image.ndim == 2:
            gray = image.astype(np.float32)
        else:
            if image.shape[2] == 4:
                alpha = image[:, :, 3].astype(np.float32) / 255.0
                color = image[:, :, :3].astype(np.float32)
                color = color * alpha[:, :, None] + 255.0 * (1.0 - alpha[:, :, None])
                gray = cv2.cvtColor(color.astype(np.uint8), cv2.COLOR_BGR2GRAY).astype(
                    np.float32
                )
            else:
                gray = cv2.cvtColor(image[:, :, :3], cv2.COLOR_BGR2GRAY).astype(
                    np.float32
                )
        return float(gray.mean()) >= 245.0 and float(gray.std()) <= 3.0
    except Exception:
        return False


async def sku_slot_has_real_image(page: Any, slot: Any) -> bool:
    """Safely distinguish a real SKU image from the ERP's blank placeholder."""
    images = slot.locator(".sc-upload .file-img")
    image_count = await images.count()
    if image_count == 0:
        return False
    if image_count != 1:
        logging.getLogger("kuaimai_erp").warning(
            "SKU 图位存在 %s 张图，无法唯一判定占位图，全部保留",
            image_count,
        )
        return True

    image = images.first.locator("img.originImg").first
    if not await image.count():
        image = images.first.locator("img").first
    if not await image.count():
        # Unknown existing markup is preserved rather than overwritten.
        return True
    source = (await image.get_attribute("src") or "").strip()
    if not source:
        return True

    response = None
    try:
        if source.startswith("data:image/"):
            header, separator, payload = source.partition(",")
            if not separator:
                return True
            if ";base64" in header.casefold():
                body = base64.b64decode(payload, validate=True)
            else:
                body = unquote_to_bytes(payload)
            return not is_blank_sku_placeholder(body)
        if source.startswith("blob:"):
            values = await page.evaluate(
                """async source => {
                  const response = await fetch(source);
                  if (!response.ok) throw new Error(`HTTP ${response.status}`);
                  return Array.from(new Uint8Array(await response.arrayBuffer()));
                }""",
                source,
            )
            return not is_blank_sku_placeholder(bytes(values))
        request_url = source
        if not urlsplit(source).scheme:
            request_url = urljoin(str(getattr(page, "url", "")), source)
        response = await page.context.request.get(
            request_url,
            fail_on_status_code=False,
            timeout=10_000,
        )
        if not response.ok:
            return True
        body = await response.body()
        return not is_blank_sku_placeholder(body)
    except Exception as exc:
        logging.getLogger("kuaimai_erp").warning(
            "SKU 已有图片无法安全识别，按真实图片保留：%s",
            type(exc).__name__,
        )
        return True
    finally:
        if response is not None:
            try:
                await response.dispose()
            except Exception:
                pass


async def replace_image_group(
    page: Any,
    item: Any,
    paths: Sequence[Path],
    label: str,
    timeout_seconds: int,
) -> None:
    """Backward-compatible alias for :func:`sync_image_group`."""
    await sync_image_group(page, item, paths, label, timeout_seconds)


async def sync_image_group(
    page: Any,
    item: Any,
    paths: Sequence[Path],
    label: str,
    timeout_seconds: int,
    *,
    force_replace: bool = False,
    normalize_small_images: bool = False,
) -> str:
    """Make one image group match the expected local image count and order.

    The page does not expose stable image identifiers, so count is the normal
    synchronization criterion.  A single near-white default placeholder is
    treated as empty even when its count matches, then replaced by the local
    image; unknown or multiple existing images remain untouched conservatively.
    ``force_replace`` is reserved for a platform field whose explicit image
    order differs from the inherited base-data order.
    """
    if not paths:
        raise AutomationError(f"{label}没有可上传图片")

    await item.scroll_into_view_if_needed()
    images = item.locator(".sc-upload .file-img")
    existing_count = await images.count()
    expected_count = len(paths)
    if existing_count == expected_count and not force_replace:
        # 属性图片和部分单图上传位会先渲染一张近白默认占位图；它在
        # DOM 上也是 ``.file-img``，不能仅凭数量判定已经有真实图片。
        # 只有单图组可以安全复用 SKU 占位图识别；识别失败时该函数会
        # 保守返回 True，继续保留未知图片而不是误删。
        if existing_count != 1 or await sku_slot_has_real_image(page, item):
            logging.getLogger("kuaimai_erp").info(
                "%s：页面已有 %s 张图片，与本地预期一致，跳过上传",
                label,
                expected_count,
            )
            return "skipped"
        logging.getLogger("kuaimai_erp").info(
            "%s：检测到默认空白占位图，将按本地图片替换",
            label,
        )
    elif force_replace and existing_count:
        logging.getLogger("kuaimai_erp").info(
            "%s：按平台指定图片顺序替换现有 %s 张图片",
            label,
            existing_count,
        )

    deleted = await delete_uploaded_images(item, page)
    if deleted:
        logging.getLogger("kuaimai_erp").info("%s：已删除 %s 张原图", label, deleted)
    inputs = item.locator('input[type="file"]')
    if not await inputs.count():
        raise AutomationError(f"{label}区域找不到本地上传控件")
    file_input = inputs.first
    accept = str(await file_input.get_attribute("accept") or "")
    accepted_suffixes = {
        token.strip().casefold()
        for token in accept.split(",")
        if token.strip().startswith(".")
    }
    accepts_any_image = any(
        token.strip().casefold() == "image/*" for token in accept.split(",")
    )
    with tempfile.TemporaryDirectory(prefix="kuaimai-upload-") as temp_dir:
        upload_paths = []
        converted_jfif = 0
        normalized_small_images = 0
        for index, path in enumerate(paths):
            suffix = path.suffix.casefold()
            if not accepted_suffixes or accepts_any_image or suffix in accepted_suffixes:
                upload_path = path
                if normalize_small_images and suffix in {".jpg", ".jpeg", ".png"}:
                    # 快麦水洗标控件会在前端先校验最小尺寸；过小的纯白
                    # 吊牌图会被标红且完全不发上传请求。只对该显式开启
                    # 的字段生成临时放大副本，不改动共享盘原文件。
                    try:
                        import cv2

                        image = cv2.imread(str(path), cv2.IMREAD_UNCHANGED)
                        if image is not None and min(image.shape[:2]) < 800:
                            scale = 800 / min(image.shape[:2])
                            height = max(800, int(round(image.shape[0] * scale)))
                            width = max(800, int(round(image.shape[1] * scale)))
                            resized = cv2.resize(
                                image,
                                (width, height),
                                interpolation=cv2.INTER_NEAREST,
                            )
                            target = Path(temp_dir) / f"normalized-{index + 1}{suffix}"
                            if cv2.imwrite(str(target), resized):
                                upload_path = target
                                normalized_small_images += 1
                                logging.getLogger("kuaimai_erp").info(
                                    "%s：第 %s 张图片尺寸 %sx%s 过小，"
                                    "已生成临时 %sx%s 副本上传（原文件不变）",
                                    label,
                                    index + 1,
                                    image.shape[1],
                                    image.shape[0],
                                    width,
                                    height,
                                )
                    except Exception as exc:
                        logging.getLogger("kuaimai_erp").warning(
                            "%s：小尺寸图片临时放大失败，继续使用原文件：%s",
                            label,
                            exc,
                        )
                upload_paths.append(upload_path)
                continue
            if (
                suffix == ".jfif"
                and accepted_suffixes.intersection({".jpg", ".jpeg"})
                and path.read_bytes()[:3] == b"\xff\xd8\xff"
            ):
                target = Path(temp_dir) / f"upload-{index + 1}.jpg"
                shutil.copyfile(path, target)
                upload_paths.append(target)
                converted_jfif += 1
                continue
            raise AutomationError(
                f"{label}图片格式 {suffix or '无扩展名'} 不在页面允许范围 {accept}"
            )
        if converted_jfif:
            logging.getLogger("kuaimai_erp").info(
                "%s：%s 张 JFIF 将按页面允许的 JPG 格式上传",
                label,
                converted_jfif,
            )
        async def trigger_upload() -> None:
            await item.scroll_into_view_if_needed()
            # Deleting the last image can recreate the upload component.
            await item.evaluate("e => new Promise(resolve => requestAnimationFrame(() => requestAnimationFrame(resolve)))")
            current_input = item.locator('input[type="file"]').first
            # 空上传位不先写入空文件。Element Upload 会把这次清空事件
            # 当成一次校验失败，导致后续真实文件事件被吞掉且不发请求。
            # 只有控件本身仍有文件时才清空，再写入新文件。
            try:
                current_count = await current_input.evaluate(
                    "input => input.files ? input.files.length : 0"
                )
            except Exception:
                current_count = 0
            if current_count:
                await current_input.set_input_files([])
            file_values = [str(path) for path in upload_paths]
            # 快麦当前上传组件的“本地上传”按钮会在点击时注册页面侧
            # change/upload 处理器。直接给隐藏 input 设置文件虽然能回读
            # files，但改版后的组件不会因此发起上传请求；优先走和人工
            # 操作相同的 file chooser 流程，找不到按钮时再回退旧路径。
            # “本地上传”文字通常只在上传占位图 hover 后显示；人工操作时
            # 鼠标移到图片区域会先触发这一层，所以脚本也要先 hover。
            upload_root = item.locator(".file-input:visible").first
            if not await upload_root.count():
                upload_root = item.locator(".sc-upload:visible").first
            if await upload_root.count():
                try:
                    logging.getLogger("kuaimai_erp").info(
                        "%s：准备悬停上传区域",
                        label,
                    )
                    await upload_root.hover(timeout=3000, force=True)
                    await page.wait_for_timeout(100)
                    logging.getLogger("kuaimai_erp").info(
                        "%s：上传区域已悬停",
                        label,
                    )
                except Exception:
                    logging.getLogger("kuaimai_erp").warning(
                        "%s：悬停上传区域失败，继续尝试文件控件",
                        label,
                    )
            upload_button = item.locator(".location-upload:visible").first
            used_file_chooser = False
            if await upload_button.count():
                try:
                    async with page.expect_file_chooser(timeout=5000) as chooser_info:
                        await upload_button.click()
                    chooser = await chooser_info.value
                    await chooser.set_files(file_values)
                    used_file_chooser = True
                    logging.getLogger("kuaimai_erp").info(
                        "%s：已通过“本地上传”按钮触发文件选择器",
                        label,
                    )
                except Exception as exc:
                    logging.getLogger("kuaimai_erp").warning(
                        "%s：文件选择器流程失败，回退直接写入控件：%s",
                        label,
                        exc,
                    )
            if not used_file_chooser:
                await current_input.set_input_files(file_values)
            try:
                file_count = await current_input.evaluate(
                    "input => input.files ? input.files.length : 0"
                )
                logging.getLogger("kuaimai_erp").info(
                    "%s：已将 %s 个本地文件写入页面上传控件",
                    label,
                    file_count,
                )
            except Exception as exc:
                logging.getLogger("kuaimai_erp").warning(
                    "%s：已调用文件上传，但无法回读控件文件数：%s",
                    label,
                    exc,
                )

        upload_events = []

        def record_upload_response(response: Any) -> None:
            try:
                request = response.request
                if request.resource_type in {"xhr", "fetch"}:
                    upload_events.append(
                        f"{request.method} {response.status} {response.url}"
                    )
            except Exception:
                pass

        def record_upload_failure(request: Any) -> None:
            try:
                upload_events.append(f"FAILED {request.method} {request.url}")
            except Exception:
                pass

        page.on("response", record_upload_response)
        page.on("requestfailed", record_upload_failure)
        try:
            await trigger_upload()
            await wait_for_image_uploads(
                item, expected_count, label, timeout_seconds,
                retry_idle_upload=trigger_upload if label == "小红书3:4主图" else None,
            )
        except Exception:
            if upload_events:
                logging.getLogger("kuaimai_erp").error(
                    "%s上传相关请求：%s", label, "；".join(upload_events[-12:])
                )
            raise
        finally:
            page.remove_listener("response", record_upload_response)
            page.remove_listener("requestfailed", record_upload_failure)
    return "replaced"


async def base_specification_value_group(drawer: Any, spec_name: str) -> Any:
    """Locate one base-data specification value group by its exact spec name."""
    result = await drawer.evaluate(
        """
        (root, expectedName) => {
          const marker = 'data-codex-base-spec-values';
          root.querySelectorAll(`[${marker}]`).forEach(node => node.removeAttribute(marker));
          const clean = value => String(value || '').replace(/[\\s:：]+/g, '');
          const visible = element => {
            const style = getComputedStyle(element);
            return style.display !== 'none' && style.visibility !== 'hidden'
              && element.getClientRects().length > 0;
          };
          const matches = [];
          const titles = [...root.querySelectorAll('.block-specification .title-bg')]
            .filter(visible);
          for (const title of titles) {
            const nameInput = title.querySelector('input');
            if (!nameInput || clean(nameInput.value) !== clean(expectedName)) continue;
            let values = title.parentElement?.querySelector(':scope > .specification-value');
            if (!values) {
              const block = title.closest('.block-specification');
              const blockTitles = block ? [...block.querySelectorAll('.title-bg')] : [];
              const blockValues = block ? [...block.querySelectorAll('.specification-value')] : [];
              const index = blockTitles.indexOf(title);
              values = index >= 0 ? blockValues[index] : null;
            }
            if (values && visible(values)) matches.push(values);
          }
          const unique = [...new Set(matches)];
          if (unique.length === 1) unique[0].setAttribute(marker, expectedName);
          return {titleCount: titles.length, matchCount: unique.length};
        }
        """,
        spec_name,
    )
    match_count = int(result.get("matchCount", 0))
    if match_count != 1:
        raise AutomationError(
            f"基础资料规格名“{spec_name}”无法唯一定位：匹配到 {match_count} 组"
        )
    return drawer.locator(
        f'[data-codex-base-spec-values="{spec_name}"]'
    ).first


async def read_base_specification_values(
    drawer: Any,
    spec_name: str,
) -> Tuple[str, ...]:
    group = await base_specification_value_group(drawer, spec_name)
    values = await group.locator(
        ".specification-value-flex_input input:not([type=checkbox])"
    ).evaluate_all("inputs => inputs.map(input => input.value.trim())")
    return tuple(str(value).strip() for value in values)


async def sync_base_color_spec_values(
    drawer: Any,
    expected_colors: Sequence[str],
    *,
    timeout_seconds: float = 5.0,
) -> Dict[str, Any]:
    """Backward-compatible wrapper for the generic specification reconciler."""
    return await sync_base_specification_values(
        drawer,
        "颜色",
        expected_colors,
        timeout_seconds=timeout_seconds,
    )


async def _delete_base_specification_value(
    drawer: Any,
    spec_name: str,
    index: int,
    *,
    timeout_seconds: float,
) -> None:
    """Delete exactly one visible specification value and verify the row count."""
    group = await base_specification_value_group(drawer, spec_name)
    before = await read_base_specification_values(drawer, spec_name)
    if not (0 <= index < len(before)):
        raise AutomationError(
            f"基础资料{spec_name}规格删除下标越界：{index} / {len(before)}"
        )
    marker = "data-codex-base-spec-delete"
    result = await group.evaluate(
        """
        (root, payload) => {
          root.querySelectorAll(`[${payload.marker}]`).forEach(node =>
            node.removeAttribute(payload.marker));
          const inputs = [...root.querySelectorAll(
            '.specification-value-flex_input input:not([type=checkbox])'
          )];
          const input = inputs[payload.index];
          if (!input) return {inputFound: false, candidateCount: 0};
          let item = input.closest('.specification-value-flex_item');
          if (!item) {
            let node = input.parentElement;
            while (node && node !== root) {
              if (node.querySelectorAll(
                '.specification-value-flex_input input:not([type=checkbox])'
              ).length === 1) {
                item = node;
                if (node.parentElement === root ||
                    node.parentElement?.classList.contains('specification-value-flex')) break;
              }
              node = node.parentElement;
            }
          }
          if (!item) return {inputFound: true, candidateCount: 0};
          const normalize = value => String(value || '').replace(/\s+/g, '').toLowerCase();
          const candidates = [...item.querySelectorAll('button,[role=button],i,span,svg')]
            .filter(node => {
              if (node.contains(input)) return false;
              const token = normalize([
                node.getAttribute('title'),
                node.getAttribute('aria-label'),
                node.className && String(node.className),
                node.textContent,
              ].join(' '));
              return /删除|移除|delete|remove|close|shanchu|jian|/.test(token);
            });
          const unique = [...new Set(candidates)].filter(node =>
            !candidates.some(other => other !== node && node.contains(other)));
          if (unique.length === 1) unique[0].setAttribute(payload.marker, '1');
          return {inputFound: true, candidateCount: unique.length};
        }
        """,
        {"marker": marker, "index": index},
    )
    if not result.get("inputFound") or int(result.get("candidateCount", 0)) != 1:
        raise AutomationError(
            f"基础资料{spec_name}第 {index + 1} 个规格值找不到唯一删除控件"
        )
    target = group.locator(f'[{marker}="1"]')
    try:
        # The close icon is intentionally hidden until hover on some ERP
        # builds.  Dispatch its native DOM click and use the resulting value
        # count as the authoritative success signal.
        await target.evaluate("node => node.click()")
    except Exception:
        current = await read_base_specification_values(drawer, spec_name)
        if len(current) != len(before) - 1:
            raise
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        current = await read_base_specification_values(drawer, spec_name)
        if len(current) == len(before) - 1:
            return
        await asyncio.sleep(0.05)
    raise AutomationError(f"删除基础资料{spec_name}规格值后数量未减少")


async def sync_base_specification_values(
    drawer: Any,
    spec_name: str,
    expected_values: Sequence[str],
    *,
    timeout_seconds: float = 5.0,
) -> Dict[str, Any]:
    """Reconcile one base specification to the ordered current-product values."""
    expected = tuple(
        str(value).strip() for value in expected_values if str(value).strip()
    )
    if not expected:
        raise AutomationError(f"当前商品没有可用于基础资料的{spec_name}规格值")
    if len({re.sub(r"\s+", "", value).casefold() for value in expected}) != len(expected):
        raise AutomationError(
            f"当前商品{spec_name}规格存在重复值，无法安全匹配 SKU 顺序"
        )

    before = await read_base_specification_values(drawer, spec_name)
    removed = 0
    # Remove surplus values from the end before rewriting the retained slots.
    # This makes the final sequence authoritative without depending on a fixed
    # pants/outerwear size template, while preserving the first-spec image order.
    while len(await read_base_specification_values(drawer, spec_name)) > len(expected):
        current = await read_base_specification_values(drawer, spec_name)
        await _delete_base_specification_value(
            drawer,
            spec_name,
            len(current) - 1,
            timeout_seconds=timeout_seconds,
        )
        removed += 1

    group = await base_specification_value_group(drawer, spec_name)
    inputs = group.locator(
        ".specification-value-flex_input input:not([type=checkbox])"
    )
    current_editability = await inputs.evaluate_all(
        "inputs => inputs.map(input => ({disabled: input.disabled, readOnly: input.readOnly}))"
    )
    if any(item.get("disabled") or item.get("readOnly") for item in current_editability):
        raise AutomationError(f"基础资料{spec_name}规格值存在不可编辑输入框，未强制覆盖")

    current_count = len(await read_base_specification_values(drawer, spec_name))
    missing = len(expected) - current_count
    if missing:
        add_buttons = group.get_by_role("button", name="添加规格值", exact=True)
        if await add_buttons.count() != 1:
            raise AutomationError(
                f"基础资料还需新增 {missing} 个{spec_name}，但找不到唯一的“添加规格值”按钮"
            )
        for _index in range(missing):
            previous_count = len(await read_base_specification_values(drawer, spec_name))
            await add_buttons.click()
            deadline = time.monotonic() + timeout_seconds
            while time.monotonic() < deadline:
                current = await read_base_specification_values(drawer, spec_name)
                if len(current) == previous_count + 1:
                    break
                await asyncio.sleep(0.05)
            else:
                raise AutomationError(
                    f"点击“添加规格值”后，基础资料没有新增{spec_name}输入框"
                )
            group = await base_specification_value_group(drawer, spec_name)
            add_buttons = group.get_by_role("button", name="添加规格值", exact=True)

    changed = 0
    for index, expected_value in enumerate(expected):
        group = await base_specification_value_group(drawer, spec_name)
        inputs = group.locator(
            ".specification-value-flex_input input:not([type=checkbox])"
        )
        if await inputs.count() != len(expected):
            raise AutomationError(f"基础资料{spec_name}规格在填写过程中数量发生变化")
        input_box = inputs.nth(index)
        current_value = (await input_box.input_value()).strip()
        if current_value == expected_value:
            continue
        if await input_box.is_disabled() or await input_box.is_editable() is False:
            raise AutomationError(
                f"基础资料第 {index + 1} 个{spec_name}规格值不可编辑"
            )
        await input_box.fill(expected_value)
        await input_box.press("Tab")
        deadline = time.monotonic() + timeout_seconds
        while time.monotonic() < deadline:
            actual = await read_base_specification_values(drawer, spec_name)
            if len(actual) == len(expected) and actual[index] == expected_value:
                break
            await asyncio.sleep(0.05)
        else:
            raise AutomationError(
                f"基础资料{spec_name}规格值填写后回读失败：期望 {expected_value!r}"
            )
        changed += 1

    after = await read_base_specification_values(drawer, spec_name)
    if after != expected:
        raise AutomationError(
            f"基础资料{spec_name}规格与当前商品顺序不一致："
            f"页面 {after}，目标 {expected}"
        )
    return {
        "before": before,
        "after": after,
        "changed": changed,
        "added": missing,
        "removed": removed,
    }


async def replace_sku_images(
    page: Any,
    item: Any,
    paths: Sequence[Path],
    timeout_seconds: int,
) -> int:
    """Upload only missing SKU images and preserve images already on the page."""
    await item.scroll_into_view_if_needed()
    first_spec = item.locator(".block-specification .specification-value").first
    slots = first_spec.locator(".specification-value-flex_img")
    slot_count = await slots.count()
    if slot_count == 0:
        raise AutomationError("商品规格中找不到 SKU 图上传位")
    if slot_count != len(paths):
        raise AutomationError(
            f"SKU 图数量（{len(paths)}）与第一规格值数量（{slot_count}）不一致，"
            "为避免图片和颜色错配，程序已停止。"
        )
    uploaded_count = 0
    for index, path in enumerate(paths):
        slot = slots.nth(index)
        if await slot.locator(".sc-upload .file-img").count():
            if await sku_slot_has_real_image(page, slot):
                logging.getLogger("kuaimai_erp").info(
                    "SKU 图 %s：页面已有真实图片，保留并跳过上传",
                    index + 1,
                )
                continue
            deleted = await delete_uploaded_images(slot, page)
            if deleted != 1:
                raise AutomationError(
                    f"第 {index + 1} 个 SKU 图位识别为空白占位图，但未能安全删除"
                )
            logging.getLogger("kuaimai_erp").info(
                "SKU 图 %s：已移除默认空白占位图",
                index + 1,
            )
        file_input = slot.locator('input[type="file"]')
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline and not await file_input.count():
            await asyncio.sleep(0.1)
        if not await file_input.count():
            raise AutomationError(f"第 {index + 1} 个 SKU 图位找不到本地上传控件")
        await file_input.first.set_input_files(str(path))
        await wait_for_image_uploads(slot, 1, f"SKU 图 {index + 1}", timeout_seconds)
        uploaded_count += 1
    return uploaded_count


async def set_base_price(drawer: Any, price: str) -> Any:
    block = await first_visible(drawer.locator(".block-specification-list"))
    if block is None:
        raise AutomationError("页面没有可见的基础资料规格明细")
    try:
        await block.scroll_into_view_if_needed(timeout=10_000)
    except Exception as exc:
        raise AutomationError("基础资料规格明细在 10 秒内无法滚动到可见区域") from exc
    batch_items = block.locator(".sku-batch-item")
    price_item = None
    for index in range(await batch_items.count()):
        candidate = batch_items.nth(index)
        label = candidate.locator(".sku-batch-item_label")
        if await label.count() and "基本售价" in (await label.inner_text()):
            price_item = candidate
            break
    if price_item is None:
        raise AutomationError("规格明细中找不到“基本售价”批量输入框")

    input_box = price_item.locator("input").first

    async def rows_match() -> bool:
        """在页面内一次扫描大表格，避免虚拟滚动表逐元素等待。"""
        try:
            values = await block.evaluate(
                """
                (root) => {
                  const visible = element => {
                    const style = getComputedStyle(element);
                    return style.display !== 'none' && style.visibility !== 'hidden'
                      && element.getClientRects().length > 0;
                  };
                  for (const table of root.querySelectorAll('.el-table')) {
                    if (!visible(table)) continue;
                    const headerRoot = table.querySelector(':scope > .el-table__main-wrapper > .el-table__header-wrapper')
                      || table.querySelector(':scope > .el-table__header-wrapper');
                    const bodyRoot = table.querySelector(':scope > .el-table__main-wrapper > .el-table__body-wrapper')
                      || table.querySelector(':scope > .el-table__body-wrapper');
                    if (!headerRoot || !bodyRoot) continue;
                    const headers = [...headerRoot.querySelectorAll('thead th')];
                    const indexes = headers.map((header, index) => {
                      const cell = header.querySelector(':scope > .cell');
                      const text = (cell?.getAttribute('title') || cell?.innerText || header.innerText || '')
                        .replace(/[\\s*:：]+/g, '');
                      return text.startsWith('基本售价') ? index : -1;
                    }).filter(index => index >= 0);
                    if (indexes.length !== 1) continue;
                    const rows = [...bodyRoot.querySelectorAll('tbody > tr')];
                    if (!rows.length) continue;
                    const result = [];
                    for (const row of rows) {
                      const cell = row.querySelectorAll(':scope > td')[indexes[0]];
                      const inputs = cell ? cell.querySelectorAll('input:not([disabled])') : [];
                      if (inputs.length !== 1) return [];
                      result.push(inputs[0].value);
                    }
                    return result;
                  }
                  return [];
                }
                """,
                timeout=10_000,
            )
        except Exception as exc:
            raise AutomationError("基本售价 SKU 列在 10 秒内无法读取") from exc
        if not values:
            return False
        try:
            return all(Decimal(str(value)) == Decimal(price) for value in values)
        except InvalidOperation:
            return False

    if await rows_match():
        await input_box.fill(price, timeout=10_000)
        return input_box

    try:
        await input_box.fill(price, timeout=10_000)
        await input_box.press("Tab", timeout=10_000)
    except Exception as exc:
        raise AutomationError("基本售价输入框在 10 秒内无法填写") from exc
    batch_button = await first_visible(block.get_by_role("button", name=re.compile(r"^\s*批量设置\s*$")))
    if batch_button is None:
        raise AutomationError("规格明细中找不到“批量设置”按钮")
    try:
        await batch_button.click(timeout=10_000)
    except Exception as exc:
        # 部分版本在 Tab 失焦后已经同步到 SKU，此时按钮可能不再可点。
        if await rows_match():
            return input_box
        raise AutomationError("基本售价“批量设置”在 10 秒内无法点击") from exc
    await asyncio.sleep(0.5)
    if not await rows_match():
        raise AutomationError(f"基本售价逐行校验失败：期望所有 SKU 为 {price}")
    return input_box


async def collect_visible_errors(drawer: Any) -> List[str]:
    errors: List[str] = []
    locator = drawer.locator(".el-form-item__error:visible")
    for index in range(await locator.count()):
        text = (await locator.nth(index).inner_text()).strip()
        if text and text not in errors:
            errors.append(text)
    return errors


async def click_save_and_confirm(
    page: Any,
    drawer: Any,
    sync_erp: bool,
    timeout_seconds: int,
    logger: logging.Logger,
    button_text: str = "保存",
    creation: bool = False,
) -> Dict[str, Any]:
    """单次点击指定保存按钮，通过同源编辑接口或成功提示确认。"""
    save_responses: List[Any] = []

    def on_response(response: Any) -> None:
        if "/item/base/edit.json" in response.url:
            save_responses.append(response)

    page.on("response", on_response)
    save_buttons = drawer.locator(".drawer-footer button")
    save_button = None
    for index in range(await save_buttons.count()):
        candidate = save_buttons.nth(index)
        if not await visible(candidate):
            continue
        text = re.sub(r"\s+", "", (await candidate.inner_text()))
        if text == re.sub(r"\s+", "", button_text):
            save_button = candidate
            break
    if save_button is None:
        raise AutomationError(f"商品编辑页找不到“{button_text}”按钮")

    await save_button.scroll_into_view_if_needed()
    await save_button.click()
    logger.info("已点击“%s”，等待后端确认", button_text)
    started_at = time.monotonic()
    deadline = time.monotonic() + timeout_seconds
    sync_dialog_handled = False
    unbound_dialog_handled = False
    distribution_relation_dialog_handled = False
    publish_dialog_handled = False
    success_toast_seen_at: Optional[float] = None

    while time.monotonic() < deadline:
        if creation:
            created = page.locator('.el-dialog:visible').filter(has_text="创建成功")
            if await visible(created.first):
                return {"result": 1, "confirmed_by": "creation_dialog", "action": button_text}
        if save_responses:
            response = save_responses[-1]
            try:
                payload = await response.json()
            except Exception:
                payload = {"http_status": response.status}
            if int(payload.get("result", 0) or 0) == 1:
                return {
                    "result": 1,
                    "confirmed_by": "api",
                    "action": button_text,
                    "payload": payload,
                }
            raise AutomationError(
                "保存接口返回失败："
                + str(payload.get("message") or payload.get("errmsg") or payload)
            )

        if button_text == "保存并铺货到平台":
            shop_dialog = page.locator('[role="dialog"]:visible').filter(
                has_text="铺货到店铺"
            ).first
            if await visible(shop_dialog):
                return {
                    "result": 1,
                    "confirmed_by": "publish_dialog_open",
                    "action": button_text,
                }

        success = page.locator(".el-message--success").filter(
            has_text=re.compile("保存成功|铺货成功|已提交铺货")
        )
        if await visible(success.first):
            if success_toast_seen_at is None:
                success_toast_seen_at = time.monotonic()
            elif time.monotonic() - success_toast_seen_at >= 2.0:
                # 某些旧页面没有可监听的保存接口；保留成功提示兜底，
                # 但先给真实接口响应留出时间，优先产出后端成功凭证。
                return {"result": 1, "confirmed_by": "toast", "action": button_text}

        sync_dialog = page.locator(".el-dialog:visible").filter(has_text="以下商品将同步至系统资料")
        if not sync_dialog_handled and await visible(sync_dialog.first):
            action = "确 定" if sync_erp else "跳过同步"
            button = sync_dialog.first.get_by_role("button", name=re.compile(r"^\s*" + re.escape(action).replace(r"\ ", r"\s*") + r"\s*$"))
            if not await button.count() and sync_erp:
                button = sync_dialog.first.get_by_text("确 定", exact=True)
            if not await button.count():
                raise AutomationError(f"ERP 同步确认框中找不到“{action}”按钮")
            await button.first.click()
            sync_dialog_handled = True
            logger.info("ERP 同步确认：%s", "同步" if sync_erp else "跳过同步，仅保存快麦通资料")

        unbound_dialog = page.locator(".el-dialog:visible").filter(has_text="商品规格未关联ERP商品")
        if not unbound_dialog_handled and await visible(unbound_dialog.first):
            await unbound_dialog.first.get_by_role("button", name="我知道了").click()
            unbound_dialog_handled = True
            logger.warning("页面提示 SKU 未关联 ERP，已按页面流程继续保存")

        distribution_relation_dialog = page.locator(".el-message-box:visible").filter(
            has_text=re.compile(
                r"商品已加入分销小店[\s\S]*"
                r"规格明细已变更[\s\S]*"
                r"是否确认保存商品资料"
            )
        )
        if (
            not distribution_relation_dialog_handled
            and await visible(distribution_relation_dialog.first)
        ):
            confirm = distribution_relation_dialog.first.get_by_role(
                "button", name=re.compile(r"^\s*(?:确定|确\s*定)\s*$")
            )
            if await confirm.count() != 1:
                raise AutomationError("分销关系变更确认框中找不到唯一的“确定”按钮")
            await confirm.click()
            distribution_relation_dialog_handled = True
            logger.info("已确认分销关系变更，继续保存商品资料")

        if button_text == "保存并铺货到平台" and not publish_dialog_handled:
            # 这里只处理保存后的简短确认框；“铺货到店铺”大弹窗必须在选择并
            # 校验店铺后单独提交，不能用唯一“确定”按钮直接略过。
            publish_dialog = page.locator(".el-message-box:visible").filter(
                has_text=re.compile("确认.*铺货|铺货.*平台")
            )
            if await visible(publish_dialog.first):
                confirm = publish_dialog.first.get_by_role(
                    "button", name=re.compile(r"^\s*(?:确定|确 定)\s*$")
                )
                if await confirm.count() != 1:
                    raise AutomationError("铺货确认框中找不到唯一的“确定”按钮")
                await confirm.click()
                publish_dialog_handled = True
                logger.info("已确认铺货对话框")

        duplicate_dialog = page.locator(".el-message-box:visible").filter(has_text=re.compile("编码.*重复|存在SPU编码相同"))
        if await visible(duplicate_dialog.first):
            raise AutomationError("保存时出现编码重复确认框，程序未自动创建新编码，请人工复核")

        if time.monotonic() - started_at > 3:
            errors = await collect_visible_errors(drawer)
            if errors:
                raise AutomationError(f"“{button_text}”被页面校验拦截：" + "；".join(errors))

        await asyncio.sleep(0.25)

    errors = await collect_visible_errors(drawer)
    if errors:
        raise AutomationError(f"“{button_text}”超时，页面校验错误：" + "；".join(errors))
    warnings: List[str] = []
    warning_locator = page.locator(".el-message--warning:visible, .el-message--error:visible")
    for index in range(await warning_locator.count()):
        text = (await warning_locator.nth(index).inner_text()).strip()
        if text:
            warnings.append(text)
    suffix = "：" + "；".join(warnings) if warnings else ""
    raise AutomationError(f"等待“{button_text}”成功确认超时" + suffix)


async def open_new_product_drawer(page: Any, timeout_seconds: int) -> Any:
    # 新增抽屉没有编辑页的 #prod-center-edit-dialog，按自己的标题定位。
    drawer = page.locator('.el-drawer:visible').filter(
        has=page.get_by_text("手工新增商品", exact=True)
    )
    if await drawer.count():
        raise AutomationError("页面已有未完成的手工新增商品，请先关闭或处理后重试")
    await page.get_by_role("button", name=re.compile(r"^新增商品")).click()
    new_product_item = page.locator(
        ".el-popover:visible .drop-list-item"
    ).filter(has_text=re.compile(r"^\s*手工新增商品\s*$"))
    await new_product_item.wait_for(state="visible", timeout=timeout_seconds * 1000)
    if await new_product_item.count() != 1:
        raise AutomationError("新增商品下拉菜单中找不到唯一的“手工新增商品”入口")
    await new_product_item.click()
    await drawer.wait_for(state="visible", timeout=timeout_seconds * 1000)
    return drawer


async def fill_new_product_form(page: Any, drawer: Any, product: ProductData, args: Any) -> Dict[str, Any]:
    await fill_input(await form_item(drawer, "款式编码", args.timeout), product.style_code, "款式编码")
    await fill_input(await form_item(drawer, "商品名称", args.timeout), "1", "商品名称")
    await sync_image_group(
        page, await form_item(drawer, "商品主图", args.timeout),
        product.main_images[:1], "新增主图", args.upload_timeout,
    )
    spec = drawer.locator('.block-specification')
    names = spec.locator('.title-bg > .el-input > input')
    if await names.count() != 2 or await names.nth(0).input_value() != "颜色" or await names.nth(1).input_value() != "尺码":
        raise AutomationError("新增页面默认规格不是颜色、尺码，请复核页面")
    values = spec.locator('.specification-value')
    color_report = await sync_base_color_spec_values(drawer, product.colors)

    # 按用户指定的两个默认模板操作，不逐个录入尺码、不改编码规则。
    await spec.get_by_role("button", name="填充常用规格").nth(1).click()
    template = page.locator('.el-dialog:visible').filter(
        has=page.locator('.el-dialog__title').filter(has_text=re.compile(r"^常用规格$"))
    )
    await template.wait_for(state="visible", timeout=args.timeout * 1000)
    apply = template.get_by_text("应用至资料", exact=True).first
    await apply.wait_for(state="visible", timeout=args.timeout * 1000)
    await apply.click()
    await template.wait_for(state="hidden", timeout=args.timeout * 1000)
    sizes = await values.nth(1).locator('.specification-value-flex_input input').evaluate_all(
        "inputs => inputs.map(input => input.value)"
    )
    if tuple(sizes) != NEW_PRODUCT_SIZES:
        raise AutomationError(f"默认尺码模板与视频不一致：{sizes}，未保存")
    await drawer.locator('.block-specification-list').get_by_role("button", name="批量生成", exact=True).click()
    generator = page.locator('.el-dialog:visible').filter(
        has=page.locator('.el-dialog__title').filter(has_text=re.compile(r"^批量生成商品编码$"))
    )
    await generator.wait_for(state="visible", timeout=args.timeout * 1000)
    rule = await generator.locator('input[type="radio"]:checked').evaluate(
        "input => input.closest('label').innerText.trim()"
    )
    await generator.get_by_role("button", name=re.compile(r"^确\s*定$")).click()
    await generator.wait_for(state="hidden", timeout=args.timeout * 1000)
    return {
        "colors": list(color_report["after"]),
        "color_specification": color_report,
        "size_template": sizes,
        "code_rule": rule,
    }


async def validate_new_product_form(drawer: Any, product: ProductData) -> Dict[str, Any]:
    for label, expected in (("款式编码", product.style_code), ("商品名称", "1")):
        value = await (await form_item(drawer, label)).locator('input').first.input_value()
        if value != expected:
            raise AutomationError(f"新增复核失败：{label}与预期不符")
    image_count = await (await form_item(drawer, "商品主图")).locator('.sc-upload .file-img').count()
    if image_count != 1:
        raise AutomationError(f"新增主图应为 1 张，实际 {image_count} 张")
    rows = await drawer.locator('.block-specification-list').evaluate(
        """root => {
          const table = [...root.querySelectorAll('.el-table')].find(e => e.getClientRects().length);
          if (!table) return [];
          const header = table.querySelector('.el-table__header-wrapper');
          const body = table.querySelector('.el-table__body-wrapper');
          if (!header || !body) return [];
          const labels = [...header.querySelectorAll('th')].map(e =>
            (e.querySelector('.cell')?.getAttribute('title') || e.innerText).replace(/[\\s*：:]/g, ''));
          return [...body.querySelectorAll('tbody > tr')].map(row => {
            const result = {};
            [...row.querySelectorAll(':scope > td')].forEach((cell, i) => {
              if (labels[i]) result[labels[i]] = cell.querySelector('input[type="text"]')?.value ?? cell.innerText.trim();
            });
            return result;
          });
        }"""
    )
    expected_rows = {
        (color, size): product.style_code + color + size
        for color in product.colors
        for size in NEW_PRODUCT_SIZES
    }
    if len(rows) != len(expected_rows):
        raise AutomationError(
            f"新增 SKU 行数应为 {len(expected_rows)}，实际 {len(rows)}"
        )
    seen = set()
    for row in rows:
        color = str(row.get("颜色") or "").strip()
        size = str(row.get("尺码") or "").strip()
        key = (color, size)
        code = expected_rows.get(key)
        if code is None or key in seen or row.get("商品编码") != code:
            raise AutomationError(
                f"新增 SKU 规格或默认编码与 Excel 颜色不一致：{color}/{size}"
            )
        seen.add(key)
        for label in ("基本售价", "销售价", "市场价", "成本价", "库存", "重量(kg)"):
            try:
                matches = Decimal(row.get(label, "")) == 0
            except InvalidOperation:
                matches = False
            if not matches:
                raise AutomationError(f"新增 SKU {color}/{size} 的{label}未保持视频默认值 0")
    errors = await collect_visible_errors(drawer)
    if errors:
        raise AutomationError("新增页面存在校验错误：" + "；".join(errors))
    return {
        "title": "1",
        "colors": list(product.colors),
        "main_image_count": image_count,
        "sku_count": len(rows),
        "rows": rows,
    }


async def run_new_product(page: Any, args: Any, product: ProductData, artifact_dir: Path, logger: logging.Logger) -> None:
    report: Dict[str, Any] = {"style_code": product.style_code, "saved": False, "published": False}
    result_path = artifact_dir / "create-product-result.json"
    try:
        record = await api_find_product(page, product.style_code, logger, strict=True)
        if record is not None:
            report.update(status="already_exists", confirmed_by="query", base_item_id=record.get("baseItemId") or record.get("id"))
            logger.info("款式编码 %s 已存在，本次不重复新增、不修改原商品", product.style_code)
            return
        drawer = await open_new_product_drawer(page, args.timeout)
        report.update(await fill_new_product_form(page, drawer, product, args))
        report["before_save"] = await validate_new_product_form(drawer, product)
        await safe_screenshot(page, artifact_dir / "create-product-preview.png")
        if not args.save:
            report.update(status="preview", confirmed_by="form_readback")
            logger.info(
                "新增链接预览完成，未保存：名称 1、Excel 颜色 %s、默认尺码、默认商品编码",
                " / ".join(product.colors),
            )
            return
        # 上传可能耗时，提交前再次查重。查询失败或重复均不点击保存。
        if await api_find_product(page, product.style_code, logger, strict=True) is not None:
            raise AutomationError("填写期间该款式已被创建，停止保存以避免重复")
        report["save_attempted"] = True
        result = await click_save_and_confirm(page, drawer, args.sync_erp, args.timeout, logger, creation=True)
        report.update(
            status="created",
            saved=True,
            confirmed_by=result.get("confirmed_by"),
            save_confirmation=result.get("confirmed_by"),
        )
        await safe_screenshot(page, artifact_dir / "create-product-after-save.png")
        logger.info(
            "新增链接保存成功，已由创建成功提示确认；按配置不再刷新和重开复核（%s 个 SKU）",
            len(product.colors) * len(NEW_PRODUCT_SIZES),
        )
    except Exception as exc:
        report.update(status="verification_required" if report.get("save_attempted") else "failed", error=str(exc))
        raise
    finally:
        result_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


async def publish_to_selected_douyin_shops(
    page: Any,
    selected_shops: Sequence[str],
    timeout_seconds: int,
    logger: logging.Logger,
) -> Dict[str, Any]:
    """在铺货弹窗中精确选择抖音店铺，确认店铺资料后提交铺货。"""
    expected = tuple(dict.fromkeys(shop.strip() for shop in selected_shops if shop.strip()))
    if not expected:
        raise AutomationError("没有指定要铺货的抖音店铺")

    dialog = page.locator('[role="dialog"]:visible').filter(has_text="铺货到店铺").first
    await dialog.wait_for(state="visible", timeout=timeout_seconds * 1000)
    douyin_platform = dialog.get_by_text("抖音", exact=True)
    if not await douyin_platform.count():
        raise AutomationError("铺货弹窗中找不到抖音平台")
    await douyin_platform.first.click()
    logger.info("铺货弹窗已选择抖音平台")

    async def checkbox_by_text(label: str) -> tuple[Any, Any]:
        candidates = dialog.locator(".el-checkbox").filter(has_text=label)
        for index in range(await candidates.count()):
            candidate = candidates.nth(index)
            text = re.sub(r"\s+", " ", (await candidate.inner_text()).strip())
            if text == label or text.startswith(label):
                checkbox = candidate.locator('input[type="checkbox"]').first
                if await checkbox.count():
                    return candidate, checkbox
        raise AutomationError(f"铺货弹窗中找不到唯一复选框：{label}")

    # 等待店铺文字出现，不依赖 Element UI 未绑定的 accessibility name。
    first_shop_visible = None
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        try:
            first_shop_visible, _ = await checkbox_by_text(KNOWN_DOUYIN_SHOPS[0])
            if await first_shop_visible.is_visible():
                break
        except AutomationError:
            first_shop_visible = None
        await asyncio.sleep(0.25)
    if first_shop_visible is None:
        raise AutomationError("抖音平台店铺列表未加载")
    unknown = [shop for shop in expected if shop not in KNOWN_DOUYIN_SHOPS]
    if unknown:
        raise AutomationError("未配置的抖音店铺：" + "、".join(unknown))

    async def has_published_tag(control: Any, shop: str) -> bool:
        container = control
        own_text = re.sub(r"\s+", " ", (await control.inner_text()).strip())
        if "已铺货" in own_text:
            return True
        for _ in range(4):
            container = container.locator("xpath=..").first
            text = re.sub(r"\s+", " ", (await container.inner_text()).strip())
            shops_in_container = [name for name in KNOWN_DOUYIN_SHOPS if name in text]
            if "已铺货" in text and shop in text and len(shops_in_container) == 1:
                return True
            if len(shops_in_container) > 1:
                break
        return False

    already_published: List[str] = []
    pending_expected: List[str] = []
    for shop in KNOWN_DOUYIN_SHOPS:
        control, checkbox = await checkbox_by_text(shop)
        is_published = await checkbox.is_disabled() or await has_published_tag(control, shop)
        logger.info("店铺状态：%s=%s", shop, "已铺货" if is_published else "可铺货")
        if is_published:
            if shop in expected:
                already_published.append(shop)
            continue
        should_check = shop in expected
        if should_check:
            pending_expected.append(shop)
        if await checkbox.is_checked() != should_check:
            await control.scroll_into_view_if_needed()
            await control.click()
            if await checkbox.is_checked() != should_check:
                raise AutomationError(f"店铺“{shop}”复选框状态切换失败")

    ai_control, ai_checkbox = await checkbox_by_text("使用AI裂变规则")
    if await ai_checkbox.is_checked():
        await ai_control.scroll_into_view_if_needed()
        await ai_control.click()
        if await ai_checkbox.is_checked():
            raise AutomationError("AI 裂变规则复选框无法关闭")

    actual_list: List[str] = []
    for shop in KNOWN_DOUYIN_SHOPS:
        control, checkbox = await checkbox_by_text(shop)
        if not await has_published_tag(control, shop) and await checkbox.is_checked():
            actual_list.append(shop)
    actual = tuple(actual_list)
    if set(actual) != set(pending_expected):
        raise AutomationError(
            "待铺货店铺复核失败：期望 " + "、".join(pending_expected)
            + "；实际 " + "、".join(actual)
        )
    if not pending_expected:
        await close_publish_preview(dialog)
        logger.info(
            "指定的抖音店铺均已铺货，已安全关闭弹窗并跳过重复提交：%s",
            "、".join(already_published),
        )
        return {
            "platform": "抖音",
            "shops": list(expected),
            "submitted_shops": [],
            "already_published": already_published,
            "result": 1,
            "confirmed_by": "already_published",
            "submitted": False,
        }
    logger.info(
        "抖音铺货店铺已精确复核：待铺货=%s；已铺货跳过=%s",
        "、".join(actual),
        "、".join(already_published) or "无",
    )

    # 用户指定的真实流程：选完店铺直接确定，不进入“确认发布信息”内页。
    final_button = dialog.get_by_role("button", name=re.compile(r"^\s*确\s*定\s*$"))
    if await final_button.count() != 1:
        raise AutomationError("铺货到店铺弹窗中找不到唯一的“确定”按钮")
    submit_responses: List[Any] = []

    def on_submit_response(response: Any) -> None:
        request = response.request
        if request.method == "POST" and any(
            marker in response.url.casefold()
            for marker in ("/item/base/edit.json", "publish", "distribute", "supply")
        ):
            submit_responses.append(response)

    page.on("response", on_submit_response)
    await final_button.click()
    logger.info("已点击铺货弹窗最终“确定”")

    continue_button = None
    completion_dialog = page.locator(
        '[role="dialog"]:visible, .el-dialog:visible, .el-message-box:visible'
    ).filter(has_text=re.compile("铺货完成"))
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        if await visible(completion_dialog.first):
            await asyncio.sleep(2)
            completion_text = re.sub(
                r"\s+", " ", (await completion_dialog.first.inner_text()).strip()
            )
            counts = re.search(
                r"商品总数\s*[：:]\s*(\d+)\s*款.*?成功\s*(\d+)\s*款.*?失败\s*(\d+)\s*款",
                completion_text,
            )
            if counts is None:
                raise AutomationError("无法解析抖音铺货完成结果：" + completion_text)
            total, succeeded, failed = (int(value) for value in counts.groups())
            if failed != 0 or succeeded != total:
                raise AutomationError(
                    f"抖音铺货未全部成功：总数 {total}，成功 {succeeded}，失败 {failed}"
                )
            close_button = completion_dialog.first.locator(
                ".el-dialog__headerbtn, .el-message-box__headerbtn, "
                'button[aria-label="Close"], button[aria-label="close"]'
            ).first
            if await visible(close_button):
                await close_button.click()
            logger.info(
                "已确认抖音铺货完成：总数 %s，成功 %s，失败 %s",
                total,
                succeeded,
                failed,
            )
            return {
                "result": 1,
                "confirmed_by": "completion_dialog",
                "shops": list(expected),
                "submitted_shops": list(actual),
                "already_published": already_published,
                "total": total,
                "succeeded": succeeded,
                "failed": failed,
                "responses": [response.url for response in submit_responses],
            }
        continue_button = await first_visible(
            page.get_by_role("button", name=re.compile("继续铺货"))
        )
        if continue_button is not None:
            break
        page_errors = page.locator(".el-message--error:visible")
        if await visible(page_errors.first):
            raise AutomationError("铺货提交失败：" + (await page_errors.first.inner_text()).strip())
        progress = await dismiss_publish_progress_dialog(page, logger, settle_seconds=2)
        if progress['dismissed'] or not await dialog.is_visible():
            if not progress['dismissed']:
                await asyncio.sleep(2)
            logger.info("抖音铺货已提交，页面稳定 2 秒，继续后续平台；后台结果不在此等待")
            return {
                "result": 1, "confirmed_by": "background_task_submitted",
                "submitted": True, "shops": list(expected),
                "submitted_shops": list(actual), "already_published": already_published,
                "responses": [response.url for response in submit_responses],
            }
        await asyncio.sleep(0.25)
    if continue_button is None:
        raise AutomationError("点击“确定”后未出现“继续铺货”确认弹窗")
    await continue_button.click()
    logger.info("已点击确认弹窗“继续铺货”")

    success_pattern = re.compile("铺货成功|提交成功|已提交铺货|任务已创建")
    success = page.locator(".el-message--success:visible").filter(has_text=success_pattern)
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        for response in submit_responses:
            try:
                payload = await response.json()
            except Exception:
                continue
            if "result" in payload and int(payload.get("result", 0) or 0) != 1:
                raise AutomationError(
                    "铺货接口返回失败："
                    + str(payload.get("message") or payload.get("errmsg") or payload)
                )
        if await visible(success.first):
            await asyncio.sleep(2)
            return {
                "result": 1,
                "confirmed_by": "toast",
                "shops": list(expected),
                "submitted_shops": list(actual),
                "already_published": already_published,
                "responses": [response.url for response in submit_responses],
            }
        progress = await dismiss_publish_progress_dialog(page, logger, settle_seconds=2)
        if progress['dismissed'] or (not await dialog.is_visible() and not await continue_button.is_visible()):
            if not progress['dismissed']:
                await asyncio.sleep(2)
            return {
                "result": 1,
                "confirmed_by": "background_task_submitted",
                "shops": list(expected),
                "submitted_shops": list(actual),
                "already_published": already_published,
                "responses": [response.url for response in submit_responses],
            }
        page_errors = page.locator(".el-message--error:visible")
        if await visible(page_errors.first):
            raise AutomationError("铺货提交失败：" + (await page_errors.first.inner_text()).strip())
        await asyncio.sleep(0.25)
    raise AutomationError("点击最终“确定”后，铺货弹窗未关闭且没有成功提示")


async def prepare_taobao_publish_dialog(
    page: Any,
    selected_shops: Sequence[str],
    timeout_seconds: int,
    logger: logging.Logger,
    *,
    platform_name: str = "淘宝",
) -> Tuple[Any, Dict[str, Any]]:
    """在铺货弹窗中选择指定平台，并且只勾选指定店铺。

    这一步不点击最终“确定”，可用于真实页面的无铺货预览。
    """
    expected = tuple(dict.fromkeys(shop.strip() for shop in selected_shops if shop.strip()))
    if not expected:
        raise AutomationError(f"没有指定要铺货的{platform_name}店铺")

    def canonical_shop_name(name: str) -> str:
        normalized = re.sub(r"\s+", " ", str(name).strip())
        if platform_name == "小红书":
            return XHS_PUBLISH_SHOP_ALIASES.get(normalized, normalized)
        return normalized

    expected_canonical = tuple(canonical_shop_name(shop) for shop in expected)

    dialog = page.locator('[role="dialog"]:visible').filter(has_text="铺货到店铺").first
    await dialog.wait_for(state="visible", timeout=timeout_seconds * 1000)
    platform = await first_visible(dialog.get_by_text(platform_name, exact=True))
    if platform is None:
        raise AutomationError(f"铺货弹窗中找不到{platform_name}平台")
    await platform.click()
    logger.info("铺货弹窗已选择%s平台", platform_name)

    def is_non_shop_label(text: str) -> bool:
        normalized = re.sub(r"\s+", "", text)
        return normalized in {"全选", "取消全选", "全部", "反选"} or "AI裂变规则" in normalized

    async def visible_shop_controls() -> List[Tuple[str, Any, Any]]:
        controls: List[Tuple[str, Any, Any]] = []
        candidates = dialog.locator(".el-checkbox:visible")
        for index in range(await candidates.count()):
            control = candidates.nth(index)
            text = re.sub(r"\s+", " ", (await control.inner_text()).strip())
            if not text or is_non_shop_label(text):
                continue
            checkbox = control.locator('input[type="checkbox"]').first
            if not await checkbox.count():
                continue
            controls.append((text, control, checkbox))
        return controls

    def shop_name_from_label(label: str) -> str:
        # “已铺货”是店铺后的状态标签，真实 DOM 有时与店名直接相连、
        # 有时由空白分隔；先剥离状态再做精确店名匹配，避免前缀误选。
        return re.sub(r"\s*已铺货\s*$", "", label).strip()

    controls: List[Tuple[str, Any, Any]] = []
    deadline = time.monotonic() + timeout_seconds
    last_shop_signature: Optional[Tuple[str, ...]] = None
    shop_list_stable_since: Optional[float] = None
    while time.monotonic() < deadline:
        controls = await visible_shop_controls()
        loaded_shop_names = {
            canonical_shop_name(shop_name_from_label(label))
            for label, _, _ in controls
        }
        if all(shop in loaded_shop_names for shop in expected_canonical):
            break
        if controls:
            signature = tuple(sorted(loaded_shop_names))
            now = time.monotonic()
            if signature != last_shop_signature:
                last_shop_signature = signature
                shop_list_stable_since = now
            elif (
                shop_list_stable_since is not None
                and now - shop_list_stable_since >= min(3.0, timeout_seconds)
            ):
                missing = [
                    shop
                    for shop, canonical in zip(expected, expected_canonical)
                    if canonical not in loaded_shop_names
                ]
                actual = [shop_name_from_label(label) for label, _, _ in controls]
                raise AutomationError(
                    f"{platform_name}平台店铺列表已加载，但找不到目标店铺："
                    + "、".join(missing)
                    + "；页面实际店铺："
                    + "、".join(actual)
                )
        await asyncio.sleep(0.25)
    else:
        actual = [shop_name_from_label(label) for label, _, _ in controls]
        raise AutomationError(
            f"{platform_name}平台店铺列表未加载或找不到："
            + "、".join(expected)
            + ("；页面实际店铺：" + "、".join(actual) if actual else "")
        )

    def matching_expected(label: str) -> Optional[str]:
        normalized = canonical_shop_name(shop_name_from_label(label))
        matches = [
            shop for shop, canonical in zip(expected, expected_canonical)
            if normalized == canonical
        ]
        if len(matches) > 1:
            raise AutomationError(f"{platform_name}店铺文字匹配不唯一：{label}")
        return matches[0] if matches else None

    found_expected: List[str] = []
    already_published: List[str] = []
    unavailable_shops: List[str] = []
    available_shops: List[str] = []
    for label, control, checkbox in controls:
        target = matching_expected(label)
        display_name = shop_name_from_label(label)
        available_shops.append(display_name)
        disabled = await checkbox.is_disabled()
        published = "已铺货" in label
        if target:
            found_expected.append(target)
            if disabled and not published:
                unavailable_shops.append(target)
                continue
            if published:
                # Keep the page's current spelling in the report and submit
                # payload; aliases are only for matching old cached markup.
                already_published.append(display_name)
                continue
        if disabled or published:
            continue
        should_check = target is not None
        if await checkbox.is_checked() != should_check:
            await control.scroll_into_view_if_needed()
            await control.click()
            if await checkbox.is_checked() != should_check:
                raise AutomationError(f"店铺“{display_name}”复选框状态切换失败")

    missing = [shop for shop in expected if shop not in found_expected]
    if missing:
        raise AutomationError(
            f"铺货弹窗缺少{platform_name}店铺：" + "、".join(missing)
        )
    if unavailable_shops:
        raise AutomationError(
            f"{platform_name}目标店铺当前不可选（且未标记已铺货）："
            + "、".join(unavailable_shops)
        )

    actual: List[str] = []
    for label, _, checkbox in await visible_shop_controls():
        target = matching_expected(label)
        if await checkbox.is_disabled() or "已铺货" in label:
            continue
        if await checkbox.is_checked():
            actual.append(shop_name_from_label(label))
    already_published_canonical = {
        canonical_shop_name(shop) for shop in already_published
    }
    pending_expected = [
        shop for shop, canonical in zip(expected, expected_canonical)
        if canonical not in already_published_canonical
    ]
    if {
        canonical_shop_name(shop) for shop in actual
    } != {canonical_shop_name(shop) for shop in pending_expected}:
        raise AutomationError(
            f"{platform_name}待铺货店铺复核失败：期望 " + "、".join(pending_expected)
            + "；实际 " + "、".join(actual)
        )

    report = {
        "platform": platform_name,
        "requested_shops": list(expected),
        "selected_shops": actual,
        "already_published": already_published,
        "available_shops": list(dict.fromkeys(available_shops)),
        "submitted": False,
    }
    logger.info(
        "%s铺货店铺已精确复核：待铺货=%s；已铺货=%s；未点击最终确定",
        platform_name,
        "、".join(actual) or "无",
        "、".join(already_published) or "无",
    )
    return dialog, report


async def submit_taobao_publish_dialog(
    page: Any,
    dialog: Any,
    selection: Mapping[str, Any],
    timeout_seconds: int,
    logger: logging.Logger,
    *,
    platform_name: str = "淘宝",
) -> Dict[str, Any]:
    """提交已经精确复核的平台铺货弹窗。"""
    selected = tuple(str(shop) for shop in selection.get("selected_shops") or ())
    already_published = list(selection.get("already_published") or ())
    if not selected:
        if already_published:
            await close_publish_preview(dialog)
            logger.info(
                "指定的%s店铺均已铺货，已安全关闭弹窗并跳过重复提交：%s",
                platform_name,
                "、".join(already_published),
            )
            return {
                **dict(selection),
                "result": 1,
                "confirmed_by": "already_published",
                "submitted": False,
                "submitted_shops": [],
            }
        raise AutomationError(f"{platform_name}铺货弹窗没有已勾选的目标店铺")

    final_button = dialog.get_by_role("button", name=re.compile(r"^\s*确\s*定\s*$"))
    if await final_button.count() != 1:
        raise AutomationError(
            f"{platform_name}铺货弹窗中找不到唯一的“确定”按钮"
        )

    submit_responses: List[Any] = []

    def on_submit_response(response: Any) -> None:
        request = response.request
        if request.method == "POST" and any(
            marker in response.url.casefold()
            for marker in ("/item/base/edit.json", "publish", "distribute", "supply")
        ):
            submit_responses.append(response)

    page.on("response", on_submit_response)
    await final_button.click()
    logger.info("已点击%s铺货弹窗最终“确定”", platform_name)

    success_pattern = re.compile("铺货成功|提交成功|已提交铺货|任务已创建")
    success = page.locator(".el-message--success:visible").filter(has_text=success_pattern)
    continue_button = None
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        continue_button = await first_visible(
            page.get_by_role("button", name=re.compile("继续铺货"))
        )
        if continue_button is not None or await visible(success.first):
            break
        page_errors = page.locator(".el-message--error:visible")
        if await visible(page_errors.first):
            raise AutomationError(
                f"{platform_name}铺货提交失败："
                + (await page_errors.first.inner_text()).strip()
            )
        progress = await dismiss_publish_progress_dialog(page, logger, settle_seconds=2)
        if progress['dismissed'] or not await dialog.is_visible():
            # 最终“确定”后任务由后台继续执行；
            # 店铺选择弹窗关闭即表示本次提交已被接受。
            if not progress['dismissed']:
                await asyncio.sleep(2)
            for response in submit_responses:
                try:
                    payload = await response.json()
                except Exception:
                    continue
                if "result" in payload and int(payload.get("result", 0) or 0) != 1:
                    raise AutomationError(
                        f"{platform_name}铺货接口返回失败："
                        + str(payload.get("message") or payload.get("errmsg") or payload)
                    )
            return {
                **dict(selection),
                "result": 1,
                "confirmed_by": "background_task_submitted",
                "submitted": True,
                "submitted_shops": list(selected),
                "responses": [response.url for response in submit_responses],
            }
        await asyncio.sleep(0.25)

    if continue_button is not None:
        await continue_button.click()
        logger.info("已点击%s铺货确认框“继续铺货”", platform_name)
    elif not await visible(success.first):
        raise AutomationError(
            f"点击{platform_name}铺货“确定”后未出现“继续铺货”或成功提示"
        )

    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        for response in submit_responses:
            try:
                payload = await response.json()
            except Exception:
                continue
            if "result" in payload and int(payload.get("result", 0) or 0) != 1:
                raise AutomationError(
                    f"{platform_name}铺货接口返回失败："
                    + str(payload.get("message") or payload.get("errmsg") or payload)
                )
        if await visible(success.first):
            await asyncio.sleep(2)
            return {
                **dict(selection),
                "result": 1,
                "confirmed_by": "toast",
                "submitted": True,
                "submitted_shops": list(selected),
                "responses": [response.url for response in submit_responses],
            }
        progress = await dismiss_publish_progress_dialog(page, logger, settle_seconds=2)
        if progress['dismissed'] or (not await dialog.is_visible() and (
            continue_button is None or not await continue_button.is_visible()
        )):
            if not progress['dismissed']:
                await asyncio.sleep(2)
            return {
                **dict(selection),
                "result": 1,
                "confirmed_by": "background_task_submitted",
                "submitted": True,
                "submitted_shops": list(selected),
                "responses": [response.url for response in submit_responses],
            }
        page_errors = page.locator(".el-message--error:visible")
        if await visible(page_errors.first):
            raise AutomationError(
                f"{platform_name}铺货提交失败："
                + (await page_errors.first.inner_text()).strip()
            )
        await asyncio.sleep(0.25)
    raise AutomationError(f"{platform_name}铺货提交后未收到成功确认")


async def close_publish_preview(dialog: Any) -> None:
    """关闭铺货预览弹窗，不触发最终提交。"""
    cancel = await first_visible(
        dialog.get_by_role("button", name=re.compile(r"^\s*(?:取消|关闭)\s*$"))
    )
    if cancel is None:
        cancel = await first_visible(dialog.locator(".el-dialog__headerbtn"))
    if cancel is not None:
        await cancel.click()
        await dialog.wait_for(state="hidden", timeout=10_000)
        return

    # 真实铺货窗口的右上角是无文字图标，且按钮可能不在
    # role=dialog 的内部 DOM 中。Escape 只能关闭弹窗，不会触发铺货。
    await dialog.press("Escape")
    try:
        await dialog.wait_for(state="hidden", timeout=5_000)
    except Exception as exc:
        raise AutomationError("铺货预览弹窗无法通过取消、关闭或 Escape 安全退出") from exc


async def dismiss_publish_progress_dialog(
    page: Any,
    logger: logging.Logger,
    *,
    appearance_timeout_seconds: float = 0.0,
    settle_seconds: float = 2.0,
) -> Dict[str, Any]:
    """页面稳定后收起后台铺货进度弹窗，不等待后台任务完成。"""
    dialogs = page.locator(
        '[role="dialog"]:visible, .el-dialog:visible, .el-message-box:visible'
    ).filter(has_text=re.compile(r"铺货中|铺货进度"))
    deadline = time.monotonic() + max(0.0, appearance_timeout_seconds)
    dialog = None
    while True:
        dialog = await first_visible(dialogs)
        if dialog is not None:
            break
        if time.monotonic() >= deadline:
            return {"found": False, "dismissed": False, "action": None}
        await asyncio.sleep(0.1)

    # 弹窗刚出现时仍可能在替换进度 DOM。只给前端一个短稳定窗口，
    # 随后立即收起；后台铺货会继续，这里不轮询它的最终进度。
    if settle_seconds > 0:
        await asyncio.sleep(settle_seconds)

    dialog_text = re.sub(r"\s+", " ", (await dialog.inner_text()).strip())
    percent_match = re.search(r"(\d{1,3})\s*%", dialog_text)
    progress_percent = int(percent_match.group(1)) if percent_match else None

    action = None
    collapse = await first_visible(
        dialog.get_by_role("button", name=re.compile(r"^\s*收\s*起\s*$"))
    )
    if collapse is None:
        collapse = await first_visible(dialog.get_by_text("收起", exact=True))
    if collapse is not None:
        await collapse.click(force=True)
        action = "collapse"
        try:
            await dialog.wait_for(state="hidden", timeout=3_000)
        except Exception:
            pass

    if await visible(dialog):
        close_button = await first_visible(
            dialog.locator(
                ".el-dialog__headerbtn, .el-message-box__headerbtn, "
                'button[aria-label="Close"], button[aria-label="close"]'
            )
        )
        if close_button is not None:
            await close_button.click(force=True)
            action = "close"
            try:
                await dialog.wait_for(state="hidden", timeout=3_000)
            except Exception:
                pass

    if await visible(dialog):
        await dialog.press("Escape")
        action = "escape"
        try:
            await dialog.wait_for(state="hidden", timeout=3_000)
        except Exception as exc:
            raise AutomationError(
                "铺货进度弹窗无法通过“收起”、关闭或 Escape 安全退出"
            ) from exc

    logger.info(
        "铺货进度弹窗已收起（方式=%s，进度=%s）",
        action,
        f"{progress_percent}%" if progress_percent is not None else "未显示",
    )
    return {
        "found": True,
        "dismissed": True,
        "action": action,
        "progress_percent": progress_percent,
    }


def requires_base_save_before_platform(
    platform: str,
) -> bool:
    """基础资料是独立阶段；平台流程不得再次改写或保存基础资料。"""
    return platform == "base"


def resolve_commerce_publish_target(
    args: argparse.Namespace,
) -> Optional[Tuple[str, Sequence[str]]]:
    """返回当前平台应精确选择的平台名称与店铺；只保存模式不铺货。"""
    if args.save_only:
        return None
    if args.platform == "taobao" and not getattr(
        args, "allow_taobao_publish_once", False
    ):
        return None
    target = COMMERCE_PUBLISH_TARGETS.get(args.platform)
    if target is None:
        return None
    if getattr(args, "all_platform_one_shop_test", False):
        platform_name, shops = target
        return platform_name, tuple(shops[:1])
    return target


def resolve_platform_save_action(
    args: argparse.Namespace,
    *,
    publish_mode: bool,
) -> Tuple[str, bool]:
    """Resolve the exact footer action, honoring the platform publish gate."""
    platform = get_platform_spec(args.platform)
    commerce_target = resolve_commerce_publish_target(args)
    should_publish = platform.publish_allowed and (
        (publish_mode and not args.save_only) or commerce_target is not None
    )
    return ("保存并铺货到平台" if should_publish else "保存", should_publish)


@asynccontextmanager
async def playwright_for_browser_run(
    async_playwright_factory: Any,
    shared_session: Optional[Dict[str, Any]],
):
    """单平台正常启停；全平台则跨阶段保留同一 Playwright 会话。"""
    if shared_session is None:
        async with async_playwright_factory() as playwright:
            yield playwright
        return

    existing = shared_session.get("playwright")
    if existing is not None:
        yield existing
        return

    manager = async_playwright_factory()
    playwright = await manager.start()
    shared_session["playwright"] = playwright
    try:
        yield playwright
    except Exception:
        # 由全平台调度器的 finally 统一释放，确保启动或登录中途失败也不泄漏。
        raise


async def close_shared_browser_session(shared_session: Dict[str, Any]) -> None:
    """释放全平台共用浏览器资源，但不关闭接管的用户/CDP Chrome。"""
    context = shared_session.get("context")
    playwright = shared_session.get("playwright")
    try:
        if context is not None and shared_session.get("owns_context", False):
            try:
                await asyncio.wait_for(context.close(), timeout=5.0)
            except Exception as exc:
                logging.getLogger(__name__).warning("浏览器上下文清理未完成，继续释放驱动：%s", type(exc).__name__)
    finally:
        try:
            if playwright is not None:
                try:
                    await asyncio.wait_for(playwright.stop(), timeout=5.0)
                except Exception as exc:
                    logging.getLogger(__name__).warning("浏览器驱动清理未完成，不阻塞审核恢复：%s", type(exc).__name__)
        finally:
            shared_session.clear()


async def ensure_shared_product_editor(
    page: Any,
    shared_session: Dict[str, Any],
    product: ProductData,
    *,
    timeout_seconds: int,
    logger: logging.Logger,
) -> Tuple[Any, Any, bool]:
    """Return the live shared drawer, reopening it after publish closes it."""
    await dismiss_publish_progress_dialog(
        page,
        logger,
        appearance_timeout_seconds=0.5,
    )
    drawer = shared_session["drawer"]
    try:
        drawer_visible = await drawer.is_visible()
    except Exception:
        try:
            await drawer.wait_for(state="visible", timeout=1_000)
            drawer_visible = True
        except Exception:
            drawer_visible = False
    if drawer_visible:
        return drawer, shared_session.get("record"), False

    logger.info(
        "上一平台完成后商品编辑抽屉已关闭，按款式编码重新打开；已完成平台不重跑"
    )
    record = await api_find_product(page, product.style_code, logger)
    drawer = await open_product_editor(
        page,
        product.style_code,
        logger,
        timeout_seconds=timeout_seconds,
    )
    shared_session.update({"drawer": drawer, "record": record})
    logger.info("商品编辑抽屉已重新打开，继续当前平台")
    return drawer, record, True


async def product_editor_for_saved_readback(
    page: Any,
    drawer: Any,
    shared_session: Optional[Dict[str, Any]],
    product: ProductData,
    *,
    timeout_seconds: int,
    logger: logging.Logger,
    platform_label: str,
    force_reload: bool = False,
) -> Tuple[Any, bool]:
    """Return a live editor for save/publish readback.

    Publishing commonly closes the drawer.  Shared all-platform runs must
    reopen it before platform-specific validation instead of waiting on a
    stale hidden locator.
    """
    if shared_session is not None and not force_reload:
        persisted_drawer, _, reopened = await ensure_shared_product_editor(
            page,
            shared_session,
            product,
            timeout_seconds=timeout_seconds,
            logger=logger,
        )
        logger.info(
            "保存后正在%s商品编辑页复核%s；不重跑已完成平台",
            "重开" if reopened else "当前",
            platform_label,
        )
        return persisted_drawer, reopened

    logger.info("保存成功，正在重新打开商品复核%s", platform_label)
    await page.reload(
        wait_until="domcontentloaded",
        timeout=timeout_seconds * 1000,
    )
    persisted_drawer = await open_product_editor(
        page,
        product.style_code,
        logger,
        timeout_seconds=timeout_seconds,
    )
    if shared_session is not None:
        shared_session['drawer'] = persisted_drawer
    return persisted_drawer, True


async def run_browser_automation(
    args: argparse.Namespace,
    product: ProductData,
    artifact_dir: Path,
    logger: logging.Logger,
    redactor: Optional[SensitiveLogRedactor] = None,
    shared_session: Optional[Dict[str, Any]] = None,
    attribute_runtime: Optional[AttributeRuntime] = None,
) -> None:
    try:
        from playwright.async_api import async_playwright
    except ImportError as exc:
        raise AutomationError("缺少 Playwright，请先运行：python3 -m pip install -r requirements.txt") from exc

    inspect_only = bool(getattr(args, "inspect_only", False))
    douyin_requested = args.platform in {"all", "douyin"}
    taobao_requested = args.platform == "taobao"
    tmall_requested = args.platform == "tmall"
    pdd_requested = args.platform == "pdd"
    wxsph_requested = args.platform == "wxsph"
    xhs_requested = args.platform == "xhs"
    youzan_requested = args.platform == "youzan"
    jd_requested = args.platform == "jd"
    if attribute_runtime is not None:
        attribute_runtime.begin_review_collection(args.platform)
    if args.platform == "douyin" and product.douyin_fields is None:
        raise AutomationError("已选择抖音流程，但 Excel/产品目录中没有抖音资料")
    if taobao_requested and product.taobao_fields is None:
        raise AutomationError("已选择淘宝流程，但 Excel 中没有可用于淘宝匹配的资料")
    if tmall_requested and not inspect_only and product.tmall_fields is None:
        raise AutomationError("已选择天猫流程，但 Excel 中没有可用于天猫匹配的资料")
    if pdd_requested and not inspect_only and product.pdd_fields is None:
        raise AutomationError("已选择拼多多流程，但 Excel 中没有可用于拼多多匹配的资料")
    if wxsph_requested and not inspect_only and product.wxsph_fields is None:
        raise AutomationError("已选择微信小店流程，但 Excel 中没有可用的资料")
    if xhs_requested and not inspect_only and product.xhs_fields is None:
        raise AutomationError("已选择小红书流程，但 Excel 中没有可用于小红书匹配的资料")
    if xhs_requested and not inspect_only and not product.xhs_fields.category_path:
        raise AutomationError("已选择小红书流程，但 Excel 中缺少“商品分类”层级")
    if youzan_requested and not inspect_only and product.youzan_fields is None:
        raise AutomationError("已选择有赞流程，但 Excel 中没有可用于有赞匹配的资料")
    if jd_requested and not inspect_only and product.jd_fields is None:
        raise AutomationError("已选择京东流程，但 Excel 中没有可用于京东匹配的资料")
    if youzan_requested and not inspect_only and not product.youzan_fields.category_path:
        raise AutomationError("已选择有赞流程，但 Excel 中缺少“商品分类”层级")
    if (
        youzan_requested
        and not inspect_only
        and product.youzan_fields.garment_kind == "unknown"
    ):
        raise AutomationError(
            "有赞运费模板只能按裤子或外套判定；请检查 Excel“商品分类”"
        )
    tmall_assets: Optional[TmallAssets] = (
        read_tmall_assets(product.product_dir) if tmall_requested and not inspect_only else None
    )
    recommendations = (
        recognize_product_recommendations(product, artifact_dir)
        if (douyin_requested or taobao_requested)
        and product.douyin_fields is not None
        else ()
    )
    if recommendations:
        logger.info("本地尺码识别完成：%s", " / ".join(item.size for item in recommendations))
    taobao_category_mode = "recommended"
    taobao_garment_kind_value = ""
    taobao_size_lengths: tuple[SizeLength, ...] = ()
    if taobao_requested:
        taobao_category_mode = resolve_taobao_category_mode(
            product, getattr(args, "taobao_category_mode", "auto")
        )
        taobao_garment_kind_value, taobao_size_lengths = recognize_taobao_size_lengths(
            product,
            taobao_category_mode,
            artifact_dir,
        )
        logger.info(
            "淘宝尺码信息表识别完成：%s | %s",
            "裤长" if taobao_garment_kind_value == "pants" else "衣长",
            " / ".join(item.size for item in taobao_size_lengths),
        )

    async with playwright_for_browser_run(async_playwright, shared_session) as playwright:
        reuse_editor = bool(shared_session and shared_session.get("drawer") is not None)
        if reuse_editor:
            remote_browser = shared_session.get("remote_browser")
            attached_existing_profile = bool(
                shared_session.get("attached_existing_profile", False)
            )
            auth_state_path = shared_session.get("auth_state_path")
            restored_scm_verified = True
            context = shared_session["context"]
            page = shared_session["page"]
        else:
            remote_browser = None
            attached_existing_profile = False
            auth_state_path: Optional[Path] = None
            restored_scm_verified = False
            if args.cdp_url:
                remote_browser = await playwright.chromium.connect_over_cdp(args.cdp_url)
                if not remote_browser.contexts:
                    raise AutomationError(f"CDP 浏览器没有可用上下文：{args.cdp_url}")
                context = remote_browser.contexts[0]
                page = next((p for p in context.pages if is_scm_url(p.url)), None)
                if page is None:
                    page = await context.new_page()
            else:
                profile_dir = Path(args.user_data_dir).expanduser().resolve()
                profile_dir.mkdir(parents=True, exist_ok=True)
                auth_state_path = Path(args.auth_state).expanduser().resolve()
                await close_stale_profile_browsers(profile_dir, logger)
                launch_options: Dict[str, Any] = {
                    "user_data_dir": str(profile_dir),
                    "channel": "chrome",
                    "headless": args.headless,
                    "no_viewport": True,
                    "args": ["--start-maximized"],
                }
                if inspect_only:
                    launch_options["service_workers"] = "block"
                context = await playwright.chromium.launch_persistent_context(
                    **launch_options,
                )

                # Chrome 异常退出后会恢复上次的旧标签。保留同一专用配置的
                # 登录状态，但本次任务始终从一个新标签开始。
                stale_pages = tuple(context.pages)
                fresh_page = await context.new_page()
                for stale_page in stale_pages:
                    try:
                        await stale_page.close()
                    except Exception:
                        pass

                await restore_auth_state(context, auth_state_path, logger)
                restored_scm_verified = saved_scm_state_is_verified(auth_state_path)
                page = fresh_page
                if page is None:
                    page = next(
                        (
                            p
                            for p in context.pages
                            if urlsplit(p.url).hostname in ERP_HOSTS
                        ),
                        None,
                    )
                if page is None:
                    page = context.pages[-1] if context.pages else await context.new_page()

            if shared_session is not None:
                shared_session.update(
                    {
                        "context": context,
                        "page": page,
                        "remote_browser": remote_browser,
                        "attached_existing_profile": attached_existing_profile,
                        "auth_state_path": auth_state_path,
                        "owns_context": not args.cdp_url
                        and not attached_existing_profile,
                    }
                )

        page.set_default_timeout(args.timeout * 1000)
        wxsph_capture: Optional[WxsphFormListing] = None
        if shared_session is not None and hasattr(page, "locator"):
            wxsph_capture = shared_session.get("wxsph_api_capture")
            if wxsph_capture is None:
                wxsph_capture = WxsphFormListing(
                    page,
                    page.locator("body"),
                    logger,
                    attribute_runtime=attribute_runtime,
                )
                shared_session["wxsph_api_capture"] = wxsph_capture
        elif wxsph_requested:
            # Single-platform runs also need the listener before
            # open_product_editor(), because the drawer can fetch the schema
            # eagerly while its base tab is being rendered.
            wxsph_capture = WxsphFormListing(
                page,
                page.locator("body"),
                logger,
                attribute_runtime=attribute_runtime,
            )
        try:
            if reuse_editor:
                session_style_code = str(
                    shared_session.get("style_code") or ""
                ).strip()
                if session_style_code != product.style_code:
                    raise AutomationError(
                        "共享编辑页款式编码不匹配："
                        f"期望 {product.style_code}，实际 {session_style_code or '未记录'}"
                    )
                drawer, record, reopened = await ensure_shared_product_editor(
                    page,
                    shared_session,
                    product,
                    timeout_seconds=args.timeout,
                    logger=logger,
                )
                if reopened:
                    logger.info("重新打开商品编辑页，继续切换到 %s 平台", args.platform)
                else:
                    logger.info("复用当前商品编辑页，直接切换到 %s 平台", args.platform)
            else:
                reused_scm = False
                if restored_scm_verified or (args.cdp_url and is_scm_url(page.url)):
                    try:
                        page = await try_reuse_verified_scm_session(
                            page,
                            product.style_code,
                            args.timeout,
                            logger,
                            auth_state_path=auth_state_path,
                        )
                        reused_scm = True
                    except Exception as exc:
                        logger.info("已保存的快麦通会话不可复用，回退 ERP 单点登录：%s", exc)

                if not reused_scm:
                    page = await enter_kuaimai_from_erp(
                        page,
                        args.login_timeout,
                        args.headless,
                        logger,
                        operation_timeout_seconds=args.timeout,
                        auth_state_path=auth_state_path,
                        style_code=product.style_code,
                    )
                page.set_default_timeout(args.timeout * 1000)
                if getattr(args, "create_product", False):
                    await run_new_product(page, args, product, artifact_dir, logger)
                    return
                if not inspect_only:
                    await dismiss_publish_progress_dialog(
                        page,
                        logger,
                        appearance_timeout_seconds=0.5,
                    )
                if inspect_only:
                    record = await api_find_product(
                        page,
                        product.style_code,
                        logger,
                        redactor=redactor,
                    )
                else:
                    record = await api_find_product(page, product.style_code, logger)
                if record is None:
                    logger.warning("继续使用 DOM 查询目标商品")

                drawer = await open_product_editor(
                    page,
                    product.style_code,
                    logger,
                    timeout_seconds=args.timeout,
                )
                if shared_session is not None:
                    shared_session.update(
                        {"page": page, "drawer": drawer, "record": record}
                    )

                style_item = await form_item(
                    drawer, "款式编码", timeout_seconds=args.timeout
                )
                style_value = await style_item.locator("input").first.input_value()
                if style_value.strip() != product.style_code:
                    raise AutomationError(
                        f"编辑页款式编码不匹配：期望 {product.style_code}，实际 {style_value}"
                    )
                if shared_session is not None:
                    shared_session["style_code"] = style_value.strip()

            if inspect_only:
                if redactor is None:
                    raise AutomationError("inspect-only 未配置敏感日志过滤器")
                await run_platform_inspection(
                    page,
                    drawer,
                    args,
                    product,
                    record,
                    artifact_dir,
                    logger,
                    redactor,
                )
                return

            # 独立验证水洗标上传链路：只打开抖音资料页并上传图片，不填写
            # 其它字段、不保存、不铺货。正常抖音流程复用同一个方法，确保
            # 这里验证通过后再重跑平台时不会换一套上传逻辑。
            if getattr(args, "wash_label_upload_test", False):
                if args.platform != "douyin":
                    raise AutomationError(
                        "水洗标单项测试只允许 --platform douyin"
                    )
                if product.douyin_assets is None:
                    raise AutomationError("当前商品没有可读取的水洗标图片")
                douyin = DouyinListing(
                    page,
                    drawer,
                    logger,
                    artifact_dir,
                    attribute_runtime=attribute_runtime,
                    category_hints=product.category_hints,
                )
                await douyin.open()
                # 抖音只有在先应用商品类目后才会渲染水洗标上传组件。
                # 单项测试也执行这一步，但不填写类目属性。
                category = await douyin.apply_first_recommended_category()
                logger.info("水洗标单项测试已应用抖音商品类目：%s", category)
                upload_result = await douyin.upload_wash_label_images_only(
                    product.douyin_assets.wash_label_images
                )
                # 给页面侧上传回调和缩略图渲染留出稳定时间，再截图作为
                # 单项测试凭证；此处不会点击保存或铺货。
                await asyncio.sleep(2)
                await safe_screenshot(
                    page,
                    artifact_dir / "douyin-wash-label-upload-test.png",
                )
                logger.info(
                    "水洗标图片单项上传测试完成：%s；未填写、未保存、未铺货",
                    upload_result,
                )
                return

            publish_mode = douyin_requested and product.douyin_fields is not None
            base_save_result = None
            if requires_base_save_before_platform(args.platform):
                title_item = await form_item(
                    drawer, "商品名称", timeout_seconds=args.timeout
                )
                title_input = await fill_input(title_item, product.title, "商品名称")
                logger.info("已填写商品名称")

                await sync_image_group(
                    page,
                    await form_item(drawer, "商品主图", timeout_seconds=args.timeout),
                    product.main_images,
                    "1:1 主图",
                    args.upload_timeout,
                )
                await sync_image_group(
                    page,
                    await form_item(drawer, "3:4主图", timeout_seconds=args.timeout),
                    product.main_images_34,
                    "3:4 主图",
                    args.upload_timeout,
                )
                await sync_image_group(
                    page,
                    await form_item(drawer, "商品详情图", timeout_seconds=args.timeout),
                    product.detail_images,
                    "商品详情图",
                    args.upload_timeout,
                )
                color_report: Optional[Dict[str, Any]] = None
                size_report: Optional[Dict[str, Any]] = None
                if product.colors:
                    color_report = await sync_base_color_spec_values(
                        drawer,
                        product.colors,
                    )
                    logger.info(
                        "基础资料颜色规格已按 Excel 顺序回读：%s",
                        " / ".join(color_report["after"]),
                    )
                else:
                    logger.info("Excel 未提供颜色字段，基础资料颜色规格保持页面原值")
                if product.derived_size_names:
                    size_report = await sync_base_specification_values(
                        drawer,
                        "尺码",
                        product.derived_size_names,
                    )
                    logger.info(
                        "基础资料尺码规格已按当前商品顺序回读：%s",
                        " / ".join(size_report["after"]),
                    )
                else:
                    logger.info("当前商品未提取到尺码字段，基础资料尺码规格保持页面原值")
                sku_count = await replace_sku_images(
                    page,
                    await form_item(drawer, "商品规格", timeout_seconds=args.timeout),
                    product.sku_images,
                    args.upload_timeout,
                )
                logger.info(
                    "SKU 图检查完成：新上传 %s 张，页面已有图片均已保留",
                    sku_count,
                )

                price_input = await set_base_price(drawer, product.base_price)
                logger.info("已将基本售价批量设置为 %s", product.base_price)

                # 保存前做一次关键值复核。
                if (await title_input.input_value()).strip() != product.title:
                    raise AutomationError("保存前复核失败：商品名称发生变化")
                if Decimal(await price_input.input_value()) != Decimal(product.base_price):
                    raise AutomationError("保存前复核失败：基本售价发生变化")
                base_errors = await collect_visible_errors(drawer)
                if base_errors:
                    raise AutomationError(
                        "基础资料存在页面校验错误：" + "；".join(base_errors)
                    )

                logger.info("基础资料复核通过，保存并等待当前页面加载完成")
                base_save_result = await click_save_and_confirm(
                    page,
                    drawer,
                    args.sync_erp,
                    args.timeout,
                    logger,
                    button_text="保存",
                )
                if color_report is not None:
                    base_save_result["color_specification"] = color_report
                if size_report is not None:
                    base_save_result["size_specification"] = size_report
                (artifact_dir / "base-save-result.json").write_text(
                    json.dumps(base_save_result, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )

                await wait_for_base_form_ready_after_save(
                    drawer,
                    product.style_code,
                    product.title,
                    args.timeout,
                )
                if product.colors:
                    persisted_colors = await read_base_specification_values(
                        drawer,
                        "颜色",
                    )
                    if persisted_colors != product.colors:
                        raise AutomationError(
                            "基础资料保存后颜色规格回读不一致："
                            f"页面 {persisted_colors}，Excel {product.colors}"
                        )
                    base_save_result["color_specification"]["persisted"] = persisted_colors
                    (artifact_dir / "base-save-result.json").write_text(
                        json.dumps(base_save_result, ensure_ascii=False, indent=2),
                        encoding="utf-8",
                    )
                if product.derived_size_names:
                    persisted_sizes = await read_base_specification_values(
                        drawer,
                        "尺码",
                    )
                    if persisted_sizes != product.derived_size_names:
                        raise AutomationError(
                            "基础资料保存后尺码规格回读不一致："
                            f"页面 {persisted_sizes}，目标 {product.derived_size_names}"
                        )
                    base_save_result["size_specification"]["persisted"] = (
                        persisted_sizes
                    )
                    (artifact_dir / "base-save-result.json").write_text(
                        json.dumps(base_save_result, ensure_ascii=False, indent=2),
                        encoding="utf-8",
                    )
                logger.info("基础资料已保存，当前编辑页加载完成")
            else:
                logger.info(
                    "%s独立流程：不填写、不保存基础资料，直接切换平台资料",
                    args.platform,
                )

            if publish_mode:
                assert product.douyin_fields is not None
                assert product.douyin_assets is not None
                douyin = DouyinListing(
                    page,
                    drawer,
                    logger,
                    artifact_dir,
                    attribute_runtime=attribute_runtime,
                    category_hints=product.category_hints,
                )
                await douyin.open()
                title_prediction = await douyin.prepare_product_title_and_predictions(
                    product.title,
                    force_refresh=False,
                )
                category_fields = dict(
                    await douyin.apply_category_and_fields(product.douyin_fields)
                )
                category_fields.update(title_prediction)
                materials = await douyin.apply_materials(
                    product.douyin_fields.materials,
                    product.douyin_assets.wash_label_images,
                    product.douyin_fields.materials_text,
                )
                size_rows = (
                    await douyin.fill_size_recommendations(recommendations)
                    if recommendations
                    else {
                        "status": "skipped",
                        "reason": "non_pants_product",
                        "garment_kind": product.garment_kind,
                    }
                )
                image_actions = await douyin.sync_douyin_images(
                    product.main_images,
                    product.main_images_34,
                    product.detail_images,
                    timeout_seconds=args.upload_timeout,
                )
                delivery = await douyin.apply_delivery_mode()
                sku_rows = await douyin.fill_sku_price_inventory(
                    product.douyin_fields.price,
                    product.douyin_fields.spot_stock,
                    product.douyin_fields.presale_stock,
                )
                freight = await douyin.apply_freight_templates(
                    product.douyin_fields.freight_aliases,
                    target_shops=DOUYIN_FREIGHT_TEMPLATE_SHOPS,
                    default_untargeted_template="包邮",
                )
                douyin_expected = {
                    **category_fields,
                    "materials": materials,
                    "sizes": size_rows,
                    "images": image_actions,
                    "delivery": delivery,
                    "sku": {
                        "rows": sku_rows,
                        "price": product.douyin_fields.price,
                        "spot_stock": product.douyin_fields.spot_stock,
                        "presale_stock": product.douyin_fields.presale_stock,
                    },
                    "freight": freight,
                }
                if (
                    attribute_runtime is not None
                    and attribute_runtime.has_deferred_reviews
                ):
                    douyin_report = {
                        **douyin_expected,
                        "validation_deferred": True,
                    }
                else:
                    douyin_report = await douyin.validate_douyin_form(
                        douyin_expected
                    )
                (artifact_dir / "douyin-before-publish.json").write_text(
                    json.dumps(douyin_report, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                await safe_screenshot(page, artifact_dir / "before-publish.png")
                logger.info("抖音资料填写与铺货前复核完成")
            elif taobao_requested:
                assert product.taobao_fields is not None
                taobao = TaobaoListing(
                    page,
                    drawer,
                    logger,
                    attribute_runtime=attribute_runtime,
                    category_hints=product.category_hints,
                )
                await taobao.open()
                if args.taobao_test_scope == "category-size":
                    category = await taobao.apply_category(taobao_category_mode)
                    size_chart = await taobao.fill_size_chart_lengths(
                        taobao_size_lengths,
                        garment_kind=taobao_garment_kind_value,
                    )
                    taobao_report = {
                        "category": category,
                        "test_scope": "category-size",
                        "extended_fields": {"size_chart": size_chart},
                    }
                    logger.info("淘宝类目与尺码表独立预览完成")
                elif args.taobao_test_scope == "attributes":
                    attribute_report = await taobao.apply_excel_attributes(
                        product.taobao_fields,
                        category_mode=taobao_category_mode,
                    )
                    taobao_report = {
                        **dict(attribute_report),
                        "test_scope": "attributes",
                        "extended_fields": {},
                    }
                    logger.info("淘宝类目属性独立预览完成")
                elif args.taobao_test_scope == "sku-batch":
                    category = await taobao.apply_category(taobao_category_mode)
                    sku_batch = await taobao.fill_sku_batch(
                        product.taobao_fields.fields
                    )
                    taobao_report = {
                        "category": category,
                        "test_scope": "sku-batch",
                        "extended_fields": {"sku_batch": sku_batch},
                    }
                    logger.info("淘宝 SKU 批量字段独立预览完成")
                elif args.taobao_test_scope == "payment-service":
                    category = await taobao.apply_category(taobao_category_mode)
                    payment_service = await taobao.apply_payment_and_service(
                        product.taobao_fields.fields
                    )
                    taobao_report = {
                        "category": category,
                        "test_scope": "payment-service",
                        "extended_fields": {"payment_service": payment_service},
                    }
                    logger.info("淘宝库存扣减与售后服务独立预览完成")
                else:
                    taobao_report = await taobao.apply_excel_attributes(
                        product.taobao_fields,
                        category_mode=taobao_category_mode,
                    )
                    taobao_report = dict(taobao_report)
                    taobao_report["extended_fields"] = await taobao.apply_extended_fields(
                        product.taobao_fields,
                        size_lengths=taobao_size_lengths,
                        garment_kind=taobao_garment_kind_value,
                        size_recommendations=recommendations,
                    )
                extended_fields = taobao_report["extended_fields"]
                taobao_validation_errors = []
                if args.taobao_test_scope == "full":
                    taobao_validation_errors.extend(
                        extended_fields.get("visible_validation_errors") or ()
                    )
                    specification_validation = extended_fields.get(
                        "specification_validation"
                    ) or {}
                    if not specification_validation.get("valid"):
                        taobao_validation_errors.append(
                            str(
                                specification_validation.get("result")
                                or "商品规格校验未通过"
                            )
                        )
                    sale_prop_validation = extended_fields.get("sale_prop_validation") or {}
                    if sale_prop_validation.get("message"):
                        taobao_validation_errors.append(
                            str(sale_prop_validation["message"])
                        )
                (artifact_dir / "taobao-before-save.json").write_text(
                    json.dumps(taobao_report, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                await safe_screenshot(page, artifact_dir / "taobao-before-save.png")
                if (
                    taobao_validation_errors
                    and not (
                        attribute_runtime is not None
                        and attribute_runtime.has_deferred_reviews
                    )
                ):
                    raise AutomationError(
                        "淘宝资料保存前校验失败："
                        + "；".join(dict.fromkeys(taobao_validation_errors))
                    )
                if args.taobao_test_scope == "full":
                    logger.info("淘宝类目、Excel 属性、尺码表与基础销售资料填写复核完成")
            elif pdd_requested:
                assert product.pdd_fields is not None
                pdd = PddFormListing(
                    page,
                    drawer,
                    logger,
                    attribute_runtime=attribute_runtime,
                    category_hints=product.category_hints,
                )
                await pdd.open()
                try:
                    pdd_report = await pdd.apply_excel_fields(product.pdd_fields)
                except PddFormListingError:
                    await safe_screenshot(page, artifact_dir / "pdd-error.png")
                    raise
                (artifact_dir / "pdd-before-save.json").write_text(
                    json.dumps(pdd_report, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                await safe_screenshot(page, artifact_dir / "pdd-before-save.png")
                next_action = (
                    "即将保存拼多多资料（不铺货）"
                    if args.save and args.save_only
                    else "即将保存并铺货到拼多多指定店铺"
                    if args.save
                    else "拼多多资料未保存、未铺货"
                )
                logger.info("拼多多资料填写与保存前复核完成；%s", next_action)
            elif wxsph_requested:
                assert product.wxsph_fields is not None
                wxsph = wxsph_capture or WxsphFormListing(
                    page, drawer, logger, attribute_runtime=attribute_runtime
                )
                if isinstance(record, dict):
                    wxsph.base_item_id = str(
                        record.get("baseItemId")
                        or record.get("base_item_id")
                        or record.get("id")
                        or ""
                    )
                wxsph.drawer = drawer
                wxsph.logger = logger
                wxsph.attribute_runtime = attribute_runtime
                await wxsph.open()
                try:
                    wxsph_report = await wxsph.apply_excel_fields(
                        product.wxsph_fields
                    )
                except WxsphFormListingError:
                    await safe_screenshot(page, artifact_dir / "wxsph-error.png")
                    raise
                (artifact_dir / "wxsph-before-save.json").write_text(
                    json.dumps(wxsph_report, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                attribute_anchor = await first_visible(
                    drawer.get_by_text("类目属性", exact=True)
                )
                if attribute_anchor is not None:
                    await attribute_anchor.scroll_into_view_if_needed()
                    await page.wait_for_timeout(250)
                    await safe_screenshot(
                        page, artifact_dir / "wxsph-attributes-before-save.png"
                    )
                batch_anchor = await first_visible(
                    drawer.get_by_role("button", name="批量设置", exact=True)
                )
                if batch_anchor is not None:
                    await batch_anchor.scroll_into_view_if_needed()
                    await page.wait_for_timeout(250)
                    await safe_screenshot(
                        page, artifact_dir / "wxsph-sku-before-save.png"
                    )
                weight_anchor = await first_visible(
                    drawer.get_by_text(re.compile(r"^\s*\*?\s*重量\s*[：:]?\s*$"))
                )
                if weight_anchor is not None:
                    await weight_anchor.scroll_into_view_if_needed()
                    await page.wait_for_timeout(250)
                    await safe_screenshot(
                        page, artifact_dir / "wxsph-weight-before-save.png"
                    )
                await safe_screenshot(page, artifact_dir / "wxsph-before-save.png")
                logger.info(
                    "微信小店 SKU、类目属性、全款预售 15 天与重量填写回读完成；"
                    + (
                        "即将保存微信小店资料，不铺货"
                        if args.save
                        else "本阶段仅预览，未保存、未铺货"
                    )
                )
            elif xhs_requested:
                assert product.xhs_fields is not None
                xhs = XhsFormListing(
                    page,
                    drawer,
                    logger,
                    attribute_runtime=attribute_runtime,
                    category_hints=product.category_hints,
                )
                await xhs.open()
                try:
                    xhs_report = await xhs.apply_excel_fields(
                        product.xhs_fields,
                        title=product.title,
                        style_code=product.style_code,
                        portrait_paths=product.main_images_34,
                        timeout_seconds=args.upload_timeout,
                        uploader=sync_image_group,
                    )
                except XhsFormListingError:
                    await safe_screenshot(page, artifact_dir / "xhs-error.png")
                    raise
                (artifact_dir / "xhs-before-save.json").write_text(
                    json.dumps(xhs_report, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                await safe_screenshot(page, artifact_dir / "xhs-before-save.png")
                logger.info(
                    "小红书资料填写与保存前复核完成；%s",
                    "即将保存但不铺货"
                    if args.save and args.save_only
                    else "即将保存并铺货到小红书指定店铺"
                    if args.save
                    else "小红书资料未保存、未铺货",
                )
            elif youzan_requested:
                assert product.youzan_fields is not None
                youzan = YouzanFormListing(
                    page,
                    drawer,
                    logger,
                    attribute_runtime=attribute_runtime,
                    category_hints=product.category_hints,
                )
                await youzan.open()
                try:
                    youzan_report = await youzan.apply_excel_fields(
                        product.youzan_fields,
                        style_code=product.style_code,
                    )
                except YouzanFormListingError:
                    await safe_screenshot(page, artifact_dir / "youzan-error.png")
                    raise
                (artifact_dir / "youzan-before-save.json").write_text(
                    json.dumps(youzan_report, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                await safe_screenshot(page, artifact_dir / "youzan-before-save.png")
                logger.info(
                    "有赞资料填写与保存前复核完成；%s",
                    "即将保存但不铺货"
                    if args.save
                    else "有赞资料未保存、未铺货",
                )
            elif jd_requested:
                assert product.jd_fields is not None
                jd = JdFormListing(
                    page,
                    drawer,
                    logger,
                    attribute_runtime=attribute_runtime,
                )
                await jd.open()
                try:
                    jd_report = await jd.apply_excel_fields(
                        product.jd_fields,
                        style_code=product.style_code,
                        square_paths=product.main_images,
                        portrait_paths=product.main_images_34,
                        sku_paths=product.sku_images,
                    )
                except JdFormListingError:
                    await safe_screenshot(page, artifact_dir / "jd-error.png")
                    raise
                (artifact_dir / "jd-before-save.json").write_text(
                    json.dumps(jd_report, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                await jd.scroll_label_into_view("风格")
                await safe_screenshot(page, artifact_dir / "jd-style-before-save.png")
                await jd.scroll_color_image_groups_into_view()
                await safe_screenshot(page, artifact_dir / "jd-images-before-save.png")
                await safe_screenshot(page, artifact_dir / "jd-before-save.png")
                logger.info(
                    "京东类目、品牌、Excel 属性、SKU 价格库存、厚度、发货时效和商品图片填写复核完成；%s",
                    "即将保存但不铺货" if args.save else "本阶段仅预览，未保存、未铺货",
                )
            elif tmall_requested:
                assert product.tmall_fields is not None
                assert tmall_assets is not None
                tmall_api_index = TmallApiJsonIndex(page, logger)
                tmall_api_index.install()
                tmall = TmallFormListing(
                    page,
                    drawer,
                    logger,
                    api_index=tmall_api_index,
                    attribute_runtime=attribute_runtime,
                    category_hints=product.category_hints,
                )
                tmall_report: Dict[str, Any] = {
                    "status": "in_progress",
                }
                tmall_stage = "open"
                try:
                    await tmall.open()
                    tmall_stage = "category"
                    category = await tmall.apply_recommended_category()
                    tmall_report["category"] = category
                    tmall_stage = "product_identity"
                    identity = await tmall.fill_product_identity(product.tmall_fields)
                    tmall_report["product_identity"] = identity
                    tmall_stage = "full_form"
                    try:
                        await tmall.require_full_form()
                        tmall_report["form_mode"] = "existing"
                        expected_image_count = len(product.main_images)
                        existing_image_state = (
                            await tmall.inspect_initial_product_images(
                                expected_count=expected_image_count
                            )
                        )
                        tmall_report["existing_product_images_before"] = (
                            existing_image_state
                        )
                        if existing_image_state.get("decoded") is True:
                            logger.info(
                                "天猫已有完整表单：顶部产品图片完整（%s/%s），直接填写",
                                existing_image_state.get("count", 0),
                                expected_image_count,
                            )
                        else:
                            logger.info(
                                "天猫已有完整表单但顶部产品图片不完整（%s/%s）；"
                                "短暂等待页面回显，仍不足则按本地天猫顺序强制补齐",
                                existing_image_state.get("count", 0),
                                expected_image_count,
                            )
                            tmall_stage = "existing_product_images_grace_wait"
                            try:
                                existing_image_state = (
                                    await tmall.wait_for_initial_product_images(
                                        expected_count=expected_image_count,
                                        timeout_seconds=min(
                                            3.0, float(args.upload_timeout)
                                        ),
                                        initial_delay_seconds=0,
                                        stable_seconds=0.5,
                                        allow_existing_product_shortcut=False,
                                    )
                                )
                            except TmallFormListingError:
                                logger.info(
                                    "天猫重复打开后顶部产品图片未自行补齐；"
                                    "开始用本地图片强制覆盖"
                                )
                            if existing_image_state.get("decoded") is not True:
                                tmall_stage = "existing_product_images"
                                tmall_report["existing_product_images"] = (
                                    await tmall.sync_initial_product_images(
                                        product.main_images,
                                        timeout_seconds=args.upload_timeout,
                                        uploader=sync_image_group,
                                    )
                                )
                                tmall_stage = "existing_product_images_ready"
                                existing_image_state = (
                                    await tmall.wait_for_initial_product_images(
                                        expected_count=expected_image_count,
                                        timeout_seconds=args.upload_timeout,
                                        initial_delay_seconds=0,
                                        allow_existing_product_shortcut=False,
                                    )
                                )
                            if existing_image_state.get("decoded") is not True:
                                raise TmallFormListingError(
                                    "天猫重复打开后的顶部产品图片仍不完整，禁止继续"
                                )
                            tmall_report["existing_product_images_ready"] = (
                                existing_image_state
                            )
                            logger.info(
                                "天猫重复打开后的顶部产品图片已补齐并回读通过：%s/%s",
                                existing_image_state.get("count", 0),
                                expected_image_count,
                            )
                    except TmallProductWriteRequired:
                        tmall_report["form_mode"] = "initial"
                        logger.info(
                            "天猫首次填写：短暂等待基础图片自动继承；"
                            "仍不完整则直接用本地图片补齐"
                        )
                        tmall_stage = "initial_product_images_ready"
                        initial_image_state = (
                            await tmall.inspect_initial_product_images(
                                expected_count=len(product.main_images)
                            )
                        )
                        if initial_image_state.get("decoded") is not True:
                            try:
                                initial_image_state = (
                                    await tmall.wait_for_initial_product_images(
                                        expected_count=len(product.main_images),
                                        timeout_seconds=min(
                                            3.0, float(args.upload_timeout)
                                        ),
                                        initial_delay_seconds=0,
                                        stable_seconds=0.5,
                                    )
                                )
                            except TmallFormListingError:
                                logger.info(
                                    "天猫首次产品图片未在短暂等待内自动继承；"
                                    "开始用本地图片强制覆盖"
                                )
                        tmall_report["initial_product_images_ready"] = (
                            initial_image_state
                        )
                        tmall_stage = "initial_product_images"
                        tmall_report["initial_product_images"] = (
                            await tmall.sync_initial_product_images(
                                product.main_images,
                                timeout_seconds=args.upload_timeout,
                                uploader=sync_image_group,
                            )
                        )
                        matched_existing = tmall_report["initial_product_images_ready"].get(
                            "matched_existing_product"
                        ) is True
                        if matched_existing:
                            # 自动匹配既有天猫产品不是首次发布成功。旧产品
                            # 只有一张图时，上方已按本地顺序补齐，再复核全部图。
                            tmall_stage = "matched_product_images_ready"
                            ready = await tmall.wait_for_initial_product_images(
                                expected_count=len(product.main_images),
                                timeout_seconds=args.upload_timeout,
                                initial_delay_seconds=0,
                                allow_existing_product_shortcut=False,
                            )
                            if ready.get("decoded") is not True:
                                raise TmallFormListingError("天猫既有产品图片替换后仍不完整，禁止保存")
                            tmall_report["matched_product_images_ready"] = ready
                        # 首次天猫资料与已经生成过完整表单的资料是两套页面。
                        # 首次页先异步继承顶部五张产品图；只有图片完整且按
                        # 天猫 3、2、1、4、5 顺序替换后，才允许点击该区块的
                        # 蓝色“发布”生成下半页表单。底部保存/铺货仍完全由
                        # args.save 和 args.publish 门禁控制。
                        tmall_stage = "initial_product_publish"
                        initial_product_publish = {"clicked": False, "matched_existing_product": True} if matched_existing else (
                            await tmall.publish_product_information(
                                timeout_seconds=args.timeout,
                                expected_image_count=len(product.main_images),
                            )
                        )
                        tmall_report.update(
                            {
                                "form_mode": "matched_existing" if matched_existing else "initial",
                                "initial_product_publish": initial_product_publish,
                            }
                        )
                    tmall_report["api_json_validation"] = (
                        await tmall_api_index.safe_summary(
                            getattr(tmall, "panel", None)
                        )
                    )
                    logger.info(
                        "天猫接口 JSON 辅助校验：接口=%s，API 字段=%s，"
                        "DOM 匹配=%s",
                        tmall_report["api_json_validation"].get(
                            "json_endpoint_count", 0
                        ),
                        tmall_report["api_json_validation"].get(
                            "api_field_count", 0
                        ),
                        tmall_report["api_json_validation"].get(
                            "matched_field_count", 0
                        ),
                    )
                except TmallFormListingError:
                    tmall_report["api_json_validation"] = (
                        await tmall_api_index.safe_summary(
                            getattr(tmall, "panel", None)
                        )
                    )
                    await write_tmall_review_required(
                        page,
                        artifact_dir,
                        tmall_report,
                        reason_code=f"tmall_{tmall_stage}_review_required",
                    )
                    raise
                try:
                    tmall_stage = "size_source"
                    size_sources = resolve_tmall_size_sources(
                        product.tmall_fields,
                        category,
                        product_dir=product.product_dir,
                        size_chart_image=tmall_assets.parameter_image,
                    )
                    logger.info(
                        "天猫尺码来源校验完成：类型=%s，来源字段=%s",
                        size_sources.category_kind,
                        " / ".join(size_sources.headers),
                    )
                    tmall_stage = "attributes"
                    attributes = await tmall.fill_attributes(product.tmall_fields)
                    tmall_stage = "specifications"
                    specifications = await tmall.normalize_synced_specifications(category)
                    tmall_stage = "attribute_images"
                    attribute_image_actions = await tmall.sync_attribute_images(
                        product.sku_images,
                        timeout_seconds=args.upload_timeout,
                        uploader=sync_image_group,
                    )
                    tmall_stage = "sku_batch"
                    sku_batch = await tmall.fill_sku_batch(product.tmall_fields.fields)
                    tmall_stage = "size_chart"
                    tmall_stage_started = time.monotonic()
                    logger.info("天猫下方表单：开始处理尺码表")
                    size_chart = await tmall.fill_size_chart(size_sources.rows)
                    logger.info(
                        "天猫下方表单：尺码表已完成，耗时 %.1f 秒",
                        time.monotonic() - tmall_stage_started,
                    )
                    tmall_stage = "sales_and_logistics"
                    tmall_stage_started = time.monotonic()
                    logger.info("天猫下方表单：开始处理销售与物流")
                    sales_and_logistics = await tmall.fill_sales_and_logistics(
                        product.tmall_fields
                    )
                    logger.info(
                        "天猫下方表单：销售与物流已完成，耗时 %.1f 秒",
                        time.monotonic() - tmall_stage_started,
                    )
                    tmall_stage = "images"
                    image_actions = await tmall.sync_required_images(
                        tmall_assets,
                        timeout_seconds=args.upload_timeout,
                        uploader=sync_image_group,
                    )
                    tmall_stage = "inherited_images"
                    inherited_image_actions = await tmall.sync_inherited_main_images(
                        product.main_images,
                        product.main_images_34,
                        timeout_seconds=args.upload_timeout,
                        uploader=sync_image_group,
                    )
                    image_actions = {
                        **image_actions,
                        **attribute_image_actions,
                    }
                    tmall_stage = "after_sales"
                    after_sales = await tmall.fill_after_sales(product.tmall_fields)
                    # 图片、物流和售后字段会让 Vue 迟到重渲染数值输入框；
                    # 把纯展示格式清理放在全部写入操作之后，防止 155 再显示为
                    # 155.00。这里不改数据模型，也不触发保存。
                    tmall_stage = "size_display_cleanup"
                    await tmall.clean_size_chart_integer_displays()
                    tmall_stage = "remaining_required"
                    if (
                        attribute_runtime is not None
                        and attribute_runtime.has_deferred_reviews
                    ):
                        required_validation = {
                            "status": "deferred_for_operator_review"
                        }
                    else:
                        required_validation = (
                            await tmall.validate_remaining_required_fields()
                        )
                except TmallSizeReviewRequired as exc:
                    tmall_report["api_json_validation"] = (
                        await tmall_api_index.safe_summary(
                            getattr(tmall, "panel", None)
                        )
                    )
                    await write_tmall_review_required(
                        page,
                        artifact_dir,
                        tmall_report,
                        reason_code=exc.reason_code,
                    )
                    raise
                except (TmallSizeSourceError, TmallFormListingError):
                    tmall_report["api_json_validation"] = (
                        await tmall_api_index.safe_summary(
                            getattr(tmall, "panel", None)
                        )
                    )
                    reason_code = (
                        "tmall_required_fields_review_required"
                        if tmall_stage == "remaining_required"
                        else "tmall_{0}_review_required".format(tmall_stage)
                    )
                    await write_tmall_review_required(
                        page,
                        artifact_dir,
                        tmall_report,
                        reason_code=reason_code,
                    )
                    raise
                tmall_report.update(
                    {
                        "status": "ready_for_manual_review",
                        "attributes": attributes,
                        "specifications": specifications,
                        "attribute_images": attribute_image_actions,
                        "sku_batch": sku_batch,
                        "size_chart": size_chart,
                        "sales_and_logistics": sales_and_logistics,
                        "images": image_actions,
                        "inherited_images": inherited_image_actions,
                        "after_sales": after_sales,
                        "required_validation": required_validation,
                        "saved": False,
                        "published": False,
                    }
                )
                tmall_report["api_json_validation"] = (
                    await tmall_api_index.safe_summary(
                        getattr(tmall, "panel", None)
                    )
                )
                tmall_api_index.uninstall()
                (artifact_dir / "tmall-before-save.json").write_text(
                    json.dumps(
                        tmall_preview_summary(tmall_report),
                        ensure_ascii=False,
                        indent=2,
                    ),
                    encoding="utf-8",
                )
                await safe_screenshot(page, artifact_dir / "tmall-before-save.png")
                if args.save and not args.save_only:
                    next_action = "即将保存并铺货"
                elif args.save:
                    next_action = "即将保存但不铺货"
                else:
                    next_action = "天猫资料未保存、未铺货"
                logger.info("天猫资料填写与保存前复核完成；%s", next_action)
            else:
                await safe_screenshot(page, artifact_dir / "before-save.png")

            if attribute_runtime is not None:
                attribute_runtime.raise_deferred_reviews()

            if not args.save:
                logger.info(
                    "已完成所选范围的资料填写与校验；--no-save 已启用，"
                    "未保存平台资料或铺货"
                )
                return

            if args.platform == "base" and base_save_result is not None:
                (artifact_dir / "save-result.json").write_text(
                    json.dumps(base_save_result, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                await safe_screenshot(page, artifact_dir / "after-save.png")
                logger.info("基础资料已保存且当前页面已稳定；未选择平台资料填写")
                return

            publish_preview_target: Optional[Tuple[str, Sequence[str]]] = None
            publish_preview_slug = ""
            if taobao_requested and getattr(args, "taobao_publish_preview", False):
                publish_preview_target = ("淘宝", DEFAULT_TAOBAO_PUBLISH_SHOPS)
                publish_preview_slug = "taobao"
            elif youzan_requested and getattr(args, "youzan_publish_preview", False):
                publish_preview_target = ("有赞", DEFAULT_YOUZAN_PUBLISH_SHOPS)
                publish_preview_slug = "youzan"
            elif wxsph_requested and getattr(args, "wxsph_publish_preview", False):
                publish_preview_target = (
                    "微信小店（视频号）",
                    DEFAULT_WXSPH_PUBLISH_SHOPS,
                )
                publish_preview_slug = "wxsph"
            elif jd_requested and getattr(args, "jd_publish_preview", False):
                publish_preview_target = ("京东", DEFAULT_JD_PUBLISH_SHOPS)
                publish_preview_slug = "jd"

            if publish_preview_target is not None:
                preview_platform_name, preview_shops = publish_preview_target
                trigger_result = await click_save_and_confirm(
                    page,
                    drawer,
                    args.sync_erp,
                    args.timeout,
                    logger,
                    button_text="保存并铺货到平台",
                )
                publish_dialog, preview = await prepare_taobao_publish_dialog(
                    page,
                    preview_shops,
                    args.timeout,
                    logger,
                    platform_name=preview_platform_name,
                )
                trigger_result["shop_publish_preview"] = preview
                (artifact_dir / f"{publish_preview_slug}-publish-preview.json").write_text(
                    json.dumps(trigger_result, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                await safe_screenshot(
                    page,
                    artifact_dir / f"{publish_preview_slug}-publish-preview.png",
                )
                await close_publish_preview(publish_dialog)
                logger.info(
                    "%s铺货预览已关闭：已精确勾选%s，"
                    "未点击最终确定，未铺货",
                    preview_platform_name,
                    "、".join(preview.get("selected_shops") or ()),
                )
                return

            commerce_publish_target = resolve_commerce_publish_target(args)
            action_text, should_publish = resolve_platform_save_action(
                args,
                publish_mode=publish_mode,
            )
            result = await click_save_and_confirm(
                page,
                drawer,
                args.sync_erp,
                args.timeout,
                logger,
                button_text=action_text,
            )
            if base_save_result is not None:
                result["base_save"] = base_save_result
            if should_publish:
                if commerce_publish_target is not None:
                    platform_name, selected_shops = commerce_publish_target
                    publish_dialog, selection = await prepare_taobao_publish_dialog(
                        page,
                        selected_shops,
                        args.timeout,
                        logger,
                        platform_name=platform_name,
                    )
                    await safe_screenshot(
                        page,
                        artifact_dir / f"{args.platform}-before-publish.png",
                    )
                    publish_result = await submit_taobao_publish_dialog(
                        page,
                        publish_dialog,
                        selection,
                        args.timeout,
                        logger,
                        platform_name=platform_name,
                    )
                else:
                    publish_result = await publish_to_selected_douyin_shops(
                        page,
                        args.publish_shop,
                        args.timeout,
                        logger,
                    )
                if publish_result.get("submitted", True):
                    publish_result["progress_dialog"] = (
                        await dismiss_publish_progress_dialog(
                            page,
                            logger,
                            appearance_timeout_seconds=3.0,
                        )
                    )
                result["shop_publish"] = publish_result
                if commerce_publish_target is not None:
                    result["confirmed_by"] = publish_result.get(
                        "confirmed_by", result.get("confirmed_by")
                    )
            if publish_mode and args.save_only:
                persisted_drawer, persisted_reopened = (
                    await product_editor_for_saved_readback(
                        page,
                        drawer,
                        shared_session,
                        product,
                        timeout_seconds=args.timeout,
                        logger=logger,
                        platform_label="抖音资料持久化结果",
                    )
                )
                persisted_douyin = DouyinListing(
                    page,
                    persisted_drawer,
                    logger,
                    artifact_dir,
                )
                await persisted_douyin.open()
                persisted_report = await persisted_douyin.verify_persisted_values(
                    category_fields["category"],
                    category_fields["product_title"],
                    category_fields["short_title"],
                    category_fields["attributes"],
                    product.douyin_fields.price,
                    product.douyin_fields.spot_stock,
                    product.douyin_fields.presale_stock,
                )
                (artifact_dir / "douyin-after-save-validation.json").write_text(
                    json.dumps(persisted_report, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                logger.info(
                    "保存后%s复核通过：抖音字段与 %s 行 SKU 均已持久化",
                    "重开" if persisted_reopened else "当前页",
                    sku_rows,
                )

            if pdd_requested:
                assert product.pdd_fields is not None
                persisted_drawer, persisted_reopened = (
                    await product_editor_for_saved_readback(
                        page,
                        drawer,
                        shared_session,
                        product,
                        timeout_seconds=args.timeout,
                        logger=logger,
                        platform_label="拼多多价格库存",
                    )
                )
                persisted_pdd = PddFormListing(
                    page,
                    persisted_drawer,
                    logger,
                )
                await persisted_pdd.open()
                persisted_report = (
                    await persisted_pdd.verify_persisted_price_inventory(
                        product.pdd_fields.fields
                    )
                )
                (artifact_dir / "pdd-after-save-validation.json").write_text(
                    json.dumps(persisted_report, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                logger.info(
                    "保存后%s复核通过：拼多多拼单价、单买价与 %s 行库存均已持久化",
                    "重开" if persisted_reopened else "当前页",
                    persisted_report["row_count"],
                )

            if wxsph_requested:
                assert product.wxsph_fields is not None
                persisted_drawer, persisted_reopened = (
                    await product_editor_for_saved_readback(
                        page,
                        drawer,
                        shared_session,
                        product,
                        timeout_seconds=args.timeout,
                        logger=logger,
                        platform_label="微信小店关键字段",
                    )
                )
                persisted_wxsph = WxsphFormListing(
                    page,
                    persisted_drawer,
                    logger,
                )
                await persisted_wxsph.open()
                persisted_report = await persisted_wxsph.verify_persisted_values(
                    product.wxsph_fields,
                    expected_attributes=wxsph_report["attributes"]["attributes"],
                    expected_freight=wxsph_report.get("freight"),
                )
                (artifact_dir / "wxsph-after-save-validation.json").write_text(
                    json.dumps(persisted_report, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                logger.info(
                    "保存后%s复核通过：微信小店类目属性、售卖价、市场价、%s 行库存、全款预售 15 天和重量均已持久化",
                    "重开" if persisted_reopened else "当前页",
                    persisted_report["row_count"],
                )

            if xhs_requested:
                assert product.xhs_fields is not None
                persisted_drawer, persisted_reopened = (
                    await product_editor_for_saved_readback(
                        page,
                        drawer,
                        shared_session,
                        product,
                        timeout_seconds=args.timeout,
                        logger=logger,
                        platform_label="小红书关键字段",
                    )
                )
                persisted_xhs = XhsFormListing(page, persisted_drawer, logger)
                await persisted_xhs.open()
                persisted_report = await persisted_xhs.verify_persisted_values(
                    product.xhs_fields,
                    title=product.title,
                    style_code=product.style_code,
                    expected_attributes=xhs_report["attributes"]["attributes"],
                    expected_freight=xhs_report.get("freight"),
                )
                (artifact_dir / "xhs-after-save-validation.json").write_text(
                    json.dumps(persisted_report, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                logger.info(
                    "保存后%s复核通过：小红书标题、货号、类目属性、%s 行 SKU 价格库存和全款预售 15 天均已持久化",
                    "重开" if persisted_reopened else "当前页",
                    persisted_report["row_count"],
                )

            if youzan_requested:
                assert product.youzan_fields is not None
                persisted_drawer, persisted_reopened = (
                    await product_editor_for_saved_readback(
                        page,
                        drawer,
                        shared_session,
                        product,
                        timeout_seconds=args.timeout,
                        logger=logger,
                        platform_label="有赞关键字段",
                        # 有赞保存后旧表格可能暂时把 SKU 重量重置为 0；
                        # 刷新读取服务器保存值，不用旧 DOM 判断持久化失败。
                        force_reload=True,
                    )
                )
                persisted_youzan = YouzanFormListing(
                    page,
                    persisted_drawer,
                    logger,
                )
                await persisted_youzan.open()
                persisted_report = await persisted_youzan.verify_persisted_values(
                    product.youzan_fields
                )
                (artifact_dir / "youzan-after-save-validation.json").write_text(
                    json.dumps(persisted_report, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                logger.info(
                    "保存后%s复核通过：有赞类目、SKU、重量、库存扣减、配送和运费模板均已持久化",
                    "重开" if persisted_reopened else "当前页",
                )

            if jd_requested:
                assert product.jd_fields is not None
                persisted_drawer, persisted_reopened = (
                    await product_editor_for_saved_readback(
                        page,
                        drawer,
                        shared_session,
                        product,
                        timeout_seconds=args.timeout,
                        logger=logger,
                        platform_label="京东关键字段",
                    )
                )
                persisted_jd = JdFormListing(page, persisted_drawer, logger)
                await persisted_jd.open()
                persisted_report = await persisted_jd.verify_persisted_values(
                    product.jd_fields,
                    style_code=product.style_code,
                    expected_attributes=jd_report["attributes"]["attributes"],
                )
                (artifact_dir / "jd-after-save-validation.json").write_text(
                    json.dumps(persisted_report, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                logger.info(
                    "保存后%s复核通过：京东类目、品牌、SKU 京东价、库存、底部价格和 48 小时发货均已持久化",
                    "重开" if persisted_reopened else "当前页",
                )

            result_name = "publish-result.json" if should_publish else "save-result.json"
            (artifact_dir / result_name).write_text(
                json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            screenshot_name = "after-publish.png" if should_publish else "after-save.png"
            await safe_screenshot(page, artifact_dir / screenshot_name)
            if result.get("confirmed_by") == "background_task_submitted":
                submitted_shops = tuple(
                    result.get("shop_publish", {}).get("submitted_shops") or ()
                )
                logger.info(
                    "%s后台任务已提交（店铺：%s）",
                    action_text,
                    "、".join(submitted_shops) or "未记录",
                )
            elif result.get("confirmed_by") == "already_published":
                logger.info(
                    "平台资料保存成功；指定店铺此前均已铺货，本次未重复提交"
                )
            else:
                logger.info(
                    "%s成功（确认来源：%s）",
                    action_text,
                    result.get("confirmed_by"),
                )
        except Exception:
            await safe_screenshot(page, artifact_dir / "error.png")
            raise
        finally:
            if shared_session is not None:
                # 全平台流程由最外层统一释放；阶段之间保留当前编辑抽屉。
                pass
            elif args.cdp_url or attached_existing_profile:
                # 不关闭用户自行启动或本次接管的 CDP Chrome。
                pass
            else:
                await context.close()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="从 Excel 和产品素材自动填写快麦基础/平台资料")
    parser.add_argument("--excel-url", default=DEFAULT_EXCEL_URL, help="产品信息.xlsx 的 smb:// 或本地路径")
    parser.add_argument("--create-product", action="store_true", help="手工新增商品链接；与 --platform base 搭配，仅创建快麦商品，不铺货")
    parser.add_argument(
        "--platform",
        type=normalize_platform_selection,
        metavar="PLATFORM[,PLATFORM...]",
        default="all",
        help=(
            "运行范围：all=全部已实现平台，base=仅基础资料，"
            "douyin=抖音，taobao=淘宝，tmall=天猫，"
            "pdd=拼多多，wxsph=微信小店，xhs=小红书，youzan=有赞，jd=京东；"
            "可用逗号指定多个平台，按输入顺序执行，例如 douyin,jd,xhs"
        ),
    )
    parser.add_argument("--dry-run", action="store_true", help="只读取并校验 Excel/图片，不打开浏览器")
    parser.add_argument(
        "--inspect-only",
        action="store_true",
        help="只发现脚手平台字段结构，不填写、不保存、不铺货",
    )
    parser.add_argument(
        "--wash-label-upload-test",
        action="store_true",
        help=(
            "仅打开抖音资料并测试水洗标/吊牌图上传；不填写其它字段、"
            "不保存、不铺货，需同时使用 --platform douyin --no-save"
        ),
    )
    parser.add_argument(
        "--save",
        dest="save",
        action="store_true",
        default=True,
        help=(
            "最后保存；抖音、淘宝、天猫、拼多多、微信小店、小红书、有赞或京东会点击"
            "“保存并铺货到平台”（默认）"
        ),
    )
    parser.add_argument(
        "--no-save",
        dest="save",
        action="store_false",
        help="只填写并校验所选平台资料，不保存平台资料、不铺货，也不改动基础资料",
    )
    parser.add_argument(
        "--save-only",
        action="store_true",
        help=(
            "平台资料只保存、不铺货；"
            "全平台模式下会依次保存抖音、淘宝、天猫、拼多多、微信小店、小红书、有赞和京东资料"
        ),
    )
    parser.add_argument(
        "--allow-taobao-save-once",
        action="store_true",
        help="仅对本次命令解除淘宝开发预览的保存保护",
    )
    parser.add_argument(
        "--taobao-publish-preview",
        action="store_true",
        help=(
            "淘宝无铺货预览：打开‘保存并铺货到平台’弹窗，"
            "选择淘宝平台且只勾选‘钊叔制’，截图后关闭，不点击最终确定"
        ),
    )
    parser.add_argument(
        "--youzan-publish-preview",
        action="store_true",
        help=(
            "有赞无铺货预览：保存有赞资料并打开‘铺货到店铺’弹窗，"
            "只勾选有赞的‘NEIGBORL官方旗舰店’，截图后关闭，"
            "不点击最终确定"
        ),
    )
    parser.add_argument(
        "--wxsph-publish-preview",
        action="store_true",
        help=(
            "微信小店无铺货预览：保存微信小店资料并打开‘铺货到店铺’弹窗，"
            "只勾选‘NEIGBORL钊叔制鞋服’和‘NEIGBORL钊叔制造局’，"
            "截图后关闭，不点击最终确定"
        ),
    )
    parser.add_argument(
        "--jd-publish-preview",
        action="store_true",
        help=(
            "京东无铺货预览：保存京东资料并打开‘铺货到店铺’弹窗，"
            "只勾选‘NEIGBORL服饰旗舰店’，截图后关闭，不点击最终确定"
        ),
    )
    parser.add_argument(
        "--allow-taobao-publish-once",
        action="store_true",
        help=(
            "仅对本次命令解除淘宝铺货保护：完整填写后选择淘宝平台，"
            "只勾选‘钊叔制’并提交铺货"
        ),
    )
    parser.add_argument(
        "--allow-tmall-publish-once",
        action="store_true",
        help=(
            "兼容旧命令行参数；天猫已是正式流程，不再需要此参数"
        ),
    )
    parser.add_argument(
        "--taobao-category-mode",
        choices=("auto", "casual-pants", "recommended"),
        default="auto",
        help=(
            "淘宝类目模式：auto=按品类选择（默认，裤装沿用已验证休闲裤类目，"
            "其他品类使用页面唯一推荐）；casual-pants=固定休闲裤；"
            "recommended=页面唯一推荐类目"
        ),
    )
    parser.add_argument(
        "--taobao-test-scope",
        choices=(
            "full",
            "attributes",
            "category-size",
            "sku-batch",
            "payment-service",
        ),
        default="full",
        help=(
            "淘宝预览范围：full=完整淘宝流程（默认）；"
            "attributes=仅验证类目属性；"
            "category-size=仅验证类目搜索和尺码表；"
            "sku-batch=仅验证类目和 SKU 批量字段；"
            "payment-service=仅验证库存扣减与售后服务"
        ),
    )
    parser.add_argument("--sync-erp", action="store_true", help="出现 ERP 同步确认框时选择同步；默认跳过同步")
    parser.add_argument("--headless", action="store_true", help="无头模式（仅适用于专用 Chrome 配置已登录）")
    parser.add_argument("--cdp-url", help="连接已开启远程调试的 Chrome，例如 http://127.0.0.1:9222")
    parser.add_argument(
        "--user-data-dir",
        default=str(DEFAULT_PROFILE_DIR),
        help="自动化专用 Chrome 用户数据目录",
    )
    parser.add_argument(
        "--auth-state",
        default=str(DEFAULT_AUTH_STATE),
        help="会话登录状态保存文件",
    )
    parser.add_argument("--login-timeout", type=int, default=1200, help="首次手工登录等待秒数（默认 1200）")
    parser.add_argument("--timeout", type=int, default=300, help="页面操作/保存超时秒数（默认 300）")
    parser.add_argument("--upload-timeout", type=int, default=600, help="每组图片上传超时秒数（默认 600）")
    parser.add_argument(
        "--publish-shop",
        action="append",
        default=None,
        help="要铺货的抖音店铺，可重复传入；不传时使用已配置的 6 个店铺",
    )
    parser.add_argument(
        "--all-platform-one-shop-test",
        action="store_true",
        help=(
            "仅用于全平台真实铺货测试：每个平台只选当前正式配置的第一家店；"
            "不传时正式店铺配置完全不变"
        ),
    )
    parser.add_argument(
        "--all-platform-start-at",
        choices=("douyin", "taobao", "tmall", "pdd", "wxsph", "xhs", "youzan", "jd"),
        default=None,
        help="全平台失败恢复时从指定平台开始，不重跑前面已完成平台",
    )
    parser.add_argument(
        "--all-platform-skip",
        choices=("douyin", "taobao", "tmall", "pdd", "wxsph", "xhs", "youzan", "jd"),
        action="append",
        default=[],
        help="本次全平台流程跳过指定平台，可重复使用",
    )
    parser.add_argument(
        "--learning-enabled",
        action="store_true",
        help="启用 AI 学习检查点和审核服务",
    )
    parser.add_argument(
        "--learning-db",
        default=os.environ.get(
            "KUAIMAI_LEARNING_DB", ".local-state/learning.sqlite3"
        ),
        help="本地学习缓存数据库",
    )
    parser.add_argument(
        "--learning-api-url",
        default=os.environ.get("KUAIMAI_LEARNING_API_URL", ""),
        help="AI 学习审核服务地址",
    )
    return parser


def validate_execution_mode(args: argparse.Namespace) -> Tuple[PlatformSpec, ...]:
    try:
        args.platform = normalize_platform_selection(args.platform)
    except ValueError as exc:
        raise SystemExit(str(exc)) from exc
    if getattr(args, "wash_label_upload_test", False) and (
        args.platform != "douyin"
        or args.save
        or args.save_only
        or args.dry_run
        or args.inspect_only
    ):
        raise SystemExit(
            "--wash-label-upload-test 仅允许用于 --platform douyin --no-save，"
            "且不能与保存、铺货或 inspect/dry-run 组合"
        )
    if getattr(args, "create_product", False):
        if args.platform != "base" or args.inspect_only:
            raise SystemExit("--create-product 只允许 --platform base，不能与字段发现或平台铺货组合")
    selected = expand_platform_selection(args.platform)
    if len(selected) > 1 and not args.save and any(spec.cli_name == "base" for spec in selected):
        raise SystemExit("多平台仅填写模式不能包含基础资料；请只选择所需平台，或使用 --save-only")
    scaffold = tuple(spec for spec in selected if spec.lifecycle == "scaffold")
    save_forbidden = tuple(
        spec for spec in selected if spec.save_policy == "forbidden"
    )
    blocking_save_forbidden = () if args.platform == "all" else save_forbidden

    if getattr(args, "all_platform_one_shop_test", False) and (
        args.platform != "all"
        or not args.save
        or args.save_only
        or args.dry_run
        or args.inspect_only
    ):
        raise SystemExit(
            "--all-platform-one-shop-test 只允许用于 --platform all --save 真实铺货测试"
        )
    if getattr(args, "all_platform_start_at", None) and args.platform != "all":
        raise SystemExit("--all-platform-start-at 只允许用于 --platform all")
    if getattr(args, "all_platform_skip", None) and args.platform != "all":
        raise SystemExit("--all-platform-skip 只允许用于 --platform all")

    if args.save_only and not args.save:
        raise SystemExit("--save-only 与 --no-save 不能同时使用")
    if args.allow_taobao_save_once and args.platform != "taobao":
        raise SystemExit("--allow-taobao-save-once 仅允许用于淘宝流程")
    if args.taobao_publish_preview and args.allow_taobao_publish_once:
        raise SystemExit(
            "--taobao-publish-preview 与 --allow-taobao-publish-once 不能同时使用"
        )
    if args.taobao_publish_preview and (
        args.platform != "taobao"
        or not args.save
        or not args.allow_taobao_save_once
        or args.taobao_test_scope != "full"
    ):
        raise SystemExit(
            "--taobao-publish-preview 仅允许用于淘宝完整流程，"
            "并必须同时使用 --allow-taobao-save-once"
        )
    if args.youzan_publish_preview and (
        args.platform != "youzan"
        or not args.save
        or args.save_only
        or args.dry_run
        or args.inspect_only
    ):
        raise SystemExit(
            "--youzan-publish-preview 仅允许用于有赞完整流程："
            "会保存资料并勾选店铺，但不点击铺货弹窗的最终确定"
        )
    if args.wxsph_publish_preview and (
        args.platform != "wxsph"
        or not args.save
        or args.save_only
        or args.dry_run
        or args.inspect_only
    ):
        raise SystemExit(
            "--wxsph-publish-preview 仅允许用于微信小店完整流程："
            "会保存资料并勾选店铺，但不点击铺货弹窗的最终确定"
        )
    if args.jd_publish_preview and (
        args.platform != "jd"
        or not args.save
        or args.save_only
        or args.dry_run
        or args.inspect_only
    ):
        raise SystemExit(
            "--jd-publish-preview 仅允许用于京东完整流程："
            "会保存资料并勾选店铺，但不点击铺货弹窗的最终确定"
        )
    if args.allow_taobao_publish_once and (
        args.platform != "taobao"
        or not args.save
        or args.save_only
        or args.taobao_test_scope != "full"
    ):
        raise SystemExit(
            "--allow-taobao-publish-once 仅允许用于淘宝完整保存铺货流程"
        )
    if args.allow_tmall_publish_once and (
        args.platform != "tmall"
        or not args.save
        or args.save_only
        or args.dry_run
        or args.inspect_only
    ):
        raise SystemExit(
            "--allow-tmall-publish-once 仅作为天猫正式流程兼容参数，"
            "必须用于天猫完整保存铺货流程"
        )
    if args.inspect_only and args.dry_run:
        raise SystemExit("--inspect-only 与 --dry-run 不能同时使用")
    if args.taobao_test_scope != "full" and args.save:
        raise SystemExit("淘宝独立功能预览必须同时使用 --no-save")
    if blocking_save_forbidden and (
        args.save_only
        or args.allow_taobao_save_once
        or args.allow_taobao_publish_once
        or args.allow_tmall_publish_once
        or (args.save and not args.dry_run)
    ):
        names = "、".join(spec.display_name for spec in blocking_save_forbidden)
        raise SystemExit(
            f"{names}当前只允许预览：请使用 --no-save；"
            "保存、铺货和淘宝一次性授权都不能解除该门禁"
        )
    if scaffold and (
        args.save_only
        or args.allow_taobao_save_once
        or args.allow_taobao_publish_once
        or args.allow_tmall_publish_once
    ):
        raise SystemExit("字段发现平台不接受保存或一次性授权参数")

    if args.inspect_only:
        if len(selected) != 1 or not selected[0].supports_inspect:
            raise SystemExit("--inspect-only 只允许单独选择支持发现的平台")
        if args.save or args.save_only:
            raise SystemExit("--inspect-only 必须同时使用 --no-save")
    elif scaffold and not args.dry_run:
        raise SystemExit("脚手平台只允许 --inspect-only --no-save 或 --dry-run")

    if (
        args.platform == "taobao"
        and args.save
        and not args.dry_run
        and not args.allow_taobao_save_once
        and not args.allow_taobao_publish_once
    ):
        raise SystemExit(
            "淘宝流程仍在分步开发；预览请使用 "
            "--platform taobao --no-save，仅有本次获得明确授权时"
            "才可追加 --allow-taobao-save-once"
        )
    return selected


def learning_execution_mode(args: argparse.Namespace) -> str:
    if getattr(args, "dry_run", False):
        return "dry_run"
    if getattr(args, "inspect_only", False):
        return "inspect_only"
    if not args.save:
        return "preview"
    if getattr(args, "save_only", False) or args.platform == "base":
        return "save_only"
    return "save_and_publish"


def load_stage_readback(
    platform_id: str,
    artifact_dir: Path,
) -> Optional[Mapping[str, Any]]:
    validation_path = artifact_dir / f"{platform_id}-after-save-validation.json"
    if not validation_path.is_file():
        return None
    try:
        payload = json.loads(validation_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    if not isinstance(payload, dict):
        return None
    if payload.get("verified") is False or payload.get("status") in {
        "failed",
        "partial",
        "mismatch",
    }:
        return None
    return payload


def learning_stage_result(
    context: LearningRunContext,
    args: argparse.Namespace,
    platform_id: str,
    artifact_dir: Path,
) -> StageResult:
    if not args.save:
        return StageResult(
            context.run_id,
            platform_id,
            "previewed",
            {},
            {},
            False,
        )
    readback = load_stage_readback(platform_id, artifact_dir)
    if readback is None:
        return StageResult(
            context.run_id,
            platform_id,
            "saved_unverified",
            {},
            {},
            False,
        )
    return StageResult(
        context.run_id,
        platform_id,
        "readback_verified",
        {},
        readback,
        True,
    )


def save_learning_checkpoint(
    context: LearningRunContext,
    args: argparse.Namespace,
    platform_order: Tuple[str, ...],
    current_index: int,
    status: str,
    *,
    pending_review_id: Optional[str] = None,
) -> RunCheckpoint:
    requested = RunCheckpoint(
        context.run_id,
        context.product_version,
        learning_execution_mode(args),
        platform_order,
        current_index,
        status,
        pending_review_id,
        device_id=context.device_id,
        image_version=context.image_version,
    )
    stored = context.store.save_checkpoint(requested)
    persisted = stored if isinstance(stored, RunCheckpoint) else requested
    context.store.enqueue(
        f"checkpoint.updated:{persisted.run_id}:{persisted.version}",
        "checkpoint.updated",
        checkpoint_event_payload(persisted),
    )
    return persisted


def record_learning_stage(
    context: LearningRunContext,
    result: StageResult,
) -> None:
    context.store.record_stage(result)
    if not result.verified:
        return
    runtime = getattr(context, "attribute_runtime", None)
    attributes = (
        result.readback.get("attributes")
        if isinstance(result.readback, Mapping)
        else None
    )
    if runtime is not None and isinstance(attributes, Mapping):
        runtime.record_verified_readbacks(result.platform_id, attributes)
    payload = {
        "run_id": result.run_id,
        "product_version": context.product_version,
        "platform_id": result.platform_id,
        "status": result.status,
        "expected_json": result.expected,
        "readback_json": result.readback,
        "verified": True,
    }
    content_version = canonical_sha256(payload)
    context.store.enqueue(
        f"stage.completed:{result.run_id}:{result.platform_id}:{content_version}",
        "stage.completed",
        payload,
    )


async def drain_learning_events(
    context: Optional[LearningRunContext],
) -> None:
    """Drain background learning events at a durable platform boundary."""
    runtime = (
        getattr(context, "attribute_runtime", None)
        if context is not None
        else None
    )
    if runtime is not None:
        await runtime.drain()


def review_requirements(
    error: ReviewRequired | ReviewBatchRequired,
) -> Tuple[ReviewRequired, ...]:
    if isinstance(error, ReviewBatchRequired):
        return error.reviews
    return (error,)


async def wait_for_platform_review_batch(
    learning_context: LearningRunContext,
    args: argparse.Namespace,
    platform_order: Sequence[str],
    platform_index: int,
    reviews: Sequence[ReviewRequired],
    logger: logging.Logger,
) -> RunCheckpoint:
    """Consume a platform's review decisions in any operator-confirmed order."""
    if learning_context.client is None or not reviews:
        raise AutomationError("人工审核批次缺少审核服务或审核项")
    pending = {review.review_id: review for review in reviews}
    checkpoint = save_learning_checkpoint(
        learning_context,
        args,
        platform_order,
        platform_index,
        "waiting_review",
        pending_review_id=next(iter(pending)),
    )
    await drain_learning_events(learning_context)
    while pending:
        response = await asyncio.to_thread(
            learning_context.client.poll_resume,
            learning_context.device_id,
            30,
        )
        events = response.get("events", ())
        if not isinstance(events, (list, tuple)):
            raise AutomationError("人工审核恢复响应格式无效")
        matched = False
        for event in events:
            if not isinstance(event, Mapping):
                continue
            payload = event.get("payload")
            if not isinstance(payload, Mapping):
                payload = event
            review_id = str(payload.get("review_id") or "")
            review = pending.get(review_id)
            if review is None:
                continue
            checkpoint = save_learning_checkpoint(
                learning_context,
                args,
                platform_order,
                platform_index,
                "waiting_review",
                pending_review_id=review_id,
            )
            decision = validate_resume(
                checkpoint,
                event,
                learning_context.product_version,
                learning_context.image_version,
                learning_execution_mode(args),
            )
            await drain_learning_events(learning_context)
            checkpoint = await persist_resume_and_ack(
                learning_context.store,
                learning_context.client,
                checkpoint,
                decision,
            )
            pending.pop(review_id, None)
            matched = True
            logger.info(
                "运营审核已确认：%s（本批次剩余 %s 项）",
                review.request.field_label,
                len(pending),
            )
        if not matched and events:
            # Another run can share the same device event queue. Avoid a tight
            # loop while this batch's operator decisions are still pending.
            await asyncio.sleep(0.25)
    return save_learning_checkpoint(
        learning_context,
        args,
        platform_order,
        platform_index,
        "running",
    )


async def run_single_platform_with_learning(
    args: argparse.Namespace,
    product: ProductData,
    artifact_dir: Path,
    logger: logging.Logger,
    *,
    learning_context: Optional[LearningRunContext] = None,
    **browser_kwargs: Any,
) -> None:
    if learning_context is None:
        await run_browser_automation(
            args,
            product,
            artifact_dir,
            logger,
            **browser_kwargs,
        )
        return
    platform_order = (args.platform,)
    await initialize_learning_run(learning_context, product, logger)
    checkpoint = save_learning_checkpoint(
        learning_context, args, platform_order, 0, "running"
    )
    while True:
        try:
            await run_browser_automation(
                args,
                product,
                artifact_dir,
                logger,
                attribute_runtime=getattr(
                    learning_context, "attribute_runtime", None
                ),
                **browser_kwargs,
            )
            break
        except (ReviewRequired, ReviewBatchRequired) as exc:
            if learning_context.client is None:
                save_learning_checkpoint(
                    learning_context, args, platform_order, 0, "failed"
                )
                raise
            reviews = review_requirements(exc)
            logger.info(
                "平台 %s 已完成可确定字段，共 %s 项等待人工审核：%s；"
                "全部确认后将自动重开同一平台",
                args.platform,
                len(reviews),
                "、".join(review.request.field_label for review in reviews),
            )
            checkpoint = await wait_for_platform_review_batch(
                learning_context,
                args,
                platform_order,
                0,
                reviews,
                logger,
            )
            logger.info(
                "本平台全部审核项已确认，正在重新打开 %s 并重新获取候选",
                args.platform,
            )
            continue
        except Exception:
            save_learning_checkpoint(
                learning_context, args, platform_order, 0, "failed"
            )
            raise
    record_learning_stage(
        learning_context,
        learning_stage_result(
            learning_context,
            args,
            args.platform,
            artifact_dir,
        )
    )
    save_learning_checkpoint(
        learning_context, args, platform_order, 1, "completed"
    )
    await drain_learning_events(learning_context)


def platform_execution_stages(
    args: argparse.Namespace, product: ProductData,
) -> Tuple[str, ...]:
    """Custom selections are exact; only all mode adds a base-save stage."""
    selected = expand_platform_selection(args.platform)
    if args.platform != "all":
        return tuple(spec.cli_name for spec in selected)
    commerce_stages = (
        tuple(spec.cli_name for spec in selected
              if spec.cli_name != "douyin" or product.douyin_fields is not None)
    )
    start_at = getattr(args, "all_platform_start_at", None)
    if start_at is not None:
        if start_at not in commerce_stages:
            raise AutomationError(f"全平台恢复起点不可用：{start_at}")
        commerce_stages = commerce_stages[commerce_stages.index(start_at) :]
    skipped_platforms = set(getattr(args, "all_platform_skip", ()) or ())
    if skipped_platforms:
        commerce_stages = tuple(name for name in commerce_stages if name not in skipped_platforms)
    if not commerce_stages:
        raise AutomationError("本次全平台流程没有可运行的平台")
    # 预览模式承诺不保存，因此不运行会强制保存的基础资料阶段。
    # 只有全平台的两种保存模式会把它作为第一阶段。
    return (
        (("base",) + commerce_stages)
        if args.save and start_at is None
        else commerce_stages
    )


async def run_all_implemented_platforms(
    args: argparse.Namespace,
    product: ProductData,
    artifact_dir: Path,
    logger: logging.Logger,
    learning_context: Optional[LearningRunContext] = None,
) -> None:
    """在同一编辑页内依次运行全平台或用户指定的平台组合。"""
    stages = platform_execution_stages(args, product)
    logger.info("本次平台执行顺序：%s", " → ".join(stages))
    stage_results: List[Dict[str, Any]] = []
    shared_session: Dict[str, Any] = {}

    if learning_context is not None:
        await initialize_learning_run(learning_context, product, logger)
        save_learning_checkpoint(
            learning_context,
            args,
            stages,
            0,
            "running",
        )

    try:
        stage_index = 0
        while stage_index < len(stages):
            platform_name = stages[stage_index]
            stage_dir = artifact_dir / platform_name
            stage_dir.mkdir(parents=True, exist_ok=True)
            stage_args = argparse.Namespace(**vars(args))
            stage_args.platform = platform_name
            stage_args.taobao_publish_preview = False
            stage_args.youzan_publish_preview = False
            stage_args.wxsph_publish_preview = False
            stage_args.jd_publish_preview = False
            stage_args.allow_taobao_save_once = False
            stage_args.allow_taobao_publish_once = False
            if platform_name == "taobao" and args.save:
                if args.save_only:
                    stage_args.allow_taobao_save_once = True
                else:
                    stage_args.allow_taobao_publish_once = True
            stage_mode = (
                "保存基础资料"
                if platform_name == "base"
                else (
                    "平台仅填写校验"
                    if not stage_args.save
                    else ("只保存" if stage_args.save_only else "保存并铺货")
                )
            )
            logger.info(
                "多平台流程开始：%s（%s）",
                platform_name,
                stage_mode,
            )
            try:
                await run_browser_automation(
                    stage_args,
                    product,
                    stage_dir,
                    logger,
                    shared_session=shared_session,
                    attribute_runtime=(
                        getattr(learning_context, "attribute_runtime", None)
                        if learning_context is not None
                        else None
                    ),
                )
            except (ReviewRequired, ReviewBatchRequired) as exc:
                if learning_context is None or learning_context.client is None:
                    raise
                reviews = review_requirements(exc)
                stage_results.append(
                    {
                        "platform": platform_name,
                        "status": "waiting_review",
                        "review_ids": [review.review_id for review in reviews],
                        "fields": [
                            review.request.field_label for review in reviews
                        ],
                    }
                )
                (artifact_dir / "all-platform-result.json").write_text(
                    json.dumps(stage_results, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                logger.info(
                    "平台 %s 已完成可确定字段，共 %s 项等待人工审核：%s；"
                    "关闭当前浏览器会话",
                    platform_name,
                    len(reviews),
                    "、".join(review.request.field_label for review in reviews),
                )
                await close_shared_browser_session(shared_session)
                shared_session = {}
                await wait_for_platform_review_batch(
                    learning_context,
                    args,
                    stages,
                    stage_index,
                    reviews,
                    logger,
                )
                logger.info(
                    "本平台全部审核项已确认，重新打开 %s 并重新获取候选；"
                    "此前平台不重跑",
                    platform_name,
                )
                continue
            except Exception as exc:
                stage_results.append(
                    {"platform": platform_name, "status": "failed", "error": str(exc)}
                )
                (artifact_dir / "all-platform-result.json").write_text(
                    json.dumps(stage_results, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                if learning_context is not None:
                    save_learning_checkpoint(
                        learning_context,
                        args,
                        stages,
                        stage_index,
                        "failed",
                    )
                raise
            stage_results.append({"platform": platform_name, "status": "success"})
            (artifact_dir / "all-platform-result.json").write_text(
                json.dumps(stage_results, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            if learning_context is not None:
                record_learning_stage(
                    learning_context,
                    learning_stage_result(
                        learning_context,
                        stage_args,
                        platform_name,
                        stage_dir,
                    )
                )
                save_learning_checkpoint(
                    learning_context,
                    args,
                    stages,
                    stage_index + 1,
                    "completed" if stage_index + 1 == len(stages) else "running",
                )
                await drain_learning_events(learning_context)
            logger.info(
                "多平台流程完成：%s；下一平台将复用当前编辑页，"
                "如抽屉已关闭则自动重开",
                platform_name,
            )
            stage_index += 1
    finally:
        await close_shared_browser_session(shared_session)


def main() -> int:
    args = build_parser().parse_args()
    selected_platforms = validate_execution_mode(args)
    if args.publish_shop is None:
        args.publish_shop = list(DEFAULT_DOUYIN_PUBLISH_SHOPS)
    if args.all_platform_one_shop_test:
        args.publish_shop = args.publish_shop[:1]
    timestamp = time.strftime("%Y%m%d-%H%M%S")
    if args.inspect_only:
        artifact_dir = SCRIPT_DIR / "output/platform-schema" / timestamp
        redactor: Optional[SensitiveLogRedactor] = SensitiveLogRedactor()
        logger = setup_logging(artifact_dir, redactor=redactor)
    else:
        artifact_dir = SCRIPT_DIR / "output/kuaimai/runs" / timestamp
        redactor = None
        logger = setup_logging(artifact_dir)
    if args.all_platform_one_shop_test:
        logger.info(
            "已启用全平台单店测试：每平台仅使用正式配置的第一家店，默认配置未修改"
        )
    learning_context: Optional[LearningRunContext] = None
    try:
        if args.create_product:
            product = read_new_product_seed(resolve_excel_path(args.excel_url))
        else:
            product = read_product_data(
                resolve_excel_path(args.excel_url),
                include_douyin=any(
                    spec.cli_name in {"douyin", "taobao"} for spec in selected_platforms
                ),
            )
        learning_context = create_learning_context(args, product)
        if redactor is not None:
            redactor.add_sensitive_values(product.style_code, product.title)
            summary = None
            logger.info("inspect-only 输入校验通过")
        else:
            summary = product_summary(product)
            if args.create_product:
                summary["new_product"] = {
                    "title": "1", "colors": list(product.colors), "sizes": list(NEW_PRODUCT_SIZES),
                    "main_image": str(product.main_images[0]),
                    "sku_count": len(product.colors) * len(NEW_PRODUCT_SIZES),
                    "price": "0", "sku_images_uploaded": False,
                    "code_rule": "款式编码+规格值", "publish": False,
                }
            (artifact_dir / "input-summary.json").write_text(
                json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            logger.info(
                "输入校验通过：%s | 售价 %s | 1:1 %s 张 | 3:4 %s 张 | 详情 %s 张 | SKU %s 张",
                product.style_code,
                product.base_price,
                len(product.main_images),
                len(product.main_images_34),
                len(product.detail_images),
                len(product.sku_images),
            )
        if args.dry_run:
            assert summary is not None
            print(json.dumps(summary, ensure_ascii=False, indent=2))
            logger.info("dry-run 完成，未打开浏览器")
            return 0
        if redactor is None:
            if args.platform == "all" or len(selected_platforms) > 1:
                asyncio.run(
                    run_all_implemented_platforms(
                        args,
                        product,
                        artifact_dir,
                        logger,
                        learning_context=learning_context,
                    )
                )
            else:
                asyncio.run(
                    run_single_platform_with_learning(
                        args,
                        product,
                        artifact_dir,
                        logger,
                        learning_context=learning_context,
                    )
                )
        else:
            asyncio.run(
                run_single_platform_with_learning(
                    args,
                    product,
                    artifact_dir,
                    logger,
                    learning_context=learning_context,
                    redactor=redactor,
                )
            )
        return 0
    except KeyboardInterrupt:
        logger.error("用户中止了程序")
        return 130
    except (
        AutomationError,
        DouyinDataError,
        DouyinListingError,
        TaobaoListingError,
        TmallDataError,
        TmallFormListingError,
        TmallSizeSourceError,
        PlatformDiscoveryError,
        WxsphFormListingError,
        XhsDataError,
        XhsFormListingError,
        RecognitionError,
    ) as exc:
        logger.error("%s", exc)
        return 2
    except Exception:
        logger.exception("程序出现未预期错误")
        return 1
    finally:
        if learning_context is not None:
            learning_context.store.close()


if __name__ == "__main__":
    raise SystemExit(main())
